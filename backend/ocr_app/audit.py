"""Transactional business audit, separate from serious-only diagnostic logs."""
import base64
from contextvars import ContextVar
from datetime import date, datetime, timedelta, timezone
import io
import json
import re
import uuid

from sqlalchemy import text

from .task_errors import TaskError

CONTEXT = ContextVar('operation_audit_context', default=None)
SHANGHAI = timezone(timedelta(hours=8))
EVENTS = {
    'login': '登录', 'logout': '退出登录', 'session_replaced': '新登录使旧会话失效',
    'account_created': '创建账号', 'account_configured': '维护账号', 'account_enabled': '启用账号',
    'account_disabled': '禁用账号', 'account_deleted': '删除账号', 'account_name_changed': '修改账号姓名',
    'password_reset': '管理员重置密码', 'password_changed': '修改本人密码',
    'task_created': '上传材料', 'package_appended': '追加材料', 'package_person_confirmed': '确认人员归属',
    'human_review': '保存修正', 'cancel_requested': '请求取消识别', 'task_rerun': '重新识别',
    'task_deleted': '删除任务', 'batch_deleted': '删除批次', 'task_expired': '到期清理材料',
    'preview_full': '查看原图', 'reveal': '查看字段', 'field_region_ocr': '框选识别',
    'export_full': '请求导出材料', 'export_generated': '导出文件已生成',
    'required_fields_changed': '修改必填字段', 'retention_changed': '修改文件保留期限',
    'workbench_default_changed': '修改默认布局', 'audit_export': '导出审计记录',
    'upload': '上传材料', 'account_update': '修改账号',
    'task_claimed': '开始识别', 'task_completed': '识别完成', 'task_need_review': '识别待复核',
    'task_failed': '识别失败', 'task_rejected': '材料拒收', 'task_cancelled': '识别已取消',
    'task_interrupted': '识别中断', 'task_stop_requested': '请求停止识别',
}
OUTCOMES = {'success': '成功', 'failure': '失败', 'requested': '已请求', 'legacy': '历史记录，结果未确认'}
# Keep legacy event names readable, but apply this scope at both write and read boundaries.
ACCOUNT_EVENTS = {key: EVENTS[key] for key in (
    'login', 'logout', 'account_created', 'account_configured', 'account_enabled',
    'account_disabled', 'account_deleted', 'account_name_changed', 'password_reset', 'password_changed',
    'upload', 'task_created', 'package_appended', 'cancel_requested', 'task_rerun', 'batch_deleted', 'task_deleted',
    'export_full', 'export_generated', 'required_fields_changed', 'retention_changed',
    'workbench_default_changed', 'audit_export', 'account_update')}
ACCOUNT_EVENTS.update(export_full='导出表格', export_generated='导出表格')
EVENT_GROUPS = {'upload': ('upload', 'task_created'), 'export_generated': ('export_generated', 'export_full')}
VISIBLE_SQL = ("actor NOT IN ('worker','system') AND (actor<>'anonymous' OR event='login') "
               "AND event IN (" + ','.join("'" + key + "'" for key in ACCOUNT_EVENTS) + ") "
               "AND outcome<>'requested'")
ENDPOINT_EVENTS = {
    'login': 'login', 'logout': 'logout', 'create_user': 'account_created', 'update_user': 'account_update',
    'delete_user': 'account_deleted', 'reset_user_password': 'password_reset', 'change_password': 'password_changed',
    'upload': 'upload', 'append_upload': 'package_appended', 'confirm_person': 'package_person_confirmed',
    'review': 'human_review', 'cancel': 'cancel_requested', 'rerun': 'task_rerun', 'delete_batch': 'batch_deleted',
    'delete_task': 'task_deleted',
    'preview': 'preview_full', 'full_preview': 'preview_full', 'reveal': 'reveal', 'region_ocr': 'field_region_ocr',
    'export': 'export_full', 'update_required_fields': 'required_fields_changed',
    'update_retention_settings': 'retention_changed', 'update_workbench_default': 'workbench_default_changed',
    'audit_export': 'audit_export',
}
DETAIL_KEYS = {
    'result_revision', 'base_revision', 'status', 'new_run_id', 'code', 'batch_id', 'task_count',
    'added_tasks', 'page', 'field_key', 'fields', 'required_fields', 'old_required_fields', 'required_fields_revision',
    'doc_type', 'old_days', 'days', 'retention_revision', 'legacy_masked', 'revision', 'layout',
    'old_layout', 'old_enabled', 'enabled', 'revoked_sessions', 'records', 'http_status', 'batch_ids', 'doc_types',
}


def safe_details(values):
    def safe(value):
        return value is None or type(value) in (bool, int) or (
            isinstance(value, str) and re.fullmatch(r'[A-Za-z0-9_.:-]{0,100}', value) is not None)
    return {k: v for k, v in values.items() if k in DETAIL_KEYS and
            (safe(v) or isinstance(v, list) and len(v) <= (1000 if k == 'batch_ids' else 200) and all(safe(item) for item in v))}


def selection_details(rows):
    batches = list(dict.fromkeys(row['batch_id'] for row in rows))
    return dict(batch_id=batches[0] if len(batches) == 1 else None, batch_ids=batches,
                task_count=len(rows), doc_types=sorted({row['doc_type'] for row in rows}))


def batch_labels(connection, storage, identifiers):
    labels = []
    for identifier in identifiers:
        row = connection.execute(text('SELECT name_enc,created_at FROM batches WHERE id=:id'), {'id': identifier}).first()
        label = identifier
        if row:
            name = storage.unpack(row.name_enc, '')
            from .doc_profiles import LABELS
            kinds = connection.execute(text('SELECT DISTINCT doc_type FROM tasks WHERE batch_id=:id ORDER BY doc_type'), {'id': identifier}).scalars().all()
            if kinds:
                name = ' / '.join(LABELS.get(kind, kind) for kind in kinds)
            stamp = datetime.fromisoformat(row.created_at).astimezone(SHANGHAI).strftime('%Y-%m-%d %H:%M')
            label = f'{name} · {stamp}'
        labels.append(dict(id=identifier, label=label))
    return labels


def snapshot(row):
    return {key: row[key] for key in ('username', 'display_name', 'role') if key in row}


def bind_actor(user):
    context = CONTEXT.get()
    if context is not None:
        context['actor'] = user['user_id']
        context['identity'] = snapshot(user)


def append(connection, storage, event, actor='worker', task=None, *, outcome='success',
           target_type=None, target_id=None, identity_snapshot=None, private=None, **details):
    if event not in EVENTS or outcome not in OUTCOMES:
        raise ValueError('INVALID_AUDIT_EVENT')
    if (event not in ACCOUNT_EVENTS or actor in {'worker', 'system'} or outcome == 'requested'
            or actor == 'anonymous' and event != 'login'):
        return None
    context = CONTEXT.get() or {}
    row = connection.execute(text('SELECT username,display_name,role FROM web_users WHERE id=:id'),
                             {'id': actor}).mappings().first()
    actor_snapshot = identity_snapshot if identity_snapshot is not None else snapshot(row) if row else (
        context.get('identity') if context.get('actor') == actor else None)
    batch_id = details.get('batch_id') or (task.get('batch_id') if task else None)
    if task and not batch_id:
        batch_id = connection.execute(text('SELECT batch_id FROM tasks WHERE id=:id'), {'id': task['id']}).scalar()
    target_type = target_type or ('task' if task else 'batch' if batch_id else 'system')
    target_id = target_id or (task['id'] if task else batch_id)
    encrypted_context = {key: context[key] for key in ('peer', 'claimed_username') if context.get(key)}
    if private:
        encrypted_context.update(private)
    batches = details.get('batch_ids') or ([batch_id] if batch_id else [])
    if batches:
        encrypted_context['batches'] = batch_labels(connection, storage, batches)
    identifier = str(uuid.uuid4())
    connection.execute(text('''INSERT INTO audit
        (id,actor,task_id,run_id,event,created_at,detail_json,version,outcome,batch_id,
         target_type,target_id,request_id,actor_snapshot_enc,context_enc)
        VALUES (:id,:actor,:task,:run,:event,:at,:details,1,:outcome,:batch,:kind,:target,:request,:identity,:context)'''),
        dict(id=identifier, actor=actor, task=task['id'] if task else None, run=task['run_id'] if task else None,
             event=event, at=datetime.now(timezone.utc).isoformat(), details=json.dumps(safe_details(details)),
             outcome=outcome, batch=batch_id, kind=target_type, target=target_id,
             request=context.get('request_id'), identity=storage.pack(actor_snapshot) if actor_snapshot else None,
             context=storage.pack(encrypted_context) if encrypted_context else None))
    return identifier


def failure(business, request, code, status):
    """Runs only after the request's business transaction has unwound."""
    event = ENDPOINT_EVENTS.get(getattr(request.scope.get('endpoint'), '__name__', ''))
    if event not in ACCOUNT_EVENTS:
        return
    context = CONTEXT.get() or {}
    if context.get('failure_recorded'):
        return
    context['failure_recorded'] = True
    params = request.path_params
    # IDs come from validated route converters; never store arbitrary URL strings.
    ids = {}
    for key in ('task_id', 'batch_id', 'user_id'):
        try:
            ids[key] = str(uuid.UUID(str(params[key])))
        except (KeyError, ValueError, TypeError):
            pass
    try:
        with business.store.transaction() as connection:
            # Export has no path ID. The endpoint supplies only validated selection IDs.
            if event == 'export_full':
                targets = context.get('export_targets', [ids])
                batches = []
                for target in targets:
                    batch = target.get('batch_id') or connection.execute(text('SELECT batch_id FROM tasks WHERE id=:id'),
                        {'id': target.get('task_id')}).scalar()
                    if batch and batch not in batches:
                        batches.append(batch)
                append(connection, business.storage, event, context.get('actor', 'anonymous'), outcome='failure',
                       batch_id=batches[0] if len(batches) == 1 else None, batch_ids=batches, code=code, http_status=status)
            else:
                task = connection.execute(text('SELECT id,run_id,batch_id FROM tasks WHERE id=:id'), {'id': ids.get('task_id')}).mappings().first()
                append(connection, business.storage, event, context.get('actor', 'anonymous'), task,
                       outcome='failure', target_type='account' if ids.get('user_id') else None,
                       target_id=ids.get('user_id'), batch_id=ids.get('batch_id'), code=code, http_status=status)
    except Exception as error:
        if business.logger:
            business.logger.exception('storage_failed', error)


class AuditService:
    def __init__(self, engine, storage):
        self.engine, self.storage = engine, storage

    def options(self):
        with self.engine.connect() as c:
            rows = c.execute(text('''SELECT a.* FROM (SELECT DISTINCT actor FROM audit WHERE ''' + VISIBLE_SQL + ''') actors
                JOIN audit a ON a.id=(SELECT b.id FROM audit b WHERE b.actor=actors.actor
                                      AND ''' + VISIBLE_SQL + ''' ORDER BY b.created_at DESC,b.id DESC LIMIT 1)''')).mappings().all()
        actors = [dict(id=row['actor'], **self.storage.unpack(row['actor_snapshot_enc'],
                     {'display_name': actor_label(row), 'username': '', 'role': ''})) for row in rows]
        return dict(actors=actors, events=[dict(value=k, label=v) for k, v in ACCOUNT_EVENTS.items() if k not in {'task_created', 'export_full'}],
                    outcomes=[dict(value=k, label=v) for k, v in OUTCOMES.items() if k != 'requested'])

    def query(self, *, date_from=None, date_to=None, actor=None, event=None, outcome=None, batch_id=None,
              cursor=None, limit=50, page=None, page_size=None):
        numbered = page is not None or page_size is not None
        if numbered:
            page = 1 if page is None else page
            page_size = 50 if page_size is None else page_size
            if (cursor is not None or type(page) is not int or not 1 <= page <= 2147483647
                    or type(page_size) is not int or not 1 <= page_size <= 100):
                raise TaskError('AUDIT_FILTER_INVALID', '审计日期或分页参数无效', 422)
            limit = page_size
        today = datetime.now(SHANGHAI).date()
        try:
            start = datetime.combine(date.fromisoformat(date_from) if date_from else today - timedelta(days=6),
                                     datetime.min.time(), SHANGHAI)
            end = datetime.combine(date.fromisoformat(date_to) if date_to else today,
                                   datetime.min.time(), SHANGHAI) + timedelta(days=1)
            if start >= end or not 1 <= limit <= 10001:
                raise ValueError()
        except (TypeError, ValueError, OverflowError):
            raise TaskError('AUDIT_FILTER_INVALID', '审计日期或分页参数无效', 422) from None
        clauses = [VISIBLE_SQL, 'created_at>=:start', 'created_at<:end']
        params = dict(start=start.astimezone(timezone.utc).isoformat(), end=end.astimezone(timezone.utc).isoformat())
        for name, value in dict(actor=actor, event=event, outcome=outcome, batch_id=batch_id).items():
            if value is not None:
                if not isinstance(value, str) or len(value) > 100 or name == 'event' and value not in EVENTS or name == 'outcome' and value not in OUTCOMES:
                    raise TaskError('AUDIT_FILTER_INVALID', '审计筛选条件无效', 422)
                if name == 'event' and value in EVENT_GROUPS:
                    clauses.append('event IN (' + ','.join("'" + item + "'" for item in EVENT_GROUPS[value]) + ')')
                    continue
                clauses.append('(batch_id=:batch_id OR EXISTS (SELECT 1 FROM json_each(audit.detail_json, \'$.batch_ids\') WHERE value=:batch_id) '
                               'OR (version=0 AND task_id IN (SELECT id FROM tasks WHERE batch_id=:batch_id)))'
                               if name == 'batch_id' else name + '=:' + name)
                params[name] = value
        count_where = ' AND '.join(clauses)
        if cursor:
            try:
                values = json.loads(base64.urlsafe_b64decode(cursor + '=' * (-len(cursor) % 4)))
                if not isinstance(values, list) or len(values) != 2 or not all(isinstance(v, str) for v in values):
                    raise ValueError()
                stamp, identifier = values
                if datetime.fromisoformat(stamp).tzinfo is None:
                    raise ValueError()
                uuid.UUID(identifier)
            except (ValueError, TypeError, KeyError):
                raise TaskError('AUDIT_FILTER_INVALID', '审计分页参数无效', 422) from None
            clauses.append('(created_at<:before OR (created_at=:before AND id<:before_id))')
            params.update(before=stamp, before_id=identifier)
        params['limit'] = limit + 1
        params['offset'] = (page - 1) * limit if numbered else 0
        with self.engine.connect() as c:
            # SQLite legacy SELECT mode otherwise gives COUNT and rows different snapshots.
            c.exec_driver_sql('BEGIN')
            total = c.execute(text('SELECT COUNT(*) FROM audit WHERE ' + count_where), params).scalar_one()
            rows = c.execute(text('SELECT *, COALESCE(batch_id,(SELECT batch_id FROM tasks WHERE tasks.id=audit.task_id)) AS resolved_batch_id FROM audit WHERE ' + ' AND '.join(clauses) +
                                  ' ORDER BY created_at DESC,id DESC LIMIT :limit OFFSET :offset'), params).mappings().all()
        selected = rows[:limit]
        next_cursor = base64.urlsafe_b64encode(json.dumps([selected[-1]['created_at'], selected[-1]['id']]).encode()).decode().rstrip('=') if len(rows) > limit else None
        result = dict(items=[self.present(row) for row in selected], next_cursor=next_cursor, total=total)
        if numbered:
            result.update(page=page, page_size=page_size)
        return result

    def present(self, row):
        identity = self.storage.unpack(row['actor_snapshot_enc'], {})
        context = self.storage.unpack(row['context_enc'], {})
        details = safe_details(json.loads(row['detail_json']))
        batches = context.get('batches', [])
        batch_id = row.get('resolved_batch_id', row['batch_id'])
        if not batches and batch_id:
            with self.engine.connect() as connection:
                batches = batch_labels(connection, self.storage, [batch_id])
        target = context.get('target', {})
        object_label = ('；'.join(item['label'] for item in batches) if batches else
                        f"{target.get('display_name', '')}（{target['username']}）" if target.get('username') else
                        '系统设置' if row['event'] in {'required_fields_changed', 'retention_changed', 'workbench_default_changed'} else
                        context.get('claimed_username') or identity.get('username') or row['target_id'] or '')
        return dict(id=row['id'], actor=row['actor'], identity=identity, event=row['event'],
                    description=description(row['event'], details, row['outcome']), object_label=object_label, outcome=row['outcome'],
                    outcome_label=OUTCOMES.get(row['outcome'], OUTCOMES['legacy']), created_at=row['created_at'],
                    task_id=row['task_id'], run_id=row['run_id'], batch_id=row.get('resolved_batch_id', row['batch_id']),
                    target_type=row['target_type'], target_id=row['target_id'], request_id=row['request_id'],
                    details=details, legacy=row['version'] == 0)

    def detail(self, identifier):
        with self.engine.connect() as c:
            row = c.execute(text('SELECT *, COALESCE(batch_id,(SELECT batch_id FROM tasks WHERE tasks.id=audit.task_id)) AS resolved_batch_id FROM audit WHERE id=:id AND ' + VISIBLE_SQL), {'id': identifier}).mappings().first()
            if row is None:
                raise TaskError('NOT_FOUND', '审计记录不存在', 404)
            result = self.present(row)
            result['context'] = self.storage.unpack(row['context_enc'], {})
            result.update(changes=[], history_state='not_applicable')
        return result

    def workbook(self, **filters):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        page = self.query(limit=10000, **filters)
        if page['next_cursor']:
            raise TaskError('AUDIT_EXPORT_TOO_LARGE', '审计导出最多10000条，请缩小日期范围或增加筛选条件', 422)
        book = Workbook()
        sheet = book.active
        sheet.title = '操作审计'
        sheet.append(['时间（UTC+8）', '姓名', '账号', '操作', '关联对象', '结果'])
        for row in page['items']:
            identity = row['identity']
            stamp = datetime.fromisoformat(row['created_at']).astimezone(SHANGHAI).strftime('%Y-%m-%d %H:%M:%S')
            values = [stamp, identity.get('display_name', actor_label(row)), identity.get('username', ''),
                      row['description'], row['object_label'], row['outcome_label']]
            sheet.append(values)
        for row in sheet:
            for cell in row:
                cell.data_type = 's'
                cell.font = Font(name='微软雅黑', size=12, bold=cell.row == 1)
                cell.alignment = Alignment(vertical='center', horizontal='left', wrap_text=True)
                cell.fill = PatternFill('solid', fgColor='EF760E' if cell.row == 1 else 'FFFFFF')
        for column, width in zip('ABCDEF', [25, 20, 22, 65, 55, 32]):
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        output = io.BytesIO()
        book.save(output)
        return output.getvalue(), len(page['items'])


def actor_label(row):
    return {'worker': '系统任务', 'system': '系统维护', 'anonymous': '未登录／身份未验证'}.get(row['actor'], '历史账号（姓名未记录）')


def description(event, details, outcome):
    from .doc_profiles import LABELS
    label = ACCOUNT_EVENTS.get(event, '历史操作')
    if outcome in {'failure', 'legacy'}:
        return label
    if event in {'upload', 'package_appended', 'batch_deleted', 'export_generated'}:
        count = details.get('task_count', details.get('added_tasks'))
        if type(count) is int:
            label += f'（{count}份材料）'
    elif event == 'retention_changed' and 'old_days' in details and 'days' in details:
        days = lambda value: '永不清除' if value is None else f'{value}天'
        label = f"文件保留期限：{days(details['old_days'])} → {days(details['days'])}"
    elif event == 'required_fields_changed':
        from .required_fields import configurable_fields
        kind = details.get('doc_type')
        if kind in LABELS:
            names = {field['key']: field['label'] for field in configurable_fields(kind)}
            selected = details.get('required_fields')
            rendered = '恢复默认' if selected is None else '、'.join(names.get(key, key) for key in selected) or '无'
            label = f"{LABELS[kind]}必填字段：{rendered}"
            if 'old_required_fields' in details:
                old = '、'.join(names.get(key, key) for key in details['old_required_fields']) or '无'
                label = f"{LABELS[kind]}必填字段：{old} → {rendered}"
    elif event == 'workbench_default_changed':
        layout = lambda values: '系统默认' if values is None else '、'.join(LABELS.get(value, value) for value in values) or '全部收纳'
        label = f"默认布局：{layout(details.get('old_layout'))} → {layout(details.get('layout'))}"
    elif event == 'login' and details.get('revoked_sessions'):
        label += '（旧登录已失效）'
    return label
