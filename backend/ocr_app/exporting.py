"""Snapshot-only XLSX serialization; untrusted strings never become formulas."""
from io import BytesIO
from math import ceil

from .task_store import TaskError, public_result


def _field_values(result):
    values = {key: value.get('value') if isinstance(value, dict) else value
              for key, value in result.get('basic_info', {}).items()}
    values.update({field['key']: field.get('value') for field in result.get('fields', [])})
    return values


def _generic(result):
    return result.get('reviewed_generic') if 'reviewed_generic' in result else result.get('generic')


def _missing(value):
    return value is None or isinstance(value, str) and not value.strip()


def export_eligibility(result, required_overrides=None):
    """Return blockers and their disputes from one saved, unfiltered snapshot."""
    from .doc_profiles import profile, member_profile, generic_required, SUPPORTED
    if not result:
        return [dict(field='task', code='MISSING')], []
    kind = result.get('doc_type')
    if kind == 'TABLE':
        from .table_recognition import table_blockers
        return table_blockers(result), []
    if kind not in SUPPORTED:
        return [dict(field='doc_type', code='MISSING')], []
    values = _field_values(result)
    context = {key: value.strip() if isinstance(value, str) else value for key, value in values.items()}
    missing = [field['key'] for field in profile(kind, context, required_overrides)
               if field['required'] and _missing(values.get(field['key']))]
    if kind == 'HOUSEHOLD_REGISTER':
        members = result.get('members') or [{}]
        for index, member in enumerate(members):
            fields = member.get('fields') if isinstance(member.get('fields'), list) else [
                value for value in member.values() if isinstance(value, dict) and 'key' in value]
            member_values = {field['key']: field.get('value') for field in fields}
            missing.extend(f'members[{index}].{field["key"]}' for field in member_profile(required_overrides)
                           if field['required'] and _missing(member_values.get(field['key'])))
    if kind == 'GENERIC' and generic_required(required_overrides) and _missing((_generic(result) or {}).get('full_text')):
        missing.append('generic.full_text')
    from .export_disputes import export_disputes
    disputes = export_disputes(result)
    disputed = dict.fromkeys(d['field'] for d in disputes)
    blockers = [dict(field=key, code='MISSING') for key in missing] + [dict(field=key, code='DISPUTED') for key in disputed]
    return blockers, disputes


def required_export_blockers(result, required_overrides=None):
    """Check required presence and explicit disputes in a saved, unfiltered result."""
    return export_eligibility(result, required_overrides)[0]


def workbook_bytes(results, packages=()):
    from openpyxl import Workbook
    from openpyxl.cell.cell import Cell
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from .doc_profiles import profile, MEMBER_PROFILE, LABELS
    from .medical_ledger import HEADERS, LEDGER_KEYS, abnormal_cell
    book = Workbook()
    book.remove(book.active)
    sheets = {}

    def sheet_for(name, headers):
        if name not in sheets:
            if len(headers) > 16384:
                raise TaskError('EXPORT_VALUE_TOO_LONG', '导出字段列数超过Excel上限，请减少本次选择的材料', 422)
            sheet = sheets[name] = book.create_sheet(name[:31])
            sheet.append(headers)
        return sheets[name]

    def append(name, headers, values):
        sheet = sheet_for(name, headers)
        row = []
        for value in values:
            if isinstance(value, str) and len(value) > 32767:
                raise TaskError('EXPORT_VALUE_TOO_LONG', '某结构化字段超过Excel单元格容量，请先拆分或修正字段', 422)
            cell = Cell(sheet, value=value if value is None or isinstance(value, (str, int, float, bool)) else str(value))
            if isinstance(value, str):
                cell.data_type = 's'
                cell.number_format = '@'
            row.append(cell)
        sheet.append(row)

    def field_values(result, fields):
        values = _field_values(result)
        entries = dict(result.get('basic_info', {}))
        entries.update({f['key']: f for f in result.get('fields', [])})
        row = []
        for field in fields:
            key = field['key']
            value = values.get(key)
            entry = entries.get(key, {})
            if value is None and isinstance(entry, dict):
                value = {'NOT_ON_DOCUMENT': '未载明', 'NOT_APPLICABLE': '不适用'}.get(entry.get('absence_reason'))
                if result['doc_type'] == 'MEDICAL_REPORT' and entry.get('absence_reason') == 'NOT_ON_DOCUMENT':
                    value = '报告未载明'
            if key == 'gender':
                value = {'M': '男', 'F': '女'}.get(value, value)
            row.append(value)
        return row

    from .person_packages import name_suggestions
    package_rows, column_keys, multiplicity = [], {}, {}
    for package in packages:
        selected = package['results']
        data, occurrences = {}, {}
        for result in selected:
            kind = result['doc_type']
            number = occurrences[kind] = occurrences.get(kind, 0) + 1
            multiplicity[kind] = max(multiplicity.get(kind, 0), number)
            current = public_result(result) if kind == 'MEDICAL_REPORT' else result
            fields = profile(kind)
            pairs = list(zip([f['label'] for f in fields], field_values(current, fields)))
            if kind == 'MEDICAL_REPORT':
                pairs.append(('异常指标', abnormal_cell(current)))
            for index, member in enumerate(current.get('members', []), 1):
                member_fields = member.get('fields') if isinstance(member.get('fields'), list) else [
                    value for value in member.values() if isinstance(value, dict) and 'key' in value]
                pairs.extend(zip([f'成员{index}·{f["label"]}' for f in MEMBER_PROFILE],
                    field_values(dict(doc_type=kind, fields=member_fields), MEMBER_PROFILE)))
            if kind == 'GENERIC':
                generic = _generic(current) or {}
                body = generic.get('full_text') or ''
                pairs.extend((f'正文{index // 30000 + 1}', body[index:index + 30000])
                             for index in range(0, max(1, len(body)), 30000))
                for index, pair in enumerate(generic.get('key_values', []), 1):
                    pairs.extend([(f'键值{index}·字段', pair['key']), (f'键值{index}·内容', pair.get('value'))])
            for label, value in pairs:
                key = kind, number, label
                column_keys.setdefault(key, None)
                data[key] = value
        candidates = name_suggestions(selected)
        name = package['person_name']
        conflict = len(candidates) > 1 or bool(name and any(c['name'] != name for c in candidates))
        kinds = list(dict.fromkeys(r['doc_type'] for r in selected))
        metadata = [name or '姓名待确认', ('已确认；识别姓名有差异，需核对' if conflict else '已确认') if name else
                ('待确认；识别姓名有差异，需核对' if conflict else '待确认'),
                '\n'.join(f'{LABELS[k]} × {sum(r["doc_type"] == k for r in selected)}' for k in kinds), len(selected),
                '\n'.join(f'{c["name"]}：' + '；'.join(dict.fromkeys(LABELS[s['doc_type']] for s in c['sources'])) for c in candidates)]
        package_rows.append((metadata, data))
    kind_order = {kind: index for index, kind in enumerate(LABELS)}
    columns = sorted(column_keys, key=lambda key: (kind_order[key[0]], key[1]))
    headers = ['人员姓名', '姓名状态'] + [
        f'{LABELS[kind]}{number if multiplicity[kind] > 1 else ""}·{label}'
        for kind, number, label in columns] + ['材料类型及数量', '材料份数', '姓名候选']
    for metadata, data in package_rows:
        append('人员信息总表', headers, metadata[:2] + [data.get(key) for key in columns] + metadata[2:])

    medical_sequence = 0
    for result in results:
        values = _field_values(result)
        kind = result['doc_type']
        fields = profile(kind)
        if kind == 'MEDICAL_REPORT':
            medical_sequence += 1
            current = public_result(result)
            definitions = {f['key']:f for f in fields}
            append('审核记录台账', HEADERS, [medical_sequence] + field_values(current,[definitions[key] for key in LEDGER_KEYS]) + [abnormal_cell(current)])
            entries = {field['key']: field for field in current['fields']}
            sheet = sheets['审核记录台账']
            for column, key in enumerate(LEDGER_KEYS, 2):
                entry = entries.get(key, {})
                derivation = entry.get('derivation') or {}
                if (entry.get('value') == '已识别明细未见异常' and entry.get('value_origin') == 'derived'
                        and derivation.get('method') == 'medical_ledger'
                        and derivation.get('input_field') == 'items' and not derivation.get('conclusions')):
                    sheet.cell(sheet.max_row, column).comment = Comment(
                        '仅基于已识别明细，不代表整组检查或整份报告已完整识别；请对照原件核对检查范围与完整性。',
                        'HRSSC识别工作台')
        elif kind != 'GENERIC':
            append(LABELS[kind], [f['label'] for f in fields], field_values(result, fields))
        for index, member in enumerate(result.get('members', [])):
            fields = member.get('fields', []) if isinstance(member.get('fields'), list) else [v for v in member.values() if isinstance(v, dict) and 'key' in v]
            field_map = {field['key']: field for field in fields}
            append('户口成员', ['成员序号'] + [f['label'] for f in MEMBER_PROFILE],
                   [index + 1] + [field_map.get(f['key'], {}).get('value') for f in MEMBER_PROFILE])
        if kind == 'GENERIC' and _generic(result):
            generic = _generic(result)
            value = generic.get('full_text') or ''
            for index in range(0, max(1, len(value)), 500):
                append('通用识别文本', ['分段序号', '文本'], [index // 500 + 1, value[index:index + 500]])
            for pair in generic.get('key_values', []):
                append('通用识别字段', ['字段', '内容'], [pair['key'], pair.get('value')])
    for sheet in sheets.values():
        medical = sheet.title == '审核记录台账'
        sheet.freeze_panes = 'C2' if sheet.title == '人员信息总表' else 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.print_title_rows = '1:1'
        sheet.print_options.horizontalCentered = True
        widths = []
        row_count = sheet.max_row
        for index, cell in enumerate(sheet[1], 1):
            label = cell.value.rsplit('·', 1)[-1]
            width = (64 if label == '文本' or label.startswith('正文') else 55 if label in {'体检项目', '经营范围'} else
                     42 if label in {'异常诊断', '异常指标', '检测结果', '内容'} else
                     50 if label in {'材料类型及数量', '姓名候选'} else
                     34 if label in {'学校名称', '名称', '住址', '住所'} else
                     27 if any(term in label for term in ('编号', '号码', '身份证号', '卡号', '代码', '验证码', '异常项目')) else 18)
            if medical:
                width = [8,12,8,8,14,30,14,12,12,20,40,48,40,40,48,40,18,90][index-1]
            else:
                longest = max((max((sum(2 if ord(c) > 255 else 1 for c in line)
                                    for line in str(value if value is not None else '').splitlines()), default=0)
                               for value, in sheet.iter_rows(min_col=index, max_col=index, max_row=row_count, values_only=True)), default=0)
                # Leave room for 12-point text; bound long prose and retain wrapping.
                width = max(width, min(64, ceil(longest * 1.3) + 3))
            sheet.column_dimensions[get_column_letter(index)].width = width
            widths.append(width)
        for index, row in enumerate(sheet, 1):
            lines = 1
            for column, (cell, width) in enumerate(zip(row, widths),1):
                cell.alignment = Alignment(vertical='center' if index==1 or medical else 'top', wrap_text=True,
                    horizontal='center' if medical and (index==1 or column<11 or sheet.cell(1,column).value=='体检项目数量') else 'left')
                size = 12
                cell.font = Font(name='Microsoft YaHei', size=size, bold=index == 1, color='FFFFFF' if index == 1 else '243746')
                color = ('FF6C0D' if 11<=column<=16 else '4472C4') if medical else 'EF760E'
                cell.fill = PatternFill('solid', fgColor=color if index == 1 else 'FFFFFF')
                # Estimate visual line width, accounting for wide Chinese glyphs.
                text_lines = str(cell.value or '').split('\n')
                visual = sum(max(1, ceil(sum(2 if ord(c) > 255 else 1 for c in line) * size / 11 / max(1, width - 2))) for line in text_lines)
                lines = max(lines, visual)
            sheet.row_dimensions[index].height = min(409, max(32, 18 * lines + 10))
    target = BytesIO()
    book.save(target)
    return target.getvalue()


def export_results(store, request, actor, role, *, user=None, authorize=None, with_metadata=False):
    import json
    from sqlalchemy import text
    from .task_store import identity, now
    from .review import optional_reason, invalid
    if not isinstance(request, dict) or set(request) - {'batch_id', 'task_ids', 'versions', 'masked', 'reason', 'format'}:
        invalid()
    if bool(request.get('batch_id')) == bool(request.get('task_ids')) or request.get('format', 'xlsx') != 'xlsx':
        invalid('请选择一个批次或明确的任务列表，当前导出格式为xlsx')
    is_masked = request.get('masked', True)
    if type(is_masked) is not bool:
        invalid()
    if role not in {'operator', 'admin'}:
        raise TaskError('FORBIDDEN_REVEAL', '当前账号无完整导出权限', 403)
    optional_reason(request.get('reason'))
    snapshots, packages, batch_revisions = [], {}, {}
    with store.engine.connect() as connection:
        connection.exec_driver_sql('BEGIN')
        if authorize:
            authorize(connection)
        policy = store.required_policy(connection)
        if request.get('batch_id'):
            store._batch_row(connection, request['batch_id'], user)
            rows = connection.execute(text('SELECT * FROM tasks WHERE batch_id=:id ORDER BY created_at,id'), {'id': request['batch_id']}).mappings().all()
        else:
            ids = request['task_ids']
            if not isinstance(ids, list) or not ids or len(ids) > 1000 or len(set(ids)) != len(ids):
                invalid()
            rows = [store._row(connection, identifier, user) for identifier in ids]
        if not rows:
            invalid('导出选择为空')
        versions = request.get('versions')
        if versions is not None:
            if not isinstance(versions, list) or any(not isinstance(v, dict) or set(v) != {'task_id', 'run_id', 'result_revision'} for v in versions):
                invalid()
            version_map = {v['task_id']: v for v in versions}
            if len(version_map) != len(versions) or set(version_map) != {row['id'] for row in rows}:
                invalid('版本列表必须覆盖全部所选任务')
        blocked = []
        for row in rows:
            if versions is not None:
                version = version_map[row['id']]
                store._expected(row, version['run_id'], version['result_revision'])
            result = store.storage.unpack(row['result_enc'])
            if result is not None:
                result['run_id'] = row['run_id']
            problems = required_export_blockers(result, policy['overrides'])
            if problems:
                blocked.append(dict(task_id=row['id'], status=row['status'], blockers=problems))
                continue
            result.update(task_id=row['id'], run_id=row['run_id'], result_revision=row['revision'], status=row['status'])
            snapshot = public_result(result, policy['overrides'])
            snapshot['_source_filename'] = store.storage.unpack(row['meta_enc']).get('file_name')
            snapshots.append(snapshot)
            batch = store._batch_row(connection, row['batch_id'], user)
            batch_revisions[batch['id']] = batch['package_revision']
            if batch['is_person_package']:
                package = packages.setdefault(batch['id'], dict(
                    person_name=store.storage.unpack(batch['person_name_enc']), results=[]))
                package['results'].append(snapshot)
        if blocked:
            raise TaskError('EXPORT_BLOCKED', '所选任务必填未填、识别有异议，或表格结构及完整性需处理，未导出任何任务', 409, {'tasks': blocked})
        table_export = any(snapshot['doc_type'] == 'TABLE' for snapshot in snapshots)
        if table_export and (packages or any(snapshot['doc_type'] != 'TABLE' for snapshot in snapshots)):
            invalid('表格请单独导出，不能与证件或人员材料包混合')
    if table_export:
        from .table_export import table_export_bundle
        data, media, filename = table_export_bundle(snapshots)
    else:
        data, media, filename = workbook_bytes(snapshots, packages.values()), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'reviewed-results.xlsx'
    prepared = []
    try:
        for row in rows:
            payload = dict(event='export_full', actor=actor, reason=request.get('reason'), at=now(),
                           run_id=row['run_id'], result_revision=row['revision'], legacy_masked=is_masked)
            prepared.append((row, store.storage.put(json.dumps(payload, ensure_ascii=False).encode())))
        with store.transaction() as connection:
            if authorize:
                authorize(connection)
            if store.required_policy(connection)['revision'] != policy['revision']:
                raise TaskError('SETTINGS_VERSION_CONFLICT', '必填设置已变化，请重新核对后导出')
            for batch_id, revision in batch_revisions.items():
                if store._batch_row(connection, batch_id, user)['package_revision'] != revision:
                    raise TaskError('RESULT_VERSION_CONFLICT', '材料包已更新，请刷新后重新导出')
            for row, reference in prepared:
                current = store._row(connection, row['id'], user)
                store._expected(current, row['run_id'], row['revision'])
                if current['status'] != row['status']:
                    raise TaskError('RESULT_VERSION_CONFLICT', '任务状态已变化，请重新核对后导出')
                connection.execute(text("INSERT INTO artifacts(id,task_id,run_id,kind,reference) VALUES (:id,:task,:run,'access_history',:ref)"),
                                   dict(id=identity(), task=row['id'], run=row['run_id'], ref=reference))
            from .audit import selection_details
            store._audit(connection, 'export_generated', actor, **selection_details(rows),
                         required_fields_revision=policy['revision'])
    except BaseException:
        for _, reference in prepared:
            store.storage.delete(reference)
        raise
    return (data, media, filename) if with_metadata else data
