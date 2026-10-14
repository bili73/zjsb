"""Serialize effective photo tables without coercing source cells into formulas."""
from io import BytesIO
from math import ceil
import re
from zipfile import ZipFile, ZIP_DEFLATED

from .table_recognition import effective_tables, table_blockers, validate_table_grid

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def safe_title(title):
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '-', title).strip(' .')[:120]
    if not name or re.fullmatch(r'(?i:CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\..*)?', name):
        return '表格'
    return name


def table_workbook_bytes(table):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    validate_table_grid(table)
    book=Workbook(); sheet=book.active
    sheet.title=re.sub(r'[\[\]:*?/\\]', '-', safe_title(table['title']))[:31].strip("'") or '表格'
    headers=table['header_rows'] or 0
    widths=[]
    for c in range(len(table['rows'][0])):
        desired=max((max((sum(2 if ord(ch)>255 else 1 for ch in line) for line in row[c].split('\n')),default=0)
                     *12/11+3 for r,row in enumerate(table['rows'])),default=12)
        width=max(12,min(55,desired));widths.append(width)
        sheet.column_dimensions[get_column_letter(c+1)].width=width
    border=Border(**{side:Side(style='thin',color='D9E0E5') for side in ('left','right','top','bottom')})
    title=sheet.cell(1,1,table['title']);title.data_type='s';title.number_format='@'
    title.font=Font(name='Microsoft YaHei',size=21,bold=True,color='243746')
    title.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    if len(widths)>1:
        sheet.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(widths))
    title_lines=max(1,ceil(sum(2 if ord(ch)>255 else 1 for ch in table['title'])*21/11/max(1,sum(widths)-2)))
    sheet.row_dimensions[1].height=min(409,max(38,title_lines*28+10))
    for r,row in enumerate(table['rows'],2):
        is_header=r<=headers+1;size=12;lines=1
        for c,value in enumerate(row,1):
            cell=sheet.cell(r,c,value);cell.data_type='s';cell.number_format='@'
            cell.font=Font(name='Microsoft YaHei',size=size,bold=is_header,color='FFFFFF' if is_header else '243746')
            cell.fill=PatternFill('solid',fgColor='EF760E' if is_header else 'FFFFFF')
            cell.alignment=Alignment(horizontal='left',vertical='center',wrap_text=True)
            cell.border=border
            lines=max(lines,sum(max(1,ceil(sum(2 if ord(ch)>255 else 1 for ch in line)*size/11/max(1,widths[c-1]-2))) for line in value.split('\n')))
        sheet.row_dimensions[r].height=min(409,max(32,18*lines+10))
    for merge in table['merges']:
        r,c,h,w=(merge[k] for k in ('row','column','rowspan','colspan'))
        sheet.merge_cells(start_row=r+2,start_column=c+1,end_row=r+h+1,end_column=c+w)
    sheet.freeze_panes=f'A{headers+2}'
    sheet.print_title_rows=f'1:{headers+1}'
    if headers==1 and not table['merges']:
        sheet.auto_filter.ref=f'A2:{get_column_letter(len(widths))}{len(table["rows"])+1}'
    sheet.sheet_view.showGridLines=False
    sheet.page_setup.orientation='landscape';sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    output=BytesIO();book.save(output);return output.getvalue()


def table_export_bundle(results):
    if any(table_blockers(result) for result in results):
        raise ValueError('TABLE_EXPORT_BLOCKED')
    tables=[effective_tables(result)[0] for result in results]
    if len(tables)==1:
        return table_workbook_bytes(tables[0]), XLSX, safe_title(tables[0]['title'])+'.xlsx'
    output=BytesIO();used=set()
    with ZipFile(output,'w',ZIP_DEFLATED) as archive:
        for table in tables:
            stem=safe_title(table['title']);name=stem+'.xlsx';index=2
            while name.casefold() in used:
                name=f'{stem} ({index}).xlsx';index+=1
            used.add(name.casefold());archive.writestr(name,table_workbook_bytes(table))
    return output.getvalue(),'application/zip','表格批量导出.zip'
