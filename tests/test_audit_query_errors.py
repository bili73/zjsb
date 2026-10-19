"""Audit errors remain safe, actionable and distinct from authentication failures."""
from datetime import datetime, timezone
import io
import unittest
import uuid

from openpyxl import load_workbook
from sqlalchemy import text

import test_operation_audit as base


class AuditQueryErrorTests(unittest.TestCase):
    setUp = base.AuditTests.setUp
    account = base.AuditTests.account
    login = base.AuditTests.login
    admin = base.AuditTests.admin
    password = base.AuditTests.password
    origin = base.AuditTests.origin

    def test_query_validation_uses_actionable_filter_code_without_echoing_input(self):
        self.admin()
        queries = ['date_from=2026-10-19&date_to=2026-10-01', 'date_from=not-a-date', 'date_to=9999-12-31',
                   'event=private-invalid-event', 'outcome=private-invalid-outcome', 'batch_id=private-batch',
                   'cursor=private-invalid-cursor', 'page=0', 'page=word', 'page_size=101', 'page=1&cursor=abc']
        for query in queries:
            with self.subTest(query=query):
                response = self.client.get('/api/v1/admin/audit/events?' + query)
                self.assertEqual((response.status_code, response.json()['code']), (422, 'AUDIT_FILTER_INVALID'))
                self.assertNotIn('private-', response.text)
        response = self.client.get('/api/v1/admin/audit/export?date_from=not-a-date')
        self.assertEqual(response.json()['code'], 'AUDIT_FILTER_INVALID')

    def test_cross_month_range_and_page_beyond_end_are_valid_and_options_are_read_only(self):
        self.admin()
        with self.business.engine.connect() as c:
            count = c.execute(text('SELECT count(*) FROM audit')).scalar_one()
        for _ in range(2):
            self.assertEqual(self.client.get('/api/v1/admin/audit/options').status_code, 200)
        response = self.client.get('/api/v1/admin/audit/events?date_from=2025-01-01&date_to=2026-12-31&page=1000')
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['items'], [])
        with self.business.engine.connect() as c:
            self.assertEqual(c.execute(text('SELECT count(*) FROM audit')).scalar_one(), count)

    def test_export_exact_limit_and_over_limit_share_filter_and_snapshot_rules(self):
        self.admin()
        at = datetime.now(timezone.utc).isoformat()
        with self.business.store.transaction() as c:
            c.execute(text("INSERT INTO audit(id,actor,event,created_at,detail_json,version,outcome) VALUES (:id,'synthetic','upload',:at,'{}',1,'success')"),
                      [dict(id=str(uuid.uuid4()), at=at) for _ in range(10000)])
        query = '/api/v1/admin/audit/events?event=upload&page=1&page_size=20'
        self.assertEqual(self.client.get(query).json()['total'], 10000)
        response = self.client.get('/api/v1/admin/audit/export?event=upload')
        self.assertEqual(response.status_code, 200, response.text[:80] if response.status_code != 200 else '')
        self.assertEqual(load_workbook(io.BytesIO(response.content), read_only=True).active.max_row, 10001)
        with self.business.store.transaction() as c:
            c.execute(text("INSERT INTO audit(id,actor,event,created_at,detail_json,version,outcome) VALUES (:id,'synthetic','upload',:at,'{}',1,'success')"),
                      dict(id=str(uuid.uuid4()), at=at))
        self.assertEqual(self.client.get(query).json()['total'], 10001)
        response = self.client.get('/api/v1/admin/audit/export?event=upload')
        self.assertEqual((response.status_code, response.json()['code']), (422, 'AUDIT_EXPORT_TOO_LARGE'))
        self.assertNotIn('spreadsheetml', response.headers['Content-Type'])

    def test_non_audit_validation_and_authentication_errors_keep_their_codes(self):
        self.admin()
        response = self.client.get('/api/v1/batches?page=0')
        self.assertEqual(response.json()['code'], 'INVALID_REQUEST')
        self.client.cookies.clear()
        response = self.client.get('/api/v1/admin/audit/events?page=0')
        self.assertEqual((response.status_code, response.json()['code']), (401, 'UNAUTHORIZED'))


if __name__ == '__main__':
    unittest.main()
