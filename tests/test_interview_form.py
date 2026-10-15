"""Synthetic applicant forms, without real personnel fixtures."""
import io
from pathlib import Path
import sys
import unittest
from uuid import uuid4

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
from ocr_app.doc_profiles import profile
from ocr_app.document_classification import suggestion
from ocr_app.extraction import extract_result
from ocr_app.exporting import required_export_blockers, workbook_bytes


def block(text):
    return dict(text=text, page=1, region_id=str(uuid4()), bbox=[.1,.2,.9,.5],
                bbox_kind='region', extraction_method='ocr')


def fields(text):
    return {f['key']: f for f in extract_result([block(text)], 'INTERVIEW_FORM')['fields']}


def form(gender='☑男 □女', mobile='13800138000'):
    return ('<table><tr><td>身份证姓名</td><td>测试甲</td><td>性别</td><td>'+gender+
            '</td><td>年龄</td><td>43岁</td></tr><tr><td>身份证号码</td>'+
            ''.join('<td>'+c+'</td>' for c in '11010519491231002X')+
            '</tr><tr><td>手机号码</td><td>'+mobile+'</td><td>婚姻状况</td><td>未婚</td></tr>'+
            '<tr><td>紧急联络人</td><td>测试乙</td><td>电话</td><td>13900139000</td></tr></table>')


class InterviewFormTests(unittest.TestCase):
    def test_five_required_fields_and_classification(self):
        self.assertEqual([f['key'] for f in profile('INTERVIEW_FORM')],
                         ['name','gender','age','id_number','mobile_number'])
        self.assertTrue(all(f['required'] for f in profile('INTERVIEW_FORM')))
        self.assertEqual(suggestion('应聘登记表 身份证姓名 性别 年龄', '1.jpg')['doc_type'], 'INTERVIEW_FORM')

    def test_grid_digits_checkbox_and_applicant_mobile(self):
        f = fields(form())
        self.assertEqual({k:v['value'] for k,v in f.items()}, dict(name='测试甲', gender='M', age=43,
                          id_number='11010519491231002X', mobile_number='13800138000'))
        self.assertTrue(all(v.get('source') for v in f.values()))

    def test_no_derivation_from_id_or_emergency_phone(self):
        for gender in ['□男 □女', '☑男 ☑女', '男 女']:
            f = fields(form(gender=gender, mobile=''))
            self.assertIsNone(f['gender']['value'])
            self.assertIsNone(f['mobile_number']['value'])
        f=fields('身份证号码：11010519491231002X\n紧急联系人：测试乙 电话：13900139000')
        self.assertIsNone(f['age']['value'])
        self.assertIsNone(f['gender']['value'])
        self.assertIsNone(f['name']['value'])

    def test_missing_fields_block_export(self):
        result=extract_result([block(form(gender='☑女 □男', mobile=''))],'INTERVIEW_FORM')
        self.assertEqual(required_export_blockers(result),[dict(field='mobile_number',code='MISSING')])

    def test_relative_table_does_not_supply_applicant_fields(self):
        text=form(mobile='')+'<table><tr><td>姓名</td><td>亲属乙</td><td>年龄</td><td>60</td></tr></table>'
        f=fields(text)
        self.assertEqual(f['name']['value'],'测试甲')
        self.assertEqual(f['age']['value'],43)
        self.assertIsNone(f['mobile_number']['value'])

    def test_external_contact_headings_do_not_supply_applicant_fields(self):
        for heading in ['紧急联系人', '紧急联络人', '家庭成员', '家庭情况', '家属', '亲属信息', '工作经验']:
            for mobile in ['', '13800138000']:
                with self.subTest(heading=heading, mobile=mobile):
                    text = form(gender='☑女 □男', mobile=mobile) + '<h3>' + heading + '</h3>' + (
                        '<table><tr><td>姓名</td><td>测试乙</td><td>手机号码</td>'
                        '<td>13900139000</td><td>身份证号</td><td>11010519491231002X</td></tr></table>')
                    result = extract_result([block(text)], 'INTERVIEW_FORM')
                    f = {entry['key']: entry for entry in result['fields']}
                    self.assertEqual(f['name']['value'], '测试甲')
                    self.assertEqual(f['mobile_number']['value'], mobile or None)
                    self.assertEqual(required_export_blockers(result), [] if mobile else
                                     [dict(field='mobile_number', code='MISSING')])
                    self.assertIn('13900139000', result['generic']['full_text'])

    def test_generic_phone_table_needs_applicant_ownership(self):
        phone = '<table><tr><td>手机号码</td><td>13900139000</td></tr></table>'
        self.assertIsNone(fields(form(mobile='') + phone)['mobile_number']['value'])
        self.assertIsNone(fields('<h3>紧急联系人</h3>' + phone)['mobile_number']['value'])
        self.assertEqual(fields('<h3>基本信息</h3>' + phone)['mobile_number']['value'], '13900139000')
        self.assertEqual(fields(phone)['mobile_number']['value'], '13900139000')
        self.assertEqual(fields('<h3>紧急联系人</h3>' + phone + '<h3>基本信息</h3>' + form())
                         ['mobile_number']['value'], '13800138000')
        self.assertIsNone(fields('<h3>基本信息</h3>' + form(mobile='') + phone)['mobile_number']['value'])
        self.assertIsNone(fields('身份证姓名：测试甲\n紧 急 联 系 人\n手机号码：13900139000')
                          ['mobile_number']['value'])

    def test_planner_uses_basic_labels_not_largest_table_and_propagates_cancel(self):
        from unittest.mock import patch
        from PIL import Image
        from ocr_app.interview_form import plan_interview_regions
        from ocr_app.runtime import RuntimeFailure
        tables=[dict(box=(0,500,1000,900)),dict(box=(100,200,900,480))]
        lines=[dict(text=label,box=(110,220+i*30,200,240+i*30))
               for i,label in enumerate(['身份证姓名','性别','年龄','手机号码'])]
        lines.extend([dict(text='基本信息',box=(110,170,200,190)),
                      dict(text='最高学历',box=(110,400,200,420))])
        with patch('ocr_app.table_recognition.locate_table_candidates',return_value=tables), \
             patch('ocr_app.table_recognition.nearby_text',return_value=lines):
            plans,_=plan_interview_regions(Image.new('RGB',(1000,1000)),(0,0,1000,1000))
        self.assertEqual(plans,[dict(box=(84,154,916,396),task='table')])
        with patch('ocr_app.table_recognition.locate_table_candidates',side_effect=RuntimeFailure('CANCELLED')):
            with self.assertRaises(RuntimeFailure):
                plan_interview_regions(Image.new('RGB',(100,100)),(0,0,100,100))

    def test_multiline_values_and_misplaced_candidates_are_not_silently_assigned(self):
        f=fields('基本信息 测试甲\n身份证姓名 11010519491231002X 性别 ☑女 □男 年龄 43\n'
                 '身份证号码\n手机号码\n13800138000\n婚姻状况 未婚')
        self.assertIsNone(f['name']['value'])
        self.assertEqual(f['name']['candidates'][0]['value'],'测试甲')
        self.assertIsNone(f['id_number']['value'])
        self.assertEqual(f['id_number']['candidates'][0]['value'],'11010519491231002X')
        self.assertEqual(f['mobile_number']['value'],'13800138000')
        self.assertEqual(f['gender']['value'],'F')

    def test_truncated_output_is_retained_but_not_used_as_final_fields(self):
        evidence=block(form()); evidence['_interview_complete']=False
        result=extract_result([evidence],'INTERVIEW_FORM')
        self.assertTrue(all(f['value'] is None for f in result['fields']))
        self.assertIn('13800138000',result['generic']['full_text'])

    def test_conflicts_preserve_sources_and_saved_five_field_export(self):
        from openpyxl import load_workbook
        result=extract_result([block(form('☑女 □男')),block(form('☑女 □男','13900139000'))],'INTERVIEW_FORM')
        mobile=next(f for f in result['fields'] if f['key']=='mobile_number')
        self.assertIsNone(mobile['value'])
        self.assertEqual(len(mobile['candidates']),2)
        self.assertTrue(all(c.get('source') for c in mobile['candidates']))
        # Export serializes saved fields, never a model candidate or source text.
        mobile['value']='13800138000'
        book=load_workbook(io.BytesIO(workbook_bytes([result])))
        sheet=book['面试登记表']
        headers=[c.value for c in sheet[1]]
        start=headers.index('姓名')
        self.assertEqual(headers[start:start+5],['姓名','性别','年龄','身份证号','手机号'])
        for label in ['身份证号','手机号']:
            cell=sheet.cell(2,headers.index(label)+1)
            self.assertEqual(cell.data_type,'s')
            self.assertEqual(cell.font.sz,12)
        self.assertEqual(sheet.cell(1,start+1).fill.fgColor.rgb[-6:].lower(),'ef760e')

    def test_bounded_alternate_keeps_primary_fields_and_conflicts(self):
        import time
        from unittest.mock import patch
        from PIL import Image
        from ocr_app.pipeline import process_document
        from ocr_app.documents import png_bytes
        calls=[];attempts=[]
        def recognize(model, image, task, timeout, **extra):
            calls.append((model,task))
            return dict(text=form() if len(calls)==1 else form('☑女 □男'),termination_reason='eos')
        # The synthetic ID encodes female; primary male causes a scoped fallback.
        with patch('ocr_app.interview_form.plan_interview_regions',return_value=([dict(box=(0,0,100,100),task='table')],[])):
            result=process_document(png_bytes(Image.new('RGB',(100,100),'white')),'INTERVIEW_FORM',
                recognize=recognize,progress=lambda *a:None,save_page=lambda *a:str(uuid4()),
                record_attempt=attempts.append,cancelled=lambda:False,deadline=time.monotonic()+10)
        self.assertEqual(calls,[('paddleocr_vl','table'),('glm_ocr','ocr')])
        self.assertEqual(len(attempts),2)
        self.assertIsNone(next(f for f in result['fields'] if f['key']=='gender')['value'])
        self.assertEqual(len(result['generic']['blocks']),2)


class InterviewFormHTTPTests(unittest.TestCase):
    from test_business_api import BusinessAPITests as _fixtures
    setUp = _fixtures.setUp
    upload = _fixtures.upload

    def test_contact_phone_stays_missing_until_applicant_phone_is_saved(self):
        from openpyxl import load_workbook
        text = form(gender='☑女 □男', mobile='') + (
            '<h3>紧急联系人</h3><table><tr><td>手机号码</td><td>13900139000</td></tr></table>')
        created = self.upload('INTERVIEW_FORM').json()
        self.business.store.publish(self.business.store.claim_next(), extract_result([block(text)], 'INTERVIEW_FORM'))
        task = self.business.store.tasks(created['batch_id'])[0]
        current = self.business.store.result(task['task_id'])
        self.assertEqual(task['export_blockers'], [dict(field='mobile_number', code='MISSING')])
        exported = self.client.post('/api/v1/export', headers=self.headers, json={'batch_id': created['batch_id']})
        self.assertEqual(exported.status_code, 409, exported.text)
        response = self.client.post('/api/v1/tasks/' + task['task_id'] + '/review', headers=self.headers,
            json=dict(run_id=task['run_id'], expected_revision=task['result_revision'],
                      fields={'mobile_number': '13800138000'}, confirm_export_disputes=True))
        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue(response.json()['export_ready'])
        self.assertEqual(self.business.store.result(task['task_id'])['generic'], current['generic'])
        exported = self.client.post('/api/v1/export', headers=self.headers, json={'batch_id': created['batch_id']})
        self.assertEqual(exported.status_code, 200, exported.text)
        sheet = load_workbook(io.BytesIO(exported.content)).active
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(sheet.cell(2, headers.index('手机号') + 1).value, '13800138000')
