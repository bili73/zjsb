"""Pending package labels and selected-snapshot workbooks, no model calls."""
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from sqlalchemy import text
import test_business_api as base
import test_batch_ownership as ownership
from ocr_app.extraction import extract_result
from ocr_app.person_packages import name_suggestions


class PersonPackageNameTests(unittest.TestCase):
    setUp = base.BusinessAPITests.setUp
    account = ownership.BatchOwnershipTests.account
    password = 'synthetic account password'

    def create(self, headers, kinds=('DIPLOMA', 'DEGREE_CERT')):
        from PIL import Image
        files = []
        for i, kind in enumerate(kinds):
            data = io.BytesIO()
            Image.new('RGB', (300, 200), (i * 50, 100, 120)).save(data, format='PNG')
            files.append(('files', (kind + '.png', data.getvalue(), 'image/png')))
        response = self.client.post('/api/v1/batches', headers=headers,
            data={'file_types': json.dumps(kinds)}, files=files)
        self.assertEqual(response.status_code, 202, response.text)
        return response.json()

    def saved(self, task_id, kind, name):
        result = extract_result([], kind)
        for field in result['fields']:
            field.update(value=name if field['key'] == 'name' else '2020-06-01' if field['key'].endswith('_date') else '合成字段', need_review=False, review_issues=[])
        with self.business.store.transaction() as c:
            c.execute(text("UPDATE tasks SET status='NEED_REVIEW',result_enc=:value WHERE id=:id"),
                      dict(id=task_id, value=self.business.storage.pack(result)))
        return result

    def confirm(self, batch, headers, name='合成姓名', revision=0):
        return self.client.patch('/api/v1/batches/' + batch + '/person', headers=headers,
                                 json=dict(person_name=name, expected_revision=revision))

    def test_pending_package_keeps_pairing_and_can_append_before_naming(self):
        _, headers, _ = self.account('alice')
        created = self.create(headers, ('ID_CARD_FRONT',))
        batch = self.business.store.batch(created['batch_id'])
        self.assertTrue(batch['is_person_package'])
        self.assertIsNone(batch['person_name'])
        self.assertIn('姓名待确认', batch['display_name'])
        self.assertEqual(batch['id_card_pairing'], 'missing_back')
        response = self.client.post('/api/v1/batches/' + created['batch_id'] + '/append', headers=headers,
            data=dict(doc_type_hint='ID_CARD_BACK', expected_revision='0'),
            files={'files': ('back.png', self.image, 'image/png')})
        self.assertEqual(response.status_code, 202, response.text)
        self.assertEqual(self.business.store.batch(created['batch_id'])['id_card_pairing'], 'both_present')
        self.assertEqual(self.confirm(created['batch_id'], headers).status_code, 409)

    def test_interview_package_suggests_name_and_exports_all_five_fields(self):
        from test_interview_form import form, block
        _, headers, _ = self.account('alice')
        created = self.create(headers, ('INTERVIEW_FORM',))
        result = extract_result([block(form('☑女 □男'))], 'INTERVIEW_FORM')
        with self.business.store.transaction() as c:
            c.execute(text("UPDATE tasks SET status='NEED_REVIEW',result_enc=:value WHERE id=:id"),
                      dict(id=created['task_ids'][0], value=self.business.storage.pack(result)))
        batch = self.business.store.batch(created['batch_id'])
        self.assertEqual(batch['name_suggestions'][0]['name'], '测试甲')
        self.assertEqual(self.confirm(created['batch_id'], headers, name='测试甲').status_code, 200)
        response = self.client.post('/api/v1/export', headers=headers, json=dict(batch_id=created['batch_id']))
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.sheetnames[0], '人员信息总表')
        values = list(book.worksheets[0].values)
        row = dict(zip(values[0], values[1]))
        self.assertEqual([row['面试登记表·'+label] for label in ['姓名','性别','年龄','身份证号','手机号']],
                         ['测试甲','女',43,'11010519491231002X','13800138000'])

    def test_suggestions_sources_conflict_and_manual_label_are_independent(self):
        _, headers, _ = self.account('alice')
        created = self.create(headers)
        self.saved(created['task_ids'][0], 'DIPLOMA', '甲姓名')
        self.saved(created['task_ids'][1], 'DEGREE_CERT', '乙姓名')
        batch = self.business.store.batch(created['batch_id'])
        self.assertEqual([s['name'] for s in batch['name_suggestions']], ['甲姓名', '乙姓名'])
        self.assertEqual(batch['name_suggestions'][0]['sources'][0]['task_id'], created['task_ids'][0])
        self.assertIn('PERSON_NAME_CONFLICT', [w['code'] for w in batch['warnings']])
        tasks_before = self.business.store.tasks(created['batch_id'])
        response = self.confirm(created['batch_id'], headers, '人工确认姓名')
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['package_revision'], 1)
        self.assertEqual(self.business.store.tasks(created['batch_id']), tasks_before)
        self.saved(created['task_ids'][0], 'DIPLOMA', '新修正姓名')
        batch = self.business.store.batch(created['batch_id'])
        self.assertEqual(batch['person_name'], '人工确认姓名')
        self.assertEqual(batch['name_suggestions'][0]['name'], '新修正姓名')
        with self.business.engine.connect() as c:
            blob = c.execute(text('SELECT person_name_enc FROM batches WHERE id=:id'), {'id': created['batch_id']}).scalar_one()
        self.assertNotIn('人工确认姓名'.encode(), blob)
        self.assertEqual(self.confirm(created['batch_id'], headers, '过期覆盖', 0).status_code, 409)
        self.assertEqual(self.confirm(created['batch_id'], headers, '更正姓名', 1).json()['person_name'], '更正姓名')

    def test_name_confirmation_permission_validation_and_revoked_session(self):
        alice, headers, token = self.account('alice')
        _, other, _ = self.account('bob')
        _, viewer, _ = self.account('viewer', 'viewer')
        _, admin, _ = self.account('admin', 'admin')
        batch = self.create(headers)['batch_id']
        self.assertEqual(self.confirm(batch, other).status_code, 404)
        self.assertEqual(self.confirm(batch, viewer).status_code, 403)
        self.assertEqual(self.confirm(batch, {}).status_code, 401)
        for bad in ['', '  ', 'a\nb', 'x' * 81]:
            self.assertEqual(self.confirm(batch, headers, bad).status_code, 422)
        self.assertEqual(self.confirm(batch, admin).status_code, 200)
        original = self.business.store.confirm_person
        def revoke(*args, **kwargs):
            self.business.accounts.logout(token)
            return original(*args, **kwargs)
        with patch.object(self.business.store, 'confirm_person', side_effect=revoke):
            self.assertEqual(self.confirm(batch, headers, revision=1).status_code, 401)
        self.assertEqual(self.business.store.batch(batch)['package_revision'], 1)
        legacy = self.business.store.create_batch([dict(data=self.image, name='legacy.png', mime='image/png')],
                                                   'DIPLOMA', alice['user_id'])['batch_id']
        self.assertEqual(self.confirm(legacy, admin).status_code, 422)

    def test_export_overview_uses_only_selected_saved_results_and_keeps_repeated_documents(self):
        _, headers, _ = self.account('alice')
        created = self.create(headers, ('DIPLOMA', 'DIPLOMA', 'DEGREE_CERT'))
        for task, kind, name in zip(created['task_ids'], ('DIPLOMA', 'DIPLOMA', 'DEGREE_CERT'), ('=甲', '=甲', '未选择秘密姓名')):
            self.saved(task, kind, name)
        response = self.client.post('/api/v1/export', headers=headers,
                                    json={'task_ids': created['task_ids'][:2], 'masked': False})
        self.assertEqual(response.status_code, 200, response.text[:300] if response.status_code != 200 else '')
        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.sheetnames, ['人员信息总表', '毕业证'])
        overview = book.worksheets[0]
        self.assertEqual(overview['A2'].value, '姓名待确认')
        overview_values = dict(zip((cell.value for cell in overview[1]), (cell.value for cell in overview[2])))
        self.assertEqual(overview_values['材料份数'], 2)
        from ocr_app.doc_profiles import profile
        for number in (1, 2):
            for field in profile('DIPLOMA'):
                self.assertIn(f'毕业证{number}·{field["label"]}', overview_values)
        self.assertEqual(overview_values['毕业证1·姓名'], '=甲')
        self.assertEqual(overview_values['毕业证2·学校名称'], '合成字段')
        self.assertFalse(any('学位证' in key for key in overview_values))
        self.assertNotIn('秘密', '\n'.join(str(cell.value) for sheet in book for row in sheet for cell in row))
        self.assertEqual(book['毕业证'].max_row, 3)
        self.assertEqual(book['毕业证']['A2'].value, '=甲')
        self.assertEqual(book['毕业证']['A2'].data_type, 's')
        self.assertNotIn('来源文件',str(list(overview.values)))
        self.assertNotIn('DIPLOMA.png',str(list(overview.values)))
        self.assertEqual(overview_values['姓名候选'],'=甲：毕业证')
        self.assertEqual(overview['A1'].fill.fgColor.rgb[-6:], 'EF760E')
        self.assertEqual(overview['A1'].alignment.vertical, 'center')
        self.confirm(created['batch_id'], headers, '人工姓名')
        whole = self.client.post('/api/v1/export', headers=headers, json={'batch_id': created['batch_id']})
        self.assertEqual(whole.status_code, 200)
        book = load_workbook(io.BytesIO(whole.content))
        self.assertEqual(book.worksheets[0]['A2'].value, '人工姓名')
        overview_values = dict(zip((cell.value for cell in book.worksheets[0][1]),
                                   (cell.value for cell in book.worksheets[0][2])))
        self.assertEqual(overview_values['材料份数'], 3)
        self.assertEqual(overview_values['学位证·姓名'], '未选择秘密姓名')
        self.saved(created['task_ids'][1], 'DIPLOMA', None)
        self.assertEqual(self.client.post('/api/v1/export', headers=headers,
                                         json={'batch_id': created['batch_id']}).status_code, 409)

    def test_medical_package_keeps_eighteen_column_ledger(self):
        _, headers, _ = self.account('alice')
        created = self.create(headers, ('MEDICAL_REPORT',))
        result = self.saved(created['task_ids'][0], 'MEDICAL_REPORT', '合成姓名')
        response = self.client.post('/api/v1/export', headers=headers, json={'batch_id': created['batch_id']})
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.content))
        from ocr_app.medical_ledger import HEADERS
        self.assertEqual(book.sheetnames, ['人员信息总表', '审核记录台账'])
        self.assertEqual([c.value for c in book.worksheets[1][1]], HEADERS)
        self.assertEqual(book.worksheets[1].max_column, 18)
        overview_values = dict(zip((cell.value for cell in book.worksheets[0][1]),
                                   (cell.value for cell in book.worksheets[0][2])))
        self.assertEqual(overview_values['体检报告·体检项目'], '合成字段')
        self.assertEqual(overview_values['体检报告·姓名'], '合成姓名')

    def test_concurrent_names_share_one_revision_and_never_overwrite_a_winner(self):
        from concurrent.futures import ThreadPoolExecutor
        from ocr_app.task_store import TaskError
        user, headers, token = self.account('alice')
        batch_id = self.create(headers)['batch_id']
        def write(name):
            try:
                self.business.store.confirm_person(batch_id, name, 0, user['user_id'],
                    user=user, authorize=lambda c: self.business.accounts._acting_user(c, user['user_id'], token))
                return 'ok', name
            except TaskError as error:
                return error.code, name
        with ThreadPoolExecutor(2) as pool:
            outcomes = list(pool.map(write, ['甲姓名', '乙姓名']))
        self.assertEqual(sorted(code for code, _ in outcomes), ['RESULT_VERSION_CONFLICT', 'ok'])
        winner = next(name for code, name in outcomes if code == 'ok')
        batch = self.business.store.batch(batch_id)
        self.assertEqual((batch['package_revision'], batch['person_name']), (1, winner))


class PersonNameHelperTests(unittest.TestCase):
    def test_normalization_grouping_exclusions_and_explicit_bank_holder(self):
        def result(kind, key, value, identifier):
            return dict(doc_type=kind, task_id=identifier, fields=[dict(key=key, value=value)], _source_filename='synthetic.png')
        candidates = name_suggestions([
            result('DEGREE_CERT', 'name', '  合成姓名  ', 'b'), result('ID_CARD_FRONT', 'name', '合成姓名', 'a'),
            result('BANK_CARD', 'holder_name', 'OTHER PERSON', 'c'), result('HOUSEHOLD_REGISTER', 'name', '户主', 'd'),
            result('BUSINESS_LICENSE', 'name', '法人', 'e'), result('GENERIC', 'name', '正文姓名', 'f'),
            result('ID_CARD_BACK', 'name', '无来源姓名', 'g'), result('DIPLOMA', 'name', None, 'h')])
        self.assertEqual([c['name'] for c in candidates], ['合成姓名', 'OTHER PERSON'])
        self.assertEqual([s['task_id'] for s in candidates[0]['sources']], ['a', 'b'])

    def test_additive_0007_backfills_only_explicitly_named_packages(self):
        from alembic import command
        from alembic.config import Config
        from sqlalchemy import create_engine
        from ocr_app.db import open_database
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            engine = create_engine('sqlite:///' + str(root / 'tasks.sqlite3'))
            config = Config()
            config.set_main_option('script_location', str(Path(__file__).resolve().parents[1] / 'backend/migrations'))
            with engine.begin() as c:
                config.attributes['connection'] = c
                command.upgrade(config, '0006')
                for identifier, name in [('named', b'encrypted label'), ('legacy', None)]:
                    c.execute(text('INSERT INTO batches(id,created_at,name_enc,person_name_enc) VALUES (:id,:at,:name,:person)'),
                              dict(id=identifier, at='2026-10-12', name=b'original title', person=name))
                before = c.execute(text('SELECT * FROM batches ORDER BY id')).all()
            engine.dispose()
            for _ in range(2):
                engine = open_database(root)
                with engine.connect() as c:
                    rows = c.execute(text('SELECT * FROM batches ORDER BY id')).all()
                    self.assertEqual([tuple(r[:-1]) for r in rows], [tuple(r) for r in before])
                    self.assertEqual([r[-1] for r in rows], [0, 1])
                engine.dispose()
