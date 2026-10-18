"""Audit acceptance: isolated accounts/materials, no production access or inference."""
import io
from contextlib import closing
import json
from pathlib import Path
import shutil
import sqlite3
import tempfile
import unittest
import uuid
from unittest.mock import patch

from alembic import command
from alembic.config import Config
from openpyxl import load_workbook
from sqlalchemy import text

import test_business_api as base
import test_review_api as review
from ocr_app import audit
from ocr_app.db import open_database
from ocr_app.storage import Storage


class AuditTests(unittest.TestCase):
    setUp = base.BusinessAPITests.setUp
    upload = base.BusinessAPITests.upload
    task = base.BusinessAPITests.task
    ready = review.ReviewAPITests.ready
    submit = review.ReviewAPITests.submit
    password = 'synthetic-audit-password'
    origin = {'Origin': 'http://127.0.0.1'}

    def account(self, name='admin', role='admin'):
        return self.business.accounts.set_user(name, self.password, '合成姓名', role, create_only=True)

    def login(self, name='admin'):
        response = self.client.post('/api/v1/auth/login', headers=self.origin,
                                    json={'username': name, 'password': self.password})
        self.assertEqual(response.status_code, 200, response.text)
        self.headers = dict(self.origin, **{'X-CSRF-Token': response.json()['csrf_token']})
        return self.client.cookies.get('ocr_session')

    def admin(self):
        user = self.account()
        self.login()
        return user

    def records(self, event=None):
        with self.business.engine.connect() as connection:
            return [dict(row) for row in connection.execute(text('SELECT * FROM audit ORDER BY created_at,id')).mappings()
                    if event is None or row['event'] == event]

    def service(self):
        return audit.AuditService(self.business.engine, self.business.storage)

    def test_admin_boundary_all_audit_endpoints_and_no_mutation_routes(self):
        paths = ['/options', '/events', '/events?page=2&page_size=20', '/events/' + str(uuid.uuid4()), '/export']
        for path in paths:
            self.assertEqual(self.client.get('/api/v1/admin/audit' + path).status_code, 401)
        for role in ('operator', 'viewer'):
            self.account(role, role); self.login(role)
            for path in paths:
                self.assertEqual(self.client.get('/api/v1/admin/audit' + path).status_code, 403)
        self.admin()
        self.assertEqual(self.client.get('/api/v1/admin/audit/options').status_code, 200)
        self.assertEqual(self.client.get('/api/v1/admin/audit/events').status_code, 200)
        row = self.records('login')[-1]
        self.assertEqual(self.client.get('/api/v1/admin/audit/events/' + row['id']).status_code, 200)
        self.assertEqual(self.client.delete('/api/v1/admin/audit/events/' + row['id'], headers=self.headers).status_code, 405)

    def test_account_names_are_snapshotted_encrypted_and_deleted_users_not_reassigned(self):
        admin = self.admin()
        user = self.account('operator', 'operator')
        old = self.records('login')[-1]
        response = self.client.patch('/api/v1/admin/users/' + admin['user_id'], headers=self.headers,
                                     json={'display_name': '改名后的合成姓名'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.service().detail(old['id'])['identity']['display_name'], '合成姓名')
        changed = self.records('account_name_changed')[-1]
        self.assertEqual(self.service().detail(changed['id'])['context']['name_change']['after'], '改名后的合成姓名')
        self.assertEqual(self.client.delete('/api/v1/admin/users/' + user['user_id'], headers=self.headers).status_code, 204)
        removed = self.records('account_deleted')[-1]
        self.assertEqual(self.service().detail(removed['id'])['context']['target']['username'], 'operator')
        replacement = self.account('operator', 'operator')
        self.assertNotEqual(user['user_id'], replacement['user_id'])
        self.assertEqual(removed['target_id'], user['user_id'])
        plaintext = '\n'.join(str(value) for row in self.records() for value in row.values() if isinstance(value, str))
        self.assertNotIn('合成姓名', plaintext)
        self.assertNotIn(self.password, plaintext)
        self.assertNotIn(self.client.cookies.get('ocr_session'), plaintext)

    def test_failed_request_keeps_verified_identity_after_concurrent_account_deletion(self):
        self.admin()
        user = self.account('deleted-in-flight', 'operator')
        self.login('deleted-in-flight')
        token = audit.CONTEXT.set(dict(request_id=str(uuid.uuid4())))
        try:
            audit.bind_actor(user)
            with self.business.store.transaction() as connection:
                connection.execute(text('DELETE FROM web_sessions WHERE user_id=:id'), {'id': user['user_id']})
                connection.execute(text('DELETE FROM web_users WHERE id=:id'), {'id': user['user_id']})
            with self.business.store.transaction() as connection:
                identifier = audit.append(connection, self.business.storage, 'task_rerun', user['user_id'],
                                          outcome='failure', code='UNAUTHORIZED')
            self.assertEqual(self.service().detail(identifier)['identity']['username'], 'deleted-in-flight')
        finally:
            audit.CONTEXT.reset(token)

    def test_login_failures_throttle_and_untrusted_forwarded_address(self):
        self.account()
        for number in range(6):
            response = self.client.post('/api/v1/auth/login', headers=dict(self.origin, **{'X-Forwarded-For': '203.0.113.123'}),
                                        json={'username': 'admin', 'password': 'wrong-password-private'})
            self.assertEqual(response.status_code, 401 if number < 5 else 429)
        rows = self.records('login')
        self.assertEqual(len(rows), 6)
        self.assertTrue(all(row['outcome'] == 'failure' and row['actor'] == 'anonymous' for row in rows))
        detail = self.service().detail(rows[-1]['id'])
        self.assertEqual(detail['details']['code'], 'LOGIN_THROTTLED')
        self.assertEqual(detail['context']['claimed_username'], 'admin')
        self.assertNotEqual(detail['context']['peer'], '203.0.113.123')
        self.assertNotIn('wrong-password-private', json.dumps(detail))

    def test_new_login_audits_replacement_but_not_token(self):
        user = self.admin()
        first = self.client.cookies.get('ocr_session')
        second = self.login()
        self.assertEqual(self.records('session_replaced'), [])
        row = self.records('login')[-1]
        self.assertEqual(row['actor'], user['user_id'])
        self.assertEqual(json.loads(row['detail_json'])['revoked_sessions'], 1)
        self.assertIsNone(self.business.accounts.session(first))
        self.assertIsNotNone(self.business.accounts.session(second))
        self.assertNotIn(second, json.dumps(self.service().detail(row['id'])))

    def test_audit_write_failure_rolls_back_account_but_failure_record_survives(self):
        self.admin()
        original = audit.append
        def injected(connection, storage, event, *args, **kwargs):
            if event == 'account_created' and kwargs.get('outcome', 'success') == 'success':
                raise OSError('synthetic audit storage failure')
            return original(connection, storage, event, *args, **kwargs)
        with patch('ocr_app.audit.append', side_effect=injected):
            response = self.client.post('/api/v1/admin/users', headers=self.headers,
                json={'username': 'never-created', 'password': self.password, 'display_name': '测试'})
        self.assertEqual(response.status_code, 500)
        self.assertFalse(any(row['username'] == 'never-created' for row in self.business.accounts.users()))
        failed = self.records('account_created')[-1]
        self.assertEqual(failed['outcome'], 'failure')
        self.assertEqual(failed['request_id'], response.headers['X-Request-ID'])

    def test_logout_revokes_even_when_audit_append_fails(self):
        self.admin()
        token = self.client.cookies.get('ocr_session')
        with patch('ocr_app.audit.append', side_effect=OSError('synthetic isolated audit failure')):
            response = self.client.post('/api/v1/auth/logout', headers=self.headers)
        self.assertEqual(response.status_code, 204)
        self.assertIsNone(self.business.accounts.session(token))

    def test_review_history_preserved_without_audit_event_or_diff(self):
        self.admin()
        task, _ = self.ready()
        before = self.records()
        response = self.submit(task, fields=review.CERT)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(self.records(), before)
        with self.business.engine.connect() as connection:
            reference = connection.execute(text("SELECT reference FROM artifacts WHERE kind='review_revision'")).scalar_one()
        history = json.loads(self.business.storage.read(reference))
        self.assertEqual(history['request']['fields']['certificate_no'], review.CERT['certificate_no'])
        self.assertEqual(self.records('human_review'), [])
        saved = response.json()
        response = self.client.request('DELETE', '/api/v1/batches/' + task['batch_id'], headers=self.headers,
            json={'versions': [dict(task_id=saved['task_id'], run_id=saved['run_id'], expected_revision=saved['result_revision'])]})
        self.assertEqual(response.status_code, 200, response.text)
        deleted = self.records('batch_deleted')
        self.assertEqual(len(deleted), 1)
        self.assertIn('毕业证', self.service().detail(deleted[0]['id'])['object_label'])

    def test_failed_export_never_records_generated_and_success_is_version_bound(self):
        self.admin()
        task, _ = self.ready()
        saved = self.submit(task, fields=review.CERT).json()
        body = {'task_ids': [task['task_id']]}
        with patch('ocr_app.exporting.workbook_bytes', side_effect=ValueError('synthetic workbook failure')):
            response = self.client.post('/api/v1/export', headers=self.headers, json=body)
        self.assertEqual(response.status_code, 500)
        self.assertEqual([row['outcome'] for row in self.records('export_full')], ['failure'])
        self.assertEqual(self.records('export_full')[-1]['batch_id'], task['batch_id'])
        self.assertFalse(self.records('export_generated'))
        response = self.client.post('/api/v1/export', headers=self.headers, json=body)
        self.assertEqual(response.status_code, 200, response.text[:100] if response.status_code != 200 else '')
        generated = self.records('export_generated')[-1]
        self.assertEqual(json.loads(generated['detail_json'])['task_count'], 1)
        with self.business.engine.connect() as connection:
            references = connection.execute(text("SELECT reference FROM artifacts WHERE kind='access_history'")).scalars().all()
        self.assertTrue(any(json.loads(self.business.storage.read(ref)).get('result_revision') == saved['result_revision'] for ref in references))
        self.assertEqual(generated['request_id'], response.headers['X-Request-ID'])

    def test_blocked_export_keeps_batch_selection_and_no_success(self):
        self.admin()
        task, _ = self.ready()
        response = self.client.post('/api/v1/export', headers=self.headers, json={'batch_id': task['batch_id']})
        self.assertEqual(response.status_code, 409)
        row = self.records('export_full')[-1]
        self.assertEqual((row['outcome'], row['batch_id']), ('failure', task['batch_id']))
        self.assertFalse(self.records('export_generated'))

    def test_responses_match_documented_contract(self):
        import jsonschema
        import yaml
        from test_pipeline_schema import json_schema
        document = json_schema(yaml.safe_load((Path(__file__).resolve().parents[1] / 'openapi.yaml').read_text(encoding='utf-8')))
        self.admin()
        identifier = self.records('login')[-1]['id']
        for schema, value in [('AuditOptions', self.service().options()), ('AuditDetail', self.service().detail(identifier)),
                              ('AuditEvent', self.service().query()['items'][0])]:
            jsonschema.Draft7Validator(dict(document, **{'$ref': '#/components/schemas/' + schema}),
                                       format_checker=jsonschema.FormatChecker()).validate(value)

    def test_permission_rechecked_after_query_detail_options_and_export(self):
        for method, path in [('query', '/events'), ('query', '/events?page=2&page_size=20'), ('options', '/options'), ('detail', None), ('workbook', '/export')]:
            with self.subTest(method=method):
                if not self.business.accounts.configured():
                    self.account()
                token = self.login()
                path = path or '/events/' + self.records('login')[-1]['id']
                original = getattr(audit.AuditService, method)
                def revoke(service, *args, **kwargs):
                    result = original(service, *args, **kwargs)
                    self.business.accounts.logout(token)
                    return result
                with patch.object(audit.AuditService, method, revoke):
                    response = self.client.get('/api/v1/admin/audit' + path)
                self.assertEqual(response.status_code, 401)

    def test_query_ties_filters_invalid_cursor_and_formula_safe_excel(self):
        user = self.admin()
        with self.business.store.transaction() as connection:
            identifiers = [audit.append(connection, self.business.storage, 'login', user['user_id']) for _ in range(4)]
            for identifier in identifiers:
                connection.execute(text("UPDATE audit SET created_at='2026-10-16T00:00:00+00:00' WHERE id=:id"), {'id': identifier})
        service = self.service()
        filters = dict(date_from='2026-10-16', date_to='2026-10-16', event='login', actor=user['user_id'])
        found, cursor = [], None
        while True:
            page = service.query(**filters, cursor=cursor, limit=1)
            found.extend(row['id'] for row in page['items']); cursor = page['next_cursor']
            if not cursor: break
        self.assertEqual(len(found), len(set(found)))
        self.assertEqual([identifier for identifier in found if identifier in identifiers], sorted(identifiers, reverse=True))
        import base64
        bad_cursors = [base64.urlsafe_b64encode(json.dumps(value).encode()).decode() for value in
                       ([42, {}], {'2026-10-16': str(uuid.uuid4())}, ['2026-10-16', str(uuid.uuid4())])]
        for query in ('cursor=invalid', 'date_to=9999-12-31', 'event=unknown', 'outcome=invalid',
                      *('cursor=' + value for value in bad_cursors)):
            self.assertEqual(self.client.get('/api/v1/admin/audit/events?' + query).status_code, 422)
        self.business.accounts.update_display_name(user['user_id'], user['user_id'], '=1+1', self.client.cookies.get('ocr_session'))
        self.login()
        response = self.client.get('/api/v1/admin/audit/export')
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.content))
        self.assertTrue(any(cell.value == '=1+1' and cell.data_type == 's' for row in book.active for cell in row))
        self.assertEqual(book.active['A1'].fill.fgColor.rgb, '00EF760E')
        self.assertEqual(book.active['A1'].font.sz, 12)
        self.assertEqual(book.active['A1'].alignment.vertical, 'center')

    def test_numbered_pages_total_filters_and_legacy_cursor(self):
        user = self.admin()
        with self.business.store.transaction() as c:
            ids = [audit.append(c, self.business.storage, 'upload', user['user_id']) for _ in range(5)]
            for identifier in ids:
                c.execute(text("UPDATE audit SET created_at='2026-10-16T00:00:00+00:00' WHERE id=:id"), {'id': identifier})
        filters = dict(date_from='2026-10-16', date_to='2026-10-16', event='upload', actor=user['user_id'])
        service = self.service()
        page = service.query(**filters, page=2, page_size=2)
        self.assertEqual((page['total'], page['page'], page['page_size']), (5, 2, 2))
        self.assertEqual([r['id'] for r in page['items']], sorted(ids, reverse=True)[2:4])
        empty = service.query(**filters, page=100, page_size=2)
        self.assertEqual((empty['items'], empty['total'], empty['next_cursor']), ([], 5, None))
        first = service.query(**filters, limit=2)
        second = service.query(**filters, cursor=first['next_cursor'], limit=2)
        self.assertEqual(second['total'], 5, 'total is not the remaining count after a cursor')
        self.assertEqual(second['items'], page['items'])
        response = self.client.get('/api/v1/admin/audit/events', params={**filters, 'page': 2, 'page_size': 2})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), page)
        for params in ({'page': 0}, {'page': 1, 'page_size': 101}, {'page': 2**63},
                       {'page': 1, 'cursor': ''}, {'page_size': 20, 'limit': 10}):
            self.assertEqual(self.client.get('/api/v1/admin/audit/events', params=params).status_code, 422, params)
        missing = service.query(**filters, outcome='failure', page=1, page_size=2)
        self.assertEqual((missing['items'], missing['total']), ([], 0))

    def test_count_and_page_share_sqlite_snapshot_during_insert(self):
        from sqlalchemy import event
        user = self.admin()
        inserted = []
        def after_count(connection, cursor, statement, parameters, context, executemany):
            if statement.startswith('SELECT COUNT(*) FROM audit') and not inserted:
                inserted.append(True)
                with self.business.store.transaction() as writer:
                    audit.append(writer, self.business.storage, 'login', user['user_id'])
        before = len(self.records('login'))
        event.listen(self.business.engine, 'after_cursor_execute', after_count)
        try:
            page = self.service().query(event='login', page=1, page_size=100)
        finally:
            event.remove(self.business.engine, 'after_cursor_execute', after_count)
        self.assertTrue(inserted)
        self.assertEqual((page['total'], len(page['items'])), (before, before))
        self.assertEqual(len(self.records('login')), before + 1)

    def test_sqlite_and_encrypted_artifact_backup_restore(self):
        self.admin()
        task, _ = self.ready(); self.submit(task, fields=review.CERT)
        identifier = self.records('upload')[-1]['id']
        expected = self.service().detail(identifier)
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory) / 'restored'
            destination.mkdir()
            with closing(sqlite3.connect(self.business.storage.root / 'tasks.sqlite3')) as source, closing(sqlite3.connect(destination / 'tasks.sqlite3')) as copy:
                source.backup(copy)
            shutil.copytree(self.business.storage.blobs, destination / 'blobs')
            engine = open_database(destination)
            try:
                restored = audit.AuditService(engine, Storage(destination, key=b'T' * 32))
                self.assertEqual(restored.detail(identifier), expected)
            finally:
                engine.dispose()

    def test_export_limit_refuses_truncation_and_records_failure(self):
        self.admin()
        with self.business.store.transaction() as connection:
            connection.execute(text('''INSERT INTO audit(id,actor,event,created_at,detail_json,version,outcome)
                VALUES (:id,'synthetic-actor','upload',:at,'{}',1,'success')'''),
                [dict(id=str(uuid.uuid4()), at=audit.datetime.now(audit.timezone.utc).isoformat()) for _ in range(10001)])
        response = self.client.get('/api/v1/admin/audit/export?event=upload')
        self.assertEqual(response.status_code, 422)
        self.assertIn('10000', response.json()['message'])
        self.assertEqual([row['outcome'] for row in self.records('audit_export')], ['failure'])

    def test_calendar_filter_uses_shanghai_midnight_boundaries(self):
        with self.business.store.transaction() as connection:
            for stamp in ('2026-10-15T15:59:59+00:00', '2026-10-15T16:00:00+00:00',
                          '2026-10-16T15:59:59+00:00', '2026-10-16T16:00:00+00:00'):
                identifier = audit.append(connection, self.business.storage, 'upload', 'synthetic-actor')
                connection.execute(text('UPDATE audit SET created_at=:at WHERE id=:id'), dict(at=stamp, id=identifier))
        rows = self.service().query(date_from='2026-10-16', date_to='2026-10-16')['items']
        self.assertEqual([row['created_at'] for row in rows], ['2026-10-16T15:59:59+00:00', '2026-10-15T16:00:00+00:00'])


    def test_excluded_events_do_not_write_even_on_failure(self):
        user = self.admin()
        before = self.records()
        with self.business.store.transaction() as connection:
            for event in set(audit.EVENTS) - set(audit.ACCOUNT_EVENTS):
                for outcome in ('success', 'failure'):
                    self.assertIsNone(audit.append(connection, self.business.storage, event, user['user_id'], outcome=outcome))
            for actor in ('worker', 'system', 'anonymous'):
                self.assertIsNone(audit.append(connection, self.business.storage, 'upload', actor))
            self.assertIsNone(audit.append(connection, self.business.storage, 'export_full', user['user_id'], outcome='requested'))
        self.assertEqual(self.records(), before)

    def test_single_upload_audit_failure_rolls_back_entire_material_submission(self):
        self.admin()
        original = audit.append
        before = {path.name for path in self.business.storage.blobs.iterdir()}
        def injected(connection, storage, event, *args, **kwargs):
            if event == 'upload' and kwargs.get('outcome', 'success') == 'success':
                raise OSError('synthetic audit failure')
            return original(connection, storage, event, *args, **kwargs)
        with patch('ocr_app.audit.append', side_effect=injected):
            response = self.upload()
        self.assertEqual(response.status_code, 500)
        self.assertEqual(self.business.store.batches()['total'], 0)
        self.assertEqual({path.name for path in self.business.storage.blobs.iterdir()}, before)
        self.assertEqual([row['outcome'] for row in self.records('upload')], ['failure'])

    def test_historical_noise_filtered_before_pagination_options_detail_and_export(self):
        user = self.admin()
        identifiers = []
        with self.business.store.transaction() as connection:
            for event, actor, outcome in [('task_claimed', 'worker', 'success'), ('human_review', 'old-hidden-user', 'success'),
                    ('preview_full', user['user_id'], 'legacy'), ('export_full', user['user_id'], 'requested')]:
                identifier = str(uuid.uuid4()); identifiers.append(identifier)
                connection.execute(text('''INSERT INTO audit(id,actor,event,created_at,detail_json,version,outcome)
                    VALUES (:id,:actor,:event,:at,'{}',1,:outcome)'''),
                    dict(id=identifier, actor=actor, event=event, outcome=outcome, at=audit.datetime.now(audit.timezone.utc).isoformat()))
        before = self.records()
        page = self.service().query(limit=1)
        self.assertEqual(page['items'][0]['event'], 'login')
        self.assertIsNone(page['next_cursor'])
        self.assertEqual(page['total'], 1)
        self.assertEqual(self.service().query(page=1, page_size=20)['total'], 1)
        self.assertEqual(self.service().query(event='human_review')['items'], [])
        self.assertNotIn('old-hidden-user', [row['id'] for row in self.service().options()['actors']])
        self.assertNotIn('task_claimed', [row['value'] for row in self.service().options()['events']])
        for identifier in identifiers:
            self.assertEqual(self.client.get('/api/v1/admin/audit/events/' + identifier).status_code, 404)
        workbook, count = self.service().workbook()
        self.assertEqual(count, 1)
        self.assertEqual(load_workbook(io.BytesIO(workbook)).active.max_column, 6)
        self.assertEqual(self.records(), before, 'filtering must not rewrite or delete old records')

    def test_multi_file_upload_append_and_delete_are_single_account_actions(self):
        from PIL import Image
        user = self.admin()
        def image(index):
            stream = io.BytesIO(); Image.new('RGB', (24, 24), (index, 1, 2)).save(stream, format='PNG')
            return ('files', (f'{index}.png', stream.getvalue(), 'image/png'))
        response = self.client.post('/api/v1/batches', headers=self.headers,
            data=dict(doc_type_hint='DIPLOMA', person_name='合成人员'), files=[image(i) for i in range(3)])
        self.assertEqual(response.status_code, 202, response.text)
        batch = response.json()['batch_id']
        self.assertEqual(len(self.records('upload')), 1)
        self.assertEqual(json.loads(self.records('upload')[0]['detail_json'])['task_count'], 3)
        response = self.client.post(f'/api/v1/batches/{batch}/append', headers=self.headers,
            data=dict(doc_type_hint='DEGREE_CERT', expected_revision='0'), files=[image(4), image(5)])
        self.assertEqual(response.status_code, 202, response.text)
        self.assertEqual(len(self.records('package_appended')), 1)
        self.assertEqual(json.loads(self.records('package_appended')[0]['detail_json'])['task_count'], 2)
        self.assertFalse(self.records('task_created'))
        self.assertTrue(all(row['actor'] == user['user_id'] for row in self.records('upload') + self.records('package_appended')))
        self.assertEqual(self.service().query(event='upload')['items'][0]['description'], '上传材料（3份材料）')

    def test_multi_batch_export_is_one_row_and_each_batch_filter_matches_it(self):
        self.admin()
        tasks = [self.ready()[0] for _ in range(2)]
        for task in tasks:
            self.assertEqual(self.submit(task, fields=review.CERT).status_code, 200)
        body = dict(task_ids=[task['task_id'] for task in tasks])
        with patch('ocr_app.exporting.workbook_bytes', side_effect=ValueError('synthetic')):
            self.assertEqual(self.client.post('/api/v1/export', headers=self.headers, json=body).status_code, 500)
        self.assertEqual(len(self.records('export_full')), 1)
        self.assertFalse(self.records('export_generated'))
        self.assertEqual(self.client.post('/api/v1/export', headers=self.headers, json=body).status_code, 200)
        self.assertEqual(len(self.records('export_generated')), 1)
        row = self.records('export_generated')[0]
        self.assertEqual(json.loads(row['detail_json'])['task_count'], 2)
        for task in tasks:
            page = self.service().query(batch_id=task['batch_id'], event='export_generated', outcome='success')
            self.assertEqual([item['id'] for item in page['items']], [row['id']])
            self.assertEqual(page['total'], 1)
        self.assertEqual(len(self.records('export_full')), 1)

    def test_event_filter_combines_legacy_and_current_uploads_and_exports(self):
        user = self.admin()
        with self.business.store.transaction() as connection:
            for event, outcome in [('upload', 'success'), ('task_created', 'legacy'),
                                   ('export_generated', 'success'), ('export_full', 'failure')]:
                audit.append(connection, self.business.storage, event, user['user_id'], outcome=outcome)
        options = self.service().options()['events']
        self.assertEqual(sum(option['label'] == '上传材料' for option in options), 1)
        self.assertEqual(sum(option['label'] == '导出表格' for option in options), 1)
        self.assertEqual({row['event'] for row in self.service().query(event='upload')['items']}, {'upload', 'task_created'})
        self.assertEqual({row['event'] for row in self.service().query(event='export_generated')['items']}, {'export_generated', 'export_full'})
        self.assertEqual(len(self.service().query(event='export_generated', outcome='failure')['items']), 1)
        self.assertEqual(self.service().query(event='upload', page=2, page_size=1)['total'], 2)
        self.assertEqual(self.service().query(event='export_generated', outcome='failure', page=1)['total'], 1)

    def test_administrator_settings_use_chinese_action_summaries(self):
        user = self.admin()
        with self.business.store.transaction() as connection:
            identifier = audit.append(connection, self.business.storage, 'retention_changed', user['user_id'], old_days=7, days=15)
        self.assertEqual(self.service().detail(identifier)['description'], '文件保留期限：7天 → 15天')
        from ocr_app.required_fields import save_configuration
        save_configuration(self.business.store, dict(doc_type='DIPLOMA', required_fields=['name'], expected_revision=0),
                           user['user_id'], authorize=lambda connection: None)
        detail = self.service().detail(self.records('required_fields_changed')[-1]['id'])
        self.assertIn('毕业证必填字段', detail['description'])
        self.assertIn(' → 姓名', detail['description'])
        self.assertNotIn('certificate_no', detail['description'])
        self.assertEqual(detail['changes'], [])


class AuditMigrationTests(unittest.TestCase):
    def test_existing_seven_columns_unchanged_no_invented_outcome_or_identity(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            engine = open_database(root, migrate=False)
            config = Config(); config.set_main_option('script_location', str(Path(__file__).resolve().parents[1] / 'backend/migrations'))
            identifier = str(uuid.uuid4())
            with engine.begin() as connection:
                connection.exec_driver_sql('BEGIN IMMEDIATE')
                config.attributes['connection'] = connection
                command.upgrade(config, '0010')
                connection.execute(text("INSERT INTO audit VALUES (:id,'old-actor',NULL,NULL,'export_full','2026-10-16T00:00:00+00:00','{}')"), {'id': identifier})
                original = connection.execute(text('SELECT * FROM audit')).one()
            engine.dispose()
            engine = open_database(root)
            try:
                with engine.connect() as connection:
                    row = connection.execute(text('SELECT * FROM audit')).one()
                    self.assertEqual(tuple(row[:7]), tuple(original))
                    self.assertEqual((row.version, row.outcome, row.actor_snapshot_enc), (0, 'legacy', None))
                    self.assertEqual(connection.exec_driver_sql('PRAGMA foreign_key_check').all(), [])
            finally:
                engine.dispose()
