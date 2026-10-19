"""Bank-only script conversion retains raw evidence and human spelling choices."""
from copy import deepcopy
import io
from pathlib import Path
import sys
import unittest
from uuid import uuid4

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
from ocr_app.extraction import extract_result
from ocr_app.review import apply_review
from ocr_app.task_store import public_result
from ocr_app.exporting import workbook_bytes
from test_bank_photocopy import block, fields


class BankSimplificationTests(unittest.TestCase):
    def test_both_entries_simplify_non_name_text_and_keep_original_sources(self):
        text = ('姓名：陳曉華\n卡號：4111111111111111\n開戶行：中國農業銀行\n'
                '開戶支行：陝西省西安市高新區營業部\n卡組織：銀聯\n卡種：貸記卡\n有效期：長期')
        expected = dict(holder_name='陳曉華', card_number='4111111111111111',
                        bank_name='中国农业银行', branch_name='陕西省西安市高新区营业部',
                        card_scheme='银联', card_type='贷记卡', valid_thru='长期')
        for kind, label in [('BANK_CARD', '银行卡'), ('BANK_CARD_COPY', '银行卡复印件')]:
            with self.subTest(kind=kind):
                evidence = [block(text)]
                before = deepcopy(evidence)
                result = extract_result(evidence, kind)
                actual = fields(result)
                self.assertEqual({key: actual[key]['value'] for key in expected}, expected)
                self.assertEqual(evidence, before)
                self.assertEqual(result['generic']['full_text'], text)
                self.assertIn('中國農業銀行', actual['bank_name']['source']['raw_text'])
                self.assertIn('陝西省西安市高新區營業部', actual['branch_name']['source']['raw_text'])
                self.assertIn('銀聯', actual['card_scheme']['source']['raw_text'])
                result.update(task_id=str(uuid4()), run_id=str(uuid4()), revision=1)
                sheet = load_workbook(io.BytesIO(workbook_bytes([result])))[label]
                row = dict(zip((c.value for c in sheet[1]), (c.value for c in sheet[2])))
                self.assertEqual(row['姓名'], '陳曉華')
                self.assertEqual(row['开户行'], expected['bank_name'])
                self.assertEqual(row['开户支行'], expected['branch_name'])
                self.assertEqual(row['卡号'], expected['card_number'])

    def test_script_equivalent_branches_merge_and_real_conflicts_remain(self):
        for kind in ('BANK_CARD', 'BANK_CARD_COPY'):
            for texts in [('中國農業銀行東華支行', '中国农业银行东华支行'),
                          ('中国农业银行东华支行', '中國農業銀行東華支行')]:
                with self.subTest(kind=kind, texts=texts):
                    branch = fields(extract_result([block(t) for t in texts], kind))['branch_name']
                    self.assertEqual(branch['value'], '中国农业银行东华支行')
                    self.assertFalse(branch.get('candidates'))
                    self.assertEqual([s['raw_text'] for s in branch['sources']], list(texts))
            with self.subTest(kind=kind, conflict=True):
                texts = ['中國農業銀行東華支行', '中国农业银行西华支行']
                branch = fields(extract_result([block(t) for t in texts], kind))['branch_name']
                self.assertIsNone(branch['value'])
                self.assertEqual({c['value'] for c in branch['candidates']},
                                 {'中国农业银行东华支行', '中国农业银行西华支行'})
                self.assertEqual({c['source']['raw_text'] for c in branch['candidates']}, set(texts))

    def test_html_labels_names_and_multiline_branch_are_source_preserving(self):
        table = ('<table><tr><td>戶名</td><td>王國銀</td></tr>'
                 '<tr><td>開戶銀行</td><td>中國農業銀行</td></tr>'
                 '<tr><td>開戶網點</td><td>測試市東華營業部</td></tr></table>')
        for kind in ('BANK_CARD', 'BANK_CARD_COPY'):
            with self.subTest(kind=kind):
                result = extract_result([block(table)], kind)
                values = fields(result)
                self.assertEqual(values['holder_name']['value'], '王國銀')
                self.assertEqual(values['bank_name']['value'], '中国农业银行')
                self.assertEqual(values['branch_name']['value'], '测试市东华营业部')
                self.assertIn('測試市東華營業部', values['branch_name']['source']['raw_text'])
                wrapped = fields(extract_result([block('中國農業銀行測試市中華\n西路社區營業部')], kind))
                self.assertEqual(wrapped['branch_name']['value'], '中国农业银行测试市中华西路社区营业部')
                self.assertEqual(wrapped['branch_name']['source']['raw_text'], '中國農業銀行測試市中華\n西路社區營業部')

    def test_manual_and_historical_fields_keep_their_original_script(self):
        original = dict(holder_name='王國銀', card_number='4111111111111111',
                        bank_name='中國農業銀行', branch_name='陝西省東華營業部',
                        card_scheme='銀聯', card_type='貸記卡', valid_thru='長期')
        for kind, label in [('BANK_CARD', '银行卡'), ('BANK_CARD_COPY', '银行卡复印件')]:
            with self.subTest(kind=kind):
                result = extract_result([], kind)
                for key, value in original.items():
                    fields(result)[key]['value'] = value
                before = deepcopy(result)
                projected = public_result(result, {})
                self.assertEqual({k: fields(projected)[k]['value'] for k in original}, original)
                self.assertEqual(result, before)
                manual = apply_review(extract_result([], kind), {'fields': original}, actor='synthetic')
                saved = apply_review(manual, {}, actor='synthetic')
                self.assertEqual({k: fields(saved)[k]['value'] for k in original}, original)
                self.assertTrue(all(fields(saved)[k]['value_origin'] == 'manual' for k in original))
                saved.update(task_id=str(uuid4()), run_id=str(uuid4()), revision=1)
                sheet = load_workbook(io.BytesIO(workbook_bytes([saved])))[label]
                row = dict(zip((c.value for c in sheet[1]), (c.value for c in sheet[2])))
                self.assertEqual(row['开户行'], original['bank_name'])
                self.assertEqual(row['开户支行'], original['branch_name'])
                self.assertEqual(row['姓名'], original['holder_name'])

    def test_other_document_types_do_not_simplify_names_or_chinese_values(self):
        values = fields(extract_result([block('姓名：王國銀\n住址：陝西省東華區')], 'ID_CARD_FRONT'))
        self.assertEqual(values['name']['value'], '王國銀')
        self.assertEqual(values['address']['value'], '陝西省東華區')

    def test_traditional_context_keeps_number_and_name_boundaries(self):
        prefix = '11010519900101001'
        checksum = '10X98765432'[sum(int(c)*w for c, w in zip(prefix, [7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2])) % 11]
        pan = prefix + checksum
        for kind in ('BANK_CARD', 'BANK_CARD_COPY'):
            with self.subTest(kind=kind):
                text = ('持卡人簽名 AUTHORIZED SIGNATURE\n身份證號 1234567890123456\n'
                        '電話 6543210987654321\n陳曉華 電話 13800000000\n卡號：' + pan)
                result = extract_result([block(text)], kind)
                self.assertEqual(fields(result)['card_number']['value'], pan)
                self.assertEqual(fields(result)['holder_name']['value'], '陳曉華')
                self.assertEqual(result['generic']['full_text'], text)


if __name__ == '__main__':
    unittest.main()
