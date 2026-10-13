"""Real HTTP, isolated encrypted storage, synthetic images, no model calls."""
import io
import json
from uuid import uuid4
import unittest
from sqlalchemy import text
import test_business_api as business_tests
from ocr_app.extraction import extract_result, _medical_item


CERT = dict(name='测试人员', school_name='测试大学', major='计算机科学', education_level='本科',
            graduate_date='2020-06', certificate_no='0012345678901234')


class ReviewAPITests(unittest.TestCase):
    setUp = business_tests.BusinessAPITests.setUp
    upload = business_tests.BusinessAPITests.upload
    task = business_tests.BusinessAPITests.task

    def ready(self, kind='DIPLOMA', failed=False, pending=False):
        response = self.upload(kind)
        self.assertEqual(response.status_code, 202, response.text)
        identifier = response.json()['task_ids'][0]
        store = self.business.store
        worker = store.claim_next()
        page_id = store.save_page(worker, 1, self.image, 16, 12)
        region_id = str(uuid4())
        pages = [dict(page_id=page_id, page_index=1, width=16, height=12, extraction_method='ocr',
                      regions=[dict(region_id=region_id, status='pending' if pending else 'recognized', bbox=[0, 0, 1, 1])])]
        result = extract_result([], kind)
        result.update(pages=pages, status='FAILED' if failed else 'NEED_REVIEW')
        if failed:
            result['error'] = dict(code='EXTRACT_TIMEOUT', message='synthetic failure')
        self.assertTrue(store.publish(worker, result))
        return self.task(identifier), region_id

    def submit(self, task, **changes):
        return self.client.post(f"/api/v1/tasks/{task['task_id']}/review", headers=self.headers,
                                json=dict(run_id=task['run_id'], expected_revision=task['result_revision'], **changes))

    def result(self, task):
        return self.client.get(f"/api/v1/tasks/{task['task_id']}/result", headers=self.headers).json()

    def test_manual_completion_preserves_original_run_and_stale_protection(self):
        task, _ = self.ready()
        with self.business.engine.connect() as c:
            original = dict(c.execute(text('SELECT * FROM runs WHERE id=:id'), {'id': task['run_id']}).mappings().one())
        response = self.submit(task, fields=CERT)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['status'], 'COMPLETED')
        self.assertEqual(self.submit(task, fields=CERT).status_code, 409)
        result = self.result(task)
        self.assertEqual(result['result_revision'], task['result_revision'] + 1)
        self.assertTrue(all(f['confidence'] is None for f in result['fields']))
        with self.business.engine.connect() as c:
            self.assertEqual(dict(c.execute(text('SELECT * FROM runs WHERE id=:id'), {'id': task['run_id']}).mappings().one()), original)
            row = c.execute(text("SELECT reference FROM artifacts WHERE kind='review_revision'")).scalar_one()
            history = json.loads(self.business.storage.read(row))
            self.assertEqual(history['request']['fields']['certificate_no'], CERT['certificate_no'])
            self.assertNotIn(CERT['certificate_no'], '\n'.join(c.execute(text('SELECT detail_json FROM audit')).scalars()))

    def test_empty_and_whitespace_required_stay_blocked(self):
        task, _ = self.ready()
        response = self.submit(task)
        self.assertEqual(response.json()['status'], 'NEED_REVIEW')
        response = self.submit(response.json(), fields=dict(CERT, name='   '))
        self.assertEqual(response.json()['status'], 'NEED_REVIEW')
        for invalid in ({'name': True}, {'certificate_no': 123}, {'principal': 'X'}, {'name': '****1234'}):
            self.assertEqual(self.submit(response.json(), fields=invalid).status_code, 422)

    def test_failed_manual_revision_keeps_failed_run(self):
        task, _ = self.ready(failed=True)
        self.assertEqual(self.submit(task, fields=CERT).status_code, 422)
        response = self.submit(task, fields=CERT, coverage=dict(pages=[1], reason='已逐页核对完整原件'))
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['status'], 'COMPLETED')
        result = self.result(task)
        self.assertEqual(result['recognition_status'], 'FAILED')
        self.assertEqual(result['recognition_error']['code'], 'EXTRACT_TIMEOUT')
        with self.business.engine.connect() as c:
            self.assertEqual(c.execute(text('SELECT status FROM runs WHERE id=:id'), {'id': task['run_id']}).scalar_one(), 'FAILED')

    def test_failed_processed_prefix_is_not_full_coverage(self):
        task, _ = self.ready(failed=True)
        with self.business.store.transaction() as c:
            row = self.business.store._row(c, task['task_id'])
            meta = self.business.storage.unpack(row['meta_enc'])
            meta['actual_source_page_count'] = 2
            c.execute(text('UPDATE tasks SET meta_enc=:meta WHERE id=:id'), dict(meta=self.business.storage.pack(meta), id=task['task_id']))
        self.assertEqual(self.submit(task, fields=CERT, coverage=dict(pages=[1], reason='检查')).status_code, 422)

    def test_pending_region_requires_bound_correction(self):
        task, region = self.ready(failed=True, pending=True)
        coverage = dict(pages=[1], reason='逐页核对', regions=[dict(region_id=region, action='TRANSCRIBED', reason='已补录')])
        self.assertEqual(self.submit(task, fields=CERT, coverage=coverage).status_code, 422)
        coverage['regions'][0]['field_keys'] = ['name', 'certificate_no']
        self.assertEqual(self.submit(task, fields=CERT, coverage=coverage).status_code, 200)

    def test_preview_is_authenticated_version_bound_and_full_by_default(self):
        task, _ = self.ready()
        url = f"/api/v1/tasks/{task['task_id']}/preview"
        query = dict(run_id=task['run_id'], result_revision=task['result_revision'])
        self.assertEqual(self.client.get(url, params=query).status_code, 401)
        for flag in ({}, {'masked': 'true'}, {'masked': 'false'}):
            preview = self.client.get(url, headers=self.headers, params=dict(query, **flag))
            self.assertEqual(preview.status_code, 200, preview.text[:100] if preview.status_code != 200 else '')
            self.assertTrue(preview.content.startswith(b'\x89PNG'))
            self.assertEqual(preview.headers['X-Preview-Stage'], 'corrected')
        viewer_key = 'preview-synthetic-viewer-key-000000000'
        self.business.accounts.accounts.append(dict(user_id=str(uuid4()), display_name='viewer', role='viewer', key=viewer_key))
        self.assertEqual(self.client.get(url, headers={'X-API-Key': viewer_key}, params=query).status_code, 403)
        body = dict(run_id=task['run_id'], expected_revision=task['result_revision'], page=1, reason='核对原件')
        response = self.client.post(url, headers=self.headers, json=body)
        self.assertEqual(response.status_code, 200, response.text[:100])
        self.assertEqual(response.headers['X-Preview-Stage'], 'corrected')
        self.assertEqual(response.headers['Cache-Control'], 'no-store')
        self.assertTrue(response.content.startswith(b'\x89PNG'))
        body['expected_revision'] += 1
        self.assertEqual(self.client.post(url, headers=self.headers, json=body).status_code, 409)

    def test_export_all_or_nothing_and_roles(self):
        first, _ = self.ready()
        completed = self.submit(first, fields=CERT).json()
        other, _ = self.ready()
        body = dict(task_ids=[first['task_id'], other['task_id']])
        blocked = self.client.post('/api/v1/export', headers=self.headers, json=body)
        self.assertEqual(blocked.status_code, 409, blocked.text)
        self.assertNotIn(CERT['name'], blocked.text)
        body['task_ids'] = [first['task_id']]
        response = self.client.post('/api/v1/export', headers=self.headers, json=body)
        self.assertEqual(response.status_code, 200, response.text[:100] if response.status_code != 200 else '')
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(response.content))
        values = [str(cell.value) for sheet in book for row in sheet for cell in row]
        self.assertIn(CERT['certificate_no'], values)
        key = 'synthetic-viewer-key-0000000000'
        self.business.accounts.accounts.append(dict(user_id=str(uuid4()), display_name='viewer', role='viewer', key=key))
        viewer = {'X-API-Key': key}
        self.assertEqual(self.client.post('/api/v1/export', headers=viewer, json=body).status_code, 403)
        for flag in (True, False):
            self.assertEqual(self.client.post('/api/v1/export', headers=viewer, json=dict(body, masked=flag)).status_code, 403)
            response = self.client.post('/api/v1/export', headers=self.headers, json=dict(body, masked=flag))
            self.assertEqual(response.status_code, 200)
            book = load_workbook(io.BytesIO(response.content))
            self.assertIn(CERT['certificate_no'], [str(cell.value) for sheet in book for row in sheet for cell in row])
        self.assertEqual(self.client.post(f"/api/v1/tasks/{first['task_id']}/review", headers=viewer,
                                         json=dict(run_id=completed['run_id'], expected_revision=completed['result_revision'])).status_code, 403)

    def test_reveal_returns_only_selected_values_and_requires_reason(self):
        task, _ = self.ready()
        task = self.submit(task, fields=CERT).json()
        body = dict(run_id=task['run_id'], expected_revision=task['result_revision'], fields=['certificate_no'], reason='核对编号')
        url = f"/api/v1/tasks/{task['task_id']}/reveal"
        response = self.client.post(url, headers=self.headers, json=body)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['fields'], {'certificate_no': CERT['certificate_no']})
        self.assertEqual(self.client.post(url, headers=self.headers, json=dict(body, fields=['cvv'])).status_code, 422)
        self.assertEqual(self.client.post(url, headers=self.headers, json=dict(body, reason='   ')).status_code, 422)

    def test_legacy_source_count_is_read_from_original_bytes(self):
        task, _ = self.ready(failed=True)
        with self.business.store.transaction() as c:
            row = self.business.store._row(c, task['task_id'])
            meta = self.business.storage.unpack(row['meta_enc'])
            meta.pop('actual_source_page_count', None)
            c.execute(text('UPDATE tasks SET meta_enc=:meta WHERE id=:id'), dict(meta=self.business.storage.pack(meta), id=task['task_id']))
        response = self.submit(task, fields=CERT, coverage=dict(pages=[1], reason='已核对原文件全部一页'))
        self.assertEqual(response.status_code, 200, response.text)

    def test_malformed_nested_values_are_422_not_server_errors(self):
        task, _ = self.ready('MEDICAL_REPORT')
        cases = [dict(items=[{'item_id': []}]), dict(removed_item_ids=[{}]),
                 dict(findings=[{'finding_id': []}]), dict(absence_reasons={'phone': []}),
                 dict(report_summary={'removed_finding_ids': [{}]}),
                 dict(accepted_issues=[{'issue_id': [], 'reason': 'check'}])]
        for case in cases:
            with self.subTest(key=next(iter(case))):
                response = self.submit(task, **case)
                self.assertEqual(response.status_code, 422, response.text)

    def test_full_export_reason_is_encrypted_and_version_checked(self):
        task, _ = self.ready()
        task = self.submit(task, fields=CERT).json()
        body = dict(task_ids=[task['task_id']], masked=False, reason='SYNTHETIC PRIVATE REASON',
                    versions=[dict(task_id=task['task_id'], run_id=task['run_id'], result_revision=task['result_revision'])])
        response = self.client.post('/api/v1/export', headers=self.headers, json=body)
        self.assertEqual(response.status_code, 200)
        with self.business.engine.connect() as c:
            self.assertNotIn(body['reason'], '\n'.join(c.execute(text('SELECT detail_json FROM audit')).scalars()))
            histories = [json.loads(self.business.storage.read(ref)) for ref in c.execute(text("SELECT reference FROM artifacts WHERE kind='access_history'")).scalars()]
            self.assertTrue(any(entry['reason'] == body['reason'] for entry in histories))
        body['versions'][0]['result_revision'] -= 1
        self.assertEqual(self.client.post('/api/v1/export', headers=self.headers, json=body).status_code, 409)

    def test_bank_store_accepts_single_photo_without_ocr_admission(self):
        from ocr_app.task_store import TaskError
        unsafe = dict(data=self.image, name='card.png', mime='image/png', actual_source_page_count=1)
        for patch in ({'mime': 'application/pdf'}, {'actual_source_page_count': 2}, {'actual_source_page_count': None}):
            with self.assertRaises(TaskError):
                self.business.store.create_batch([dict(unsafe, **patch)], 'BANK_CARD', 'synthetic')
        self.assertFalse(self.business.store.tasks())
        self.assertFalse(list(self.business.storage.blobs.iterdir()))
        safe = unsafe
        result = self.business.store.create_batch([safe], 'BANK_CARD', self.business.accounts.accounts[0]['user_id'])
        public = self.task(result['task_ids'][0])
        self.assertNotIn('bank_admission', public['extract_options'])

    def test_changed_identifier_and_original_phone_evidence_are_full_without_private_seeds(self):
        task, _ = self.ready('DIPLOMA')
        store = self.business.store
        with store.transaction() as c:
            row = store._row(c, task['task_id'])
            result = store.storage.unpack(row['result_enc'])
            identifier = next(f for f in result['fields'] if f['key'] == 'certificate_no')
            identifier.update(value='001234567890123456', value_masked='****3456', source={'page': 1, 'raw_text': '证书编号001234567890123456'}, review_issues=[])
            result['generic']['full_text'] = '手机13800000000'
            c.execute(text('UPDATE tasks SET result_enc=:value WHERE id=:id'), dict(value=store.storage.pack(result), id=task['task_id']))
        response = self.submit(task, fields={'certificate_no': '001234567890123457'})
        self.assertEqual(response.status_code, 200, response.text)
        public = self.result(task)
        serialized = json.dumps(public)
        self.assertIn('13800000000', serialized)
        self.assertIn('001234567890123456', serialized)
        self.assertIn('001234567890123457', serialized)
        self.assertNotIn('_sensitive_history', serialized)
        self.assertNotIn('value_masked', serialized)
        with store.engine.connect() as c:
            row = store._row(c, task['task_id'])
            raw = store.storage.unpack(row['result_enc'])
            self.assertIn('001234567890123456', next(f for f in raw['fields'] if f['key'] == 'certificate_no')['_sensitive_history'])
            reference = c.execute(text("SELECT reference FROM artifacts WHERE kind='review_revision'")).scalar_one()
            history = json.loads(store.storage.read(reference))
            self.assertEqual(next(f for f in history['before']['fields'] if f['key'] == 'certificate_no')['value'], '001234567890123456')
            self.assertEqual(next(f for f in history['after']['fields'] if f['key'] == 'certificate_no')['value'], '001234567890123457')

    def test_one_click_acceptance_without_reason_keeps_history_and_blockers(self):
        task, _ = self.ready()
        with self.business.store.transaction() as c:
            row = self.business.store._row(c, task['task_id'])
            result = extract_result([dict(text='姓名：测试人员', page=1)], 'DIPLOMA')
            c.execute(text('UPDATE tasks SET result_enc=:data WHERE id=:id'), dict(data=self.business.storage.pack(result), id=task['task_id']))
        acceptable = [dict(issue_id=p['issue_id']) for f in result['fields'] for p in f['review_issues'] if p['can_accept']]
        response = self.submit(task, accepted_issues=acceptable)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['status'], 'NEED_REVIEW')
        with self.business.engine.connect() as c:
            ref = c.execute(text("SELECT reference FROM artifacts WHERE kind='review_revision'")).scalar_one()
            history = json.loads(self.business.storage.read(ref))
            self.assertEqual(history['request']['accepted_issues'], acceptable)
            self.assertTrue(history['actor'])
            self.assertTrue(history['at'])
        missing = next(p for f in result['fields'] for p in f['review_issues'] if p['code'] == 'MISSING')
        self.assertEqual(self.submit(response.json(), accepted_issues=[{'issue_id': missing['issue_id']}]).status_code, 422)

    def test_failed_saved_page_before_progress_can_be_reconstructed(self):
        task, _ = self.ready(failed=True)
        with self.business.store.transaction() as c:
            c.execute(text('UPDATE tasks SET pages_enc=NULL,result_enc=NULL WHERE id=:id'), {'id': task['task_id']})
        response = self.submit(task, fields=CERT, coverage=dict(pages=[1], reason='逐页补录全部原件'))
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['page_count'], 1)
        self.assertEqual(response.json()['pages'][0]['extraction_method'], 'pending')
        self.assertEqual(response.json()['status'], 'COMPLETED')


if __name__ == '__main__':
    unittest.main()
