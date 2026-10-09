"""Broad examination names are independent from laboratory result rows."""
import copy
import io
import unittest
from sqlalchemy import text
from test_hr_medical_profile import block, table
import test_business_api as business_tests
from ocr_app.doc_profiles import profile
from ocr_app.extraction import extract_result, _field
from ocr_app.task_store import public_result

RETIRED = {'report_no', 'package_name', 'marital_status', 'phone', 'work_unit'}
HEADER = ['序号', '项目', '临床意义']


def projects(evidence):
    return {f['key']: f for f in extract_result(evidence, 'MEDICAL_REPORT')['fields']}.get('exam_projects')


class MedicalProjectTests(unittest.TestCase):
    def test_profile_and_legacy_label_boundaries(self):
        meta = {f['key']: f for f in profile('MEDICAL_REPORT')}
        self.assertFalse(RETIRED & meta.keys())
        self.assertIn('exam_projects', meta)
        self.assertFalse(meta['exam_projects']['required'])
        result = extract_result([block('姓名：测试甲 性别：男 体检编号：00042 检查日期：2026-10-01 报告日期：2026-10-02 套餐名称：入职 体检机构：示例机构')], 'MEDICAL_REPORT')
        fields = {f['key']: f for f in result['fields']}
        self.assertEqual(fields['name']['value'], '测试甲')
        self.assertEqual(fields['gender']['value'], 'M')
        self.assertEqual(fields['exam_date']['value'], '2026-10-01')
        self.assertFalse(RETIRED & fields.keys())
        self.assertEqual(set(fields), set(result['basic_info']))
        self.assertIn('00042', result['generic']['full_text'])

    def test_scope_sources_and_distinct_detection_rows(self):
        grid = block(table([HEADER, ['1', '眼科常规', '意义文本'], ['2', '内科检查', '另一意义']]), 12, .2,
                     rows=[{}, {'cell_bboxes': [None, [.2,.3,.4,.4], None]}, {'cell_bboxes': [None, [.2,.4,.4,.5], None]}])
        grid['bbox'] = [.1,.2,.9,.5]
        evidence = [block('八、本次体检项目意义 ........ 12', 3), block(table([HEADER, ['1','目录误项','说明']]), 3, .2),
                    block('九、本次体检项目意义', 12), grid,
                    block('十、其他章节', 12, .4), block(table([HEADER, ['1','其他误项','说明']]), 12, .5),
                    block(table([['项目','检测结果'], ['白细胞','5.0']]), 13)]
        result = extract_result(evidence, 'MEDICAL_REPORT')
        f = {f['key']: f for f in result['fields']}['exam_projects']
        self.assertEqual(f['value'], '眼科常规；内科检查')
        self.assertEqual([s['raw_text'] for s in f['sources']], ['眼科常规', '内科检查'])
        self.assertTrue(all(s['page'] == 12 and s['bbox_kind'] == 'cell' for s in f['sources']))
        self.assertEqual([i['name'] for i in result['items']], ['白细胞'])

    def test_same_block_adjacent_page_reprints_and_gap(self):
        first = block('## 八、本次体检项目意义\n' + table([HEADER, ['1','内 科 检 查','说明']]), 8)
        second = block(table([HEADER, ['1','内科检查','说明'], ['2','DR 胸部正位检查','说明']]), 9)
        third = block(table([HEADER, ['1','越界误项','说明']]), 11)
        self.assertEqual(projects([first, second, third])['value'], '内科检查；DR 胸部正位检查')
        self.assertEqual(len(projects([first, second])['sources']), 3)

    def test_missing_or_ambiguous_header_does_not_guess(self):
        for rows in [[['序号','项目','项目','临床意义'], ['1','眼科','内科','说明']], [['项目','结果'],['白细胞','5.0']]]:
            with self.subTest(rows=rows):
                f = projects([block('八、本次体检项目意义'), block(table(rows), y=.2)])
                self.assertIsNotNone(f)
                self.assertIsNone(f['value'])
        self.assertIsNone(projects([block(table([HEADER, ['1','眼科','说明']]))])['value'])

    def test_historical_projection_preserves_original_and_adds_editable_blank(self):
        original = extract_result([], 'MEDICAL_REPORT')
        for key in RETIRED:
            original['fields'].append(_field(key, key, [], key == 'institution'))
        original['basic_info'] = {f['key']: f for f in original['fields']}
        original['fields'] = [f for f in original['fields'] if f['key'] != 'exam_projects']
        original['basic_info'].pop('exam_projects', None)
        before = copy.deepcopy(original)
        projected = public_result(original)
        keys = {f['key'] for f in projected['fields']}
        self.assertFalse(keys & RETIRED)
        self.assertEqual(keys, set(projected['basic_info']))
        self.assertIsNone(projected['basic_info']['exam_projects']['value'])
        self.assertTrue(projected['basic_info']['exam_projects']['editable'])
        self.assertEqual(original, before)
        self.assertNotIn('fields', public_result({'doc_type':'MEDICAL_REPORT', 'status':'QUEUED'}))


class MedicalProjectHTTPTests(unittest.TestCase):
    setUp = business_tests.BusinessAPITests.setUp
    upload = business_tests.BusinessAPITests.upload

    def test_old_result_projection_manual_entry_and_excel_columns(self):
        created = self.upload('MEDICAL_REPORT').json()
        store = self.business.store
        worker = store.claim_next()
        old = extract_result([block('姓名：测试甲 性别：男 检查日期：2026-10-01')], 'MEDICAL_REPORT')
        for key in RETIRED:
            old['fields'].append(_field(key, key, [], key == 'institution'))
        old['fields'] = [f for f in old['fields'] if f['key'] != 'exam_projects']
        old['basic_info'] = {f['key']: f for f in old['fields']}
        store.publish(worker, old)
        url = '/api/v1/tasks/' + created['task_ids'][0]
        with self.business.engine.connect() as c:
            encrypted = c.execute(text('SELECT result_enc FROM tasks WHERE id=:id'), {'id':created['task_ids'][0]}).scalar_one()
        current = self.client.get(url+'/result', headers=self.headers).json()
        self.assertFalse(RETIRED & current['basic_info'].keys())
        with self.business.engine.connect() as c:
            self.assertEqual(c.execute(text('SELECT result_enc FROM tasks WHERE id=:id'), {'id':created['task_ids'][0]}).scalar_one(), encrypted)
        body = dict(run_id=current['run_id'], expected_revision=current['result_revision'], fields={'exam_projects':'眼科常规；内科检查'})
        response = self.client.post(url+'/review', headers=self.headers, json=body)
        self.assertEqual(response.status_code, 200, response.text)
        latest = response.json()
        reloaded = self.client.get(url+'/result', headers=self.headers).json()
        self.assertFalse(RETIRED & reloaded['basic_info'].keys())
        self.assertEqual(reloaded['basic_info']['exam_projects']['value'], '眼科常规；内科检查')
        self.assertEqual(len(reloaded['fields']), 17)
        for key in RETIRED:
            invalid = dict(body, expected_revision=latest['result_revision'], fields={key:'retired'})
            self.assertEqual(self.client.post(url+'/review', headers=self.headers, json=invalid).status_code, 422)
        export = self.client.post('/api/v1/export', headers=self.headers, json={'task_ids':created['task_ids']})
        self.assertEqual(export.status_code, 200, export.text[:80] if export.status_code != 200 else '')
        from openpyxl import load_workbook
        sheet = load_workbook(io.BytesIO(export.content))['审核记录台账']
        headers = [c.value for c in sheet[1]]
        self.assertFalse({'体检编号','套餐名称','联系电话','婚姻状况','工作单位'} & set(headers))
        self.assertTrue({'报告日期','体检机构'} <= set(headers))
        self.assertEqual(sheet.cell(2, headers.index('体检项目数量')+1).value, 2)
