"""Synthetic, source-separated bank supplements; no real personal data fixtures."""
import io
from pathlib import Path
import sys
import time
import unittest
from unittest.mock import Mock, patch
from uuid import uuid4
from PIL import Image, ImageDraw

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'backend'))
from ocr_app.bank_photocopy import plan_bank_regions, card_boxes
from ocr_app.extraction import extract_result
from ocr_app.doc_profiles import profile
from ocr_app.exporting import required_export_blockers, workbook_bytes
from ocr_app.review import apply_review
from ocr_app.task_store import public_result


def block(text, role='whole', bbox=None):
    return dict(text=text,_bank_role=role,page=1,region_id=str(uuid4()),bbox=bbox or [0,0,1,1],bbox_kind='region',extraction_method='ocr')


def fields(result):return {field['key']:field for field in result['fields']}


class BankPhotocopyTests(unittest.TestCase):
    def test_non_pan_numbers_and_signature_are_not_business_fields(self):
        text = ('持卡人姓名 AUTHORIZED SIGNATURE\n客服电话：10101234 95555\n'
                '社会保障号码 11010519491231002X\n身份证号 110105194912310021\n'
                '行号 1234567890123456\n4111111111111111')
        f=fields(extract_result([block(text)],'BANK_CARD'))
        self.assertEqual(f['card_number']['value'],'4111111111111111')
        self.assertIsNone(f['holder_name']['value'])

    def test_valid_unlabelled_id_does_not_conflict_with_pan(self):
        # Synthetic dated Chinese ID with a numeric checksum; no sample fixture.
        prefix='11010519900101001'
        digit='10X98765432'[sum(int(c)*w for c,w in zip(prefix,[7,9,10,5,8,4,2,1,6,3,7,9,10,5,8,4,2]))%11]
        f=fields(extract_result([block('李小明\n'+prefix+digit+'\n4111111111111111')],'BANK_CARD'))
        self.assertEqual(f['card_number']['value'],'4111111111111111')
        self.assertEqual(f['holder_name']['value'],'李小明')
        f=fields(extract_result([block('11010519491231002X\n4111111111111111')],'BANK_CARD'))
        self.assertEqual(f['card_number']['value'],'4111111111111111')
        # Explicit card context is not discarded merely for resembling an ID.
        f=fields(extract_result([block('银行卡号：'+prefix+digit)],'BANK_CARD'))
        self.assertEqual(f['card_number']['value'],prefix+digit)

    def test_bank_script_variants_keep_original_evidence_without_conflict(self):
        f=fields(extract_result([block('中國銀行','front'),block('中国银行测试市东城支行','supplement')],'BANK_CARD'))
        self.assertEqual(f['bank_name']['value'],'中國銀行')
        self.assertFalse(f['bank_name'].get('candidates'))
        self.assertEqual(len(f['bank_name']['sources']),2)
        self.assertEqual(f['bank_name']['sources'][1]['raw_text'],'中国银行测试市东城支行')

    def test_bank_brand_branch_and_multiline_are_distinct(self):
        f=fields(extract_result([block('ABC 中国农业银行 服务一卡通\n开户行：测试市东城支行\n4111111111111111')],'BANK_CARD'))
        self.assertEqual(f['bank_name']['value'],'中国农业银行')
        self.assertEqual(f['branch_name']['value'],'测试市东城支行')
        f=fields(extract_result([block('合成银行测试市中华\n西路社区支行\n4111111111111111')],'BANK_CARD'))
        self.assertEqual(f['branch_name']['value'],'合成银行测试市中华西路社区支行')
        self.assertEqual(f['bank_name']['value'],'合成银行')

    def test_credit_union_front_and_online_banking_back(self):
        cards=[[100,100,600,410],[100,800,600,1110]]
        lines=[dict(text='测试信合',box=[150,150,400,190]),dict(text='4111111111111111',box=[140,280,540,320]),
               dict(text='网上银行 http://example.test',box=[150,900,500,940]),
               dict(text='24小时客服电话 10101234 95555',box=[150,950,540,990]),
               dict(text='??? 123 broken',box=[150,480,500,540]),dict(text='测试支行',box=[150,580,500,630])]
        with patch('ocr_app.bank_photocopy.card_boxes',return_value=cards),patch('ocr_app.table_recognition.nearby_text',return_value=lines):
            plans,_=plan_bank_regions(Image.new('RGB',(800,1200)),(0,0,800,1200),lambda:None)
        self.assertEqual(plans[0]['bank_role'],'front')
        self.assertLess(plans[0]['box'][3],800)
        self.assertTrue(any(p['box'][1]<=480 and p['box'][3]>=540 for p in plans[1:]))
        self.assertTrue(all(p['box'][3]<800 for p in plans))

    def test_side_supplements_and_branch_continuation(self):
        cards=[[400,250,900,560]]
        lines=[dict(text='合成银行',box=[500,290,800,330]),
               dict(text='李小明',box=[100,230,230,270]),dict(text='合成银行测试市中华',box=[100,310,350,360]),
               dict(text='西路社区支行',box=[100,358,340,410])]
        with patch('ocr_app.bank_photocopy.card_boxes',return_value=cards),patch('ocr_app.table_recognition.nearby_text',return_value=lines):
            plans,_=plan_bank_regions(Image.new('RGB',(1000,1400)),(0,0,1000,1400),lambda:None)
        self.assertEqual(plans[0]['bank_role'],'front')
        self.assertTrue(any(p['box'][0]<=100 and p['box'][1]<=310 and p['box'][3]>=410 for p in plans[1:]))

    def test_four_fields_and_unchanged_default_requirement(self):
        self.assertEqual([m['key'] for m in profile('BANK_CARD')][:4],['holder_name','card_number','bank_name','branch_name'])
        self.assertEqual([m['key'] for m in profile('BANK_CARD') if m['required']],['card_number'])

    def test_handwriting_phone_boundary_and_bank_provenance(self):
        result=extract_result([block('合成银行\n4111 1111 1111 1111','front'),
            block('4111111111111111\n李小明 13800000000\n测试市东城支行','supplement')],'BANK_CARD')
        values=fields(result)
        for key,value in dict(holder_name='李小明',card_number='4111111111111111',bank_name='合成银行',branch_name='测试市东城支行').items():
            self.assertEqual(values[key]['value'],value)
            self.assertNotIn('derivation',values[key])
        self.assertEqual(len(values['card_number']['sources']),2)
        self.assertEqual(values['holder_name']['source']['raw_text'],'李小明 13800000000')

    def test_traditional_bank_print_is_preserved_as_source(self):
        result=extract_result([block('中國銀行','front')],'BANK_CARD')
        self.assertEqual(fields(result)['bank_name']['value'],'中國銀行')
        self.assertEqual(fields(result)['bank_name']['source']['raw_text'],'中國銀行')

    def test_conflicting_numbers_retain_both_sources_and_no_checksum_choice(self):
        result=extract_result([block('4111111111111111','front'),block('4111111111111112','supplement')],'BANK_CARD')
        number=fields(result)['card_number']
        self.assertIsNone(number['value'])
        self.assertEqual([c['reason'] for c in number['candidates']],['卡面识别','补充文字识别'])
        self.assertTrue(all(c['source']['region_id'] for c in number['candidates']))
        self.assertEqual(required_export_blockers(result),[dict(field='card_number',code='MISSING')])
        fixed=apply_review(result,{'fields':{'card_number':'4111111111111112'}},actor='synthetic')
        self.assertFalse(required_export_blockers(fixed))
        self.assertEqual(len(fields(fixed)['card_number']['candidates']),2)

    def test_phone_and_back_never_become_card_or_owner(self):
        result=extract_result([block('13900000000','supplement'),block('持卡人签名\n6217000000000000000\n姓名 王小乙','back')],'BANK_CARD')
        self.assertIsNone(fields(result)['card_number']['value'])
        self.assertIsNone(fields(result)['holder_name']['value'])

    def test_ordinary_card_no_inferred_branch(self):
        result=extract_result([block('4111111111111111\n服务热线 13800000000\n欢迎使用')],'BANK_CARD')
        self.assertIsNone(fields(result)['branch_name']['value'])
        self.assertIsNone(fields(result)['holder_name']['value'])
        self.assertFalse(required_export_blockers(result))

    def test_standalone_handwritten_name_and_ambiguous_names(self):
        result=extract_result([block('张小甲','supplement')],'BANK_CARD')
        self.assertEqual(fields(result)['holder_name']['value'],'张小甲')
        result=extract_result([block('张小甲\n李小明','supplement')],'BANK_CARD')
        self.assertIsNone(fields(result)['holder_name']['value'])
        self.assertEqual(len(fields(result)['holder_name']['candidates']),2)

    def test_old_results_projection_preserves_manual_values_and_new_policy(self):
        old=extract_result([block('4111111111111111')],'BANK_CARD')
        old['fields']=[f for f in old['fields'] if f['key']!='branch_name']
        old=apply_review(old,{'fields':{'holder_name':'王小乙'}},actor='synthetic')
        old['fields']=[f for f in old['fields'] if f['key']!='branch_name']
        new=public_result(old,{'BANK_CARD':['card_number','branch_name']})
        self.assertEqual(fields(new)['holder_name']['value'],'王小乙')
        self.assertEqual(fields(new)['holder_name']['value_origin'],'manual')
        self.assertTrue(fields(new)['branch_name']['required'])
        self.assertNotIn('branch_name',fields(old))
        self.assertTrue(required_export_blockers(new,{'BANK_CARD':['branch_name']}))

    def test_manual_four_fields_export_as_strings(self):
        from openpyxl import load_workbook
        result=apply_review(extract_result([],'BANK_CARD'),{'fields':dict(holder_name='李小明',card_number='4111111111111111',bank_name='合成银行',branch_name='东城支行')},actor='synthetic')
        result.update(task_id=str(uuid4()),run_id=str(uuid4()),revision=1)
        book=load_workbook(io.BytesIO(workbook_bytes([result])))
        sheet=book['银行卡'];headers=[c.value for c in sheet[1]];values=[c.value for c in sheet[2]]
        for label,value in [('姓名','李小明'),('卡号','4111111111111111'),('开户行','合成银行'),('开户支行','东城支行')]:
            self.assertEqual(values[headers.index(label)],value)
        self.assertEqual(sheet.cell(2,headers.index('卡号')+1).data_type,'s')
        book=load_workbook(io.BytesIO(workbook_bytes([result],[dict(person_name='李小明',results=[result])])))
        row=list(book['人员信息总表'].values)
        self.assertEqual(dict(zip(row[0],row[1]))['银行卡·开户支行'],'东城支行')
        self.assertEqual(dict(zip(row[0],row[1]))['银行卡·姓名'],'李小明')

    def test_layout_separates_back_and_neighbor_form(self):
        image=Image.new('RGB',(1000,1400),'white')
        cards=[[250,150,850,520],[250,950,850,1320]]
        lines=[dict(text='合成银行',box=[300,180,460,220]),dict(text='持卡人签名',box=[600,1100,780,1140]),
            dict(text='4111111111111111',box=[280,550,800,600]),dict(text='李小明 13800000000',box=[300,650,790,700]),
            dict(text='东城支行',box=[300,750,700,800]),dict(text='姓名 王小乙',box=[0,650,150,700])]
        with patch('ocr_app.bank_photocopy.card_boxes',return_value=cards),patch('ocr_app.table_recognition.nearby_text',return_value=lines):
            plans,notes=plan_bank_regions(image,(0,0,1000,1400),lambda:None)
        self.assertEqual([p['bank_role'] for p in plans],['front','supplement','supplement','supplement'])
        self.assertTrue(all(p['box'][0]>150 and p['box'][3]<950 for p in plans))
        self.assertTrue(notes)

    def test_two_fronts_do_not_assign_a_person_or_branch(self):
        cards=[[100,100,600,410],[100,800,600,1110]]
        lines=[dict(text='合成银行',box=[150,150,400,190]),dict(text='合成银行',box=[150,850,400,890]),
               dict(text='李小明13800000000',box=[150,500,500,550])]
        with patch('ocr_app.bank_photocopy.card_boxes',return_value=cards),patch('ocr_app.table_recognition.nearby_text',return_value=lines):
            plans,notes=plan_bank_regions(Image.new('RGB',(800,1200)),(0,0,800,1200),lambda:None)
        self.assertTrue(all(p['bank_role']=='ambiguous_front' for p in plans));self.assertTrue(notes)
        result=extract_result([block('姓名 李小明\n开户支行 东城支行','ambiguous_front')],'BANK_CARD')
        self.assertIsNone(fields(result)['holder_name']['value']);self.assertIsNone(fields(result)['branch_name']['value'])

    def test_detector_failure_falls_back_but_deadline_propagates(self):
        image=Image.new('RGB',(500,700))
        with patch('ocr_app.bank_photocopy.card_boxes',side_effect=ValueError):
            plans,notes=plan_bank_regions(image,(0,0,500,700),lambda:None)
            self.assertEqual(plans[0]['bank_role'],'whole');self.assertTrue(notes)
            with self.assertRaises(TimeoutError):plan_bank_regions(image,(0,0,500,700),Mock(side_effect=TimeoutError))

    def test_actual_geometry_is_not_sample_coordinates(self):
        image=Image.new('RGB',(800,1200),'white');draw=ImageDraw.Draw(image)
        draw.rectangle((210,120,690,420),outline='gray',width=3)
        draw.rectangle((200,750,680,1050),outline='gray',width=3)
        self.assertEqual(len(card_boxes(image)),2)

    def test_pipeline_calls_each_region_and_retains_roles(self):
        from ocr_app.pipeline import process_document
        stream=io.BytesIO();Image.new('RGB',(100,150),'white').save(stream,format='PNG')
        plans=[dict(box=(0,0,100,70),task='ocr',bank_role='front'),dict(box=(0,70,100,140),task='ocr',bank_role='supplement')]
        recognize=Mock(side_effect=[dict(text='合成银行\n4111111111111111',termination_reason='eos'),dict(text='李小明 13800000000\n东城支行',termination_reason='eos')])
        with patch('ocr_app.bank_photocopy.plan_bank_regions',return_value=(plans,[])):
            result=process_document(stream.getvalue(),'BANK_CARD',recognize=recognize,progress=lambda *a:None,save_page=lambda *a:'page',record_attempt=lambda *a:None,cancelled=lambda:False,deadline=time.monotonic()+10)
        self.assertEqual(recognize.call_count,2);self.assertEqual(fields(result)['holder_name']['value'],'李小明')
        self.assertEqual(fields(result)['branch_name']['source']['bbox'],[0,.466667,1,.933333])


if __name__=='__main__':unittest.main()
