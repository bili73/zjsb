"""Source-bound template fields, independent counts and corrected exports."""
import copy
import io
import unittest
from openpyxl import load_workbook
from test_hr_medical_profile import block, table
from ocr_app.doc_profiles import profile
from ocr_app.extraction import extract_result, confirmed_abnormal
from ocr_app.exporting import workbook_bytes, required_export_blockers
from ocr_app.review import apply_review
from ocr_app.task_store import public_result, TaskError


def fields(result):
    return {f['key']: f for f in result['fields']}


def fixture():
    return [block('一、个人信息'), block(table([['姓名','测试甲','性别','男'],['年龄','28','体检日期','2026-10-01']]),y=.2),
        block('体检机构：示例医院\n报告日期：2026-10-02',y=.3),
        block('身高体重血压',page=8),
        block(table([['项目名称','检测结果','单位','参考范围','风险等级'],['身高','174','cm','0-300','C'],
            ['体重','78','kg','0-300','C'],['体重指数','25.76','kg/m2','18-24','C'],
            ['血压(收缩压)','129','mmHg','89-139','A'],['血压(舒张压)','84','mmHg','59-89','A']]),page=8,y=.2),
        block('尿常规检验报告单',page=9),
        block(table([['项目名称','检测结果','风险等级','异常诊断'],['潜血','+3','E3','潜血阳性'],
            ['蛋白质','+1','E3','尿蛋白质阳性'],['小结','两项异常','E3','']]),page=9,y=.2),
        block('心电图诊断科室',page=10),block(table([['项目名称','检测结果','风险等级'],
            ['心电图','窦性心律','A'],['小结','未见异常','A']]),page=10,y=.2),
        block('八、本次体检项目意义',page=12),
        block(table([['序号','项目','临床意义'],['1','身高体重血压','说明'],['2','尿常规','说明'],['3','心电图','说明']]),page=12,y=.2)]


def normal_groups():
    evidence = []
    for page, (heading, name, result) in enumerate([
        ('血常规检验报告单', '白细胞', '5.0'),
        ('尿常规检验报告单', '尿蛋白', '阴性'),
        ('心电图诊断科室', '心电图', '窦性心律'),
        ('DR影像', '胸部正位', '胸部未见异常'),
        ('腹部彩超', '肝脏', '形态正常'),
    ], start=7):
        evidence += [block(heading, page=page), block(table([
            ['项目名称', '检测结果', '提示', '风险等级'],
            [name, result, '正常', 'A'],
        ]), page=page, y=.2)]
    return evidence


class MedicalLedgerTests(unittest.TestCase):
    def test_partial_normal_evidence_is_explicitly_qualified(self):
        result = extract_result(normal_groups(), 'MEDICAL_REPORT')
        self.assertEqual(fields(result)['blood_routine']['value'], '已识别明细未见异常')
        sheet = load_workbook(io.BytesIO(workbook_bytes([result])))['审核记录台账']
        self.assertEqual(sheet['K2'].value, '已识别明细未见异常')

    def test_excel_qualifies_detail_derived_normal_without_relabeling_other_conclusions(self):
        result = extract_result(normal_groups(), 'MEDICAL_REPORT')
        before = copy.deepcopy(result)
        sheet = load_workbook(io.BytesIO(workbook_bytes([result]))).active
        self.assertEqual(sheet['K2'].value, '已识别明细未见异常')
        self.assertIsNotNone(sheet['K2'].comment)
        self.assertIn('仅基于已识别明细', sheet['K2'].comment.text)
        self.assertEqual(result, before)

        manual = copy.deepcopy(result)
        fields(manual)['blood_routine'].update(value='未见异常', value_origin='manual')
        manual_sheet = load_workbook(io.BytesIO(workbook_bytes([manual]))).active
        self.assertEqual(manual_sheet['K2'].value, '未见异常')
        self.assertIsNone(manual_sheet['K2'].comment)
        source_conclusion = extract_result(fixture(), 'MEDICAL_REPORT')
        source_sheet = load_workbook(io.BytesIO(workbook_bytes([source_conclusion]))).active
        self.assertEqual(source_sheet['M2'].value, '未见异常')
        self.assertIsNone(source_sheet['M2'].comment)

    def test_a_grade_laboratory_rows_need_no_conclusion(self):
        evidence = []
        for page, heading, name, value in [(7, '血常规检验报告单', '白细胞', '5.0'),
                (8, '尿常规检验报告单', '尿蛋白', '阴性'),
                (9, '生化检验报告单', '总胆红素', '16.05')]:
            evidence += [block(heading, page), block(table([
                ['检验项目', '检测结果', '风险等级'], [name, value, 'A']]), page, .2)]
        r = extract_result(evidence, 'MEDICAL_REPORT')
        for key in ('blood_routine', 'urine_routine', 'liver_function'):
            self.assertEqual(fields(r)[key]['value'], '已识别明细未见异常', key)
        self.assertTrue(all(i['abnormal_flag'] == 'UNKNOWN' and i['need_review'] for i in r['items']))

    def test_closed_native_catalogue_distinguishes_absent_from_unrecognized(self):
        evidence = [block('八、本次体检项目意义', 18), block(table([
            ['序号', '项目', '临床意义'], ['1', '血常规', '说明'], ['2', '肝功五项', '说明'],
            ['3', '人绒毛膜促性腺激素（HCG尿检）', '说明']]), 18, .2),
            block('检后服务介绍', 19)]
        r = extract_result(evidence, 'MEDICAL_REPORT')
        for key in ('urine_routine', 'abdominal_ultrasound'):
            self.assertIsNone(fields(r)[key]['value'])
            self.assertEqual(fields(r)[key].get('absence_reason'), 'NOT_ON_DOCUMENT')
            self.assertEqual(fields(r)[key]['sources'], fields(r)['exam_projects']['sources'])
        for key in ('blood_routine', 'liver_function'):
            self.assertIsNone(fields(r)[key]['value'])
            self.assertNotIn('absence_reason', fields(r)[key])
        before = copy.deepcopy(r)
        self.assertEqual(public_result(public_result(r)), public_result(r))
        sheet = load_workbook(io.BytesIO(workbook_bytes([r]))).active
        self.assertEqual(sheet['L2'].value, '报告未载明')
        self.assertEqual(sheet['O2'].value, '报告未载明')
        self.assertEqual(r, before)
        for incomplete in (evidence[:-1], [evidence[0], block(evidence[1]['text'].replace('<td>1</td>', '<td>2</td>'), 18, .2), evidence[2]],
                [evidence[0], evidence[1], block('检后服务介绍', 21)],
                evidence + [block('尿常规检验报告单', 7)]):
            self.assertNotIn('absence_reason', fields(extract_result(incomplete, 'MEDICAL_REPORT'))['urine_routine'])
        ocr = copy.deepcopy(evidence)
        for entry in ocr:
            entry['extraction_method'] = 'ocr'
        self.assertNotIn('absence_reason', fields(extract_result(ocr, 'MEDICAL_REPORT'))['urine_routine'])
        manual = apply_review(r, {'fields': {'urine_routine': '人工核对'}}, actor='test')
        self.assertEqual(fields(public_result(manual))['urine_routine']['value'], '人工核对')
        changed = apply_review(r, {'fields': {'exam_projects': '血常规；尿常规'}}, actor='test')
        self.assertNotIn('absence_reason', fields(public_result(changed))['urine_routine'])

    def test_normal_groups_without_conclusions_fill_source_backed_summary(self):
        r = extract_result(normal_groups(), 'MEDICAL_REPORT')
        keys = ('blood_routine', 'urine_routine', 'ecg', 'dr_imaging', 'abdominal_ultrasound')
        for key in keys:
            with self.subTest(key=key):
                field = fields(r)[key]
                self.assertEqual(field['value'], '已识别明细未见异常')
                self.assertEqual(field['value_origin'], 'derived')
                self.assertEqual(field['sources'], [i['source'] for i in r['items'] if i['category'] == field['label']])
                self.assertIsNone(field['confidence'])
                self.assertTrue(all(i['need_review'] for i in r['items'] if i['category'] == field['label']))
                self.assertEqual(field['derivation']['conclusions'], [])
        self.assertIsNone(fields(r)['liver_function']['value'])
        sheet = load_workbook(io.BytesIO(workbook_bytes([r]))).active
        self.assertEqual([sheet.cell(2, n).value for n in range(11, 16)], ['已识别明细未见异常'] * 5)
        self.assertIsNone(sheet['P2'].value)

    def test_normal_fallback_does_not_hide_uncertain_group_rows(self):
        original = extract_result(normal_groups(), 'MEDICAL_REPORT')
        variants = [
            {'abnormal_flag': 'UNKNOWN', 'abnormal_basis': 'unknown', 'risk_grade': '/'},
            {'abnormal_flag': 'POSITIVE', 'abnormal_basis': 'unknown'},
            {'abnormal_basis': 'unknown'},
            {'review_issues': [{'code': 'CONFLICT', 'message': 'conflicting evidence'}]},
            {'source': None}, {'result_raw': ''}, {'result_raw': '/'},
        ]
        for patch in variants:
            with self.subTest(patch=patch):
                r = copy.deepcopy(original)
                row = copy.deepcopy(r['items'][0])
                row.update(patch)
                r['items'].append(row)
                self.assertIsNone(fields(public_result(r))['blood_routine']['value'])
        r = copy.deepcopy(original)
        r['items'] = []
        self.assertIsNone(fields(public_result(r))['blood_routine']['value'])
        r = copy.deepcopy(original)
        for row in r['items']:
            row['category'] = None
        self.assertIsNone(fields(public_result(r))['blood_routine']['value'])

    def test_normal_fallback_preserves_abnormal_conclusion_and_manual_priority(self):
        evidence = normal_groups()
        evidence[1] = block(table([
            ['项目名称', '检测结果', '参考范围', '风险等级'],
            ['白细胞', '5.0', '3.5-9.5', 'A'], ['小结', '正常', '', 'A'],
        ]), y=.2)
        r = extract_result(evidence, 'MEDICAL_REPORT')
        self.assertEqual(fields(r)['blood_routine']['value'], '未见异常')
        r = extract_result(normal_groups(), 'MEDICAL_REPORT')
        row = r['items'][0]
        row.update(abnormal_flag='HIGH', abnormal_basis='report_marker', result_raw='12.0')
        self.assertEqual(fields(public_result(r))['blood_routine']['value'], '白细胞 12.0')
        row.update(abnormal_flag='NORMAL', result_raw='5.0')
        manual = apply_review(r, {'fields': {'blood_routine': '人工已核对'}}, actor='test')
        self.assertEqual(fields(public_result(manual))['blood_routine']['value'], '人工已核对')
        absent = apply_review(r, {'absence_reasons': {'blood_routine': 'NOT_ON_DOCUMENT'}}, actor='test')
        self.assertIsNone(fields(public_result(absent))['blood_routine']['value'])

    def test_historical_normal_summary_refresh_is_read_only_and_matches_export(self):
        r = extract_result(normal_groups(), 'MEDICAL_REPORT')
        fields(r)['blood_routine']['value'] = None
        before = copy.deepcopy(r)
        projected = public_result(r)
        self.assertEqual(fields(projected)['blood_routine']['value'], '已识别明细未见异常')
        self.assertEqual(r, before)
        self.assertEqual(projected['stats'], before['stats'])
        sheet = load_workbook(io.BytesIO(workbook_bytes([projected]))).active
        self.assertEqual(sheet['K2'].value, '已识别明细未见异常')
        self.assertEqual(r, before)

    def test_normal_biochemical_liver_measurements_fill_normal_summary(self):
        rows=[['总胆红素','TBIL','16.05','μmol/L','A'],['直接胆红素','DBIL','3.58','μmol/L','A'],
            ['间接胆红素','IBIM','12.47','μmol/L','A'],['总蛋白','TP','72.35','g/L','A'],
            ['白蛋白','ALB','50.69','g/L','A'],['球蛋白','GLO','21.66','g/L','A'],
            ['丙氨酸氨基转移酶','ALT','24.54','U/L','A'],['葡萄糖','GLU','5.0','mmol/L','A']]
        evidence=[block('生化检验报告单',page=14),
            block(table([['检验项目','缩写','检测结果','单位','风险等级'],*rows]),page=14,y=.2),
            block('肝炎病毒检测检验报告单',page=14,y=.5),
            block(table([['检验项目','检测结果','风险等级'],['乙肝表面抗原','-','A']]),page=14,y=.6)]
        r=extract_result(evidence,'MEDICAL_REPORT');field=fields(r)['liver_function']
        self.assertEqual(field['value'], '已识别明细未见异常')
        self.assertEqual(len([i for i in r['items'] if i.get('category')=='肝功五项（生化检测）']),7)
        self.assertEqual(load_workbook(io.BytesIO(workbook_bytes([r]))).active['P2'].value, '已识别明细未见异常')
        self.assertEqual(sum(confirmed_abnormal(i) for i in r['items']),0)

    def test_liver_summary_preserves_manual_value_and_absence(self):
        r=extract_result([block('生化检验报告单'),block(table([
            ['检验项目','检测结果','单位','风险等级'],['总胆红素','16.05','μmol/L','A']]),y=.2)],'MEDICAL_REPORT')
        manual=apply_review(r,{'fields':{'liver_function':'人工核对结果'}},actor='test')
        self.assertEqual(fields(public_result(manual))['liver_function']['value'],'人工核对结果')
        absent=apply_review(r,{'absence_reasons':{'liver_function':'NOT_ON_DOCUMENT'}},actor='test')
        self.assertIsNone(fields(public_result(absent))['liver_function']['value'])

    def test_liver_ledger_prioritizes_abnormal_measurements_and_preserves_normal_conclusions(self):
        abnormal=extract_result([block('生化检验报告单'),block(table([
            ['检验项目','检测结果','单位','风险等级'],['总胆红素','16.05','μmol/L','A'],
            ['丙氨酸氨基转移酶','80','U/L','C']]),y=.2)],'MEDICAL_REPORT')
        value=fields(abnormal)['liver_function']['value']
        self.assertIn('丙氨酸氨基转移酶 80 U/L',value)
        self.assertNotIn('总胆红素',value)
        self.assertEqual(load_workbook(io.BytesIO(workbook_bytes([abnormal]))).active['P2'].value,value)
        conclusion=extract_result([block('肝功能检验报告单'),block(table([
            ['检验项目','检测结果','风险等级'],['总胆红素','16.05','A'],['小结','未见异常','A']]),y=.2)],'MEDICAL_REPORT')
        self.assertEqual(fields(conclusion)['liver_function']['value'], '未见异常')
        self.assertEqual(load_workbook(io.BytesIO(workbook_bytes([conclusion]))).active['P2'].value, '未见异常')

    def test_historical_derived_normal_liver_values_refresh_only_in_read_projection(self):
        r=extract_result([block('生化检验报告单'),block(table([
            ['检验项目','检测结果','风险等级'],['总胆红素','16.05','A']]),y=.2)],'MEDICAL_REPORT')
        fields(r)['liver_function']['value']='总胆红素 16.05'
        before=copy.deepcopy(r)
        self.assertEqual(fields(public_result(r))['liver_function']['value'], '已识别明细未见异常')
        self.assertEqual(r,before)

    def test_liver_unknown_or_sourceless_rows_do_not_automatically_fill_field(self):
        for grade in ['/', '']:
            r=extract_result([block('生化检验报告单'),block(table([
                ['检验项目','检测结果','风险等级'],['总胆红素','16.05',grade],
                ['葡萄糖','9.0','C']]),y=.2)],'MEDICAL_REPORT')
            self.assertIsNone(fields(r)['liver_function']['value'])
        r=extract_result([block('生化检验报告单'),block(table([
            ['检验项目','检测结果','风险等级'],['总胆红素','30','C']]),y=.2)],'MEDICAL_REPORT')
        for item in r['items']: item.pop('source',None)
        self.assertIsNone(fields(public_result(r))['liver_function']['value'])

    def test_biochemical_heading_without_matching_liver_rows_does_not_fill_summary(self):
        r=extract_result([block('生化检验报告单'),block(table([
            ['检验项目','检测结果','单位','风险等级'],['葡萄糖','5.0','mmol/L','A']]),y=.2)],'MEDICAL_REPORT')
        self.assertIsNone(fields(r)['liver_function']['value'])

    def test_profile_source_values_and_separate_project_count(self):
        r=extract_result(fixture(),'MEDICAL_REPORT')
        f=fields(r)
        self.assertEqual(len(f),17)
        self.assertEqual({m['key'] for m in profile('MEDICAL_REPORT') if m['required']},{'name','gender','exam_date'})
        self.assertFalse({'marital_status','phone','work_unit','report_no','package_name'} & f.keys())
        self.assertEqual(f['institution']['value'],'示例医院')
        self.assertEqual(f['report_date']['value'],'2026-10-02')
        self.assertEqual(f['height']['value'],'174 cm')
        self.assertEqual(f['weight']['value'],'78 kg')
        self.assertEqual(f['blood_pressure']['value'],'129/84 mmHg')
        self.assertEqual(f['exam_project_count']['value'],3)
        self.assertFalse(f['exam_project_count']['editable'])
        self.assertEqual(f['ecg']['value'],'未见异常')
        self.assertEqual(f['urine_routine']['source']['page'],9)
        self.assertIn('潜血 +3',f['urine_routine']['value'])
        self.assertIn('蛋白质 +1',f['urine_routine']['value'])
        self.assertFalse(any(i['name']=='小结' for i in r['items']))
        self.assertEqual(sum(confirmed_abnormal(i) for i in r['items']),3)
        self.assertFalse(required_export_blockers(r))

    def test_missing_or_misleading_evidence_never_defaults_normal(self):
        r=extract_result([block('目录：血常规 ...... 9\n健康建议：尿常规正常'),
            block(table([['项目名称','检测结果'],['白细胞','4.0']])),
            block('八、本次体检项目意义 ...... 12'),
            block(table([['项目','临床意义'],['血常规','正常项目示例']]),y=.3)],'MEDICAL_REPORT')
        f=fields(r)
        self.assertIsNone(f['exam_project_count']['value'])
        for key in ('blood_routine','urine_routine','ecg','dr_imaging','abdominal_ultrasound','liver_function'):
            self.assertIsNone(f[key]['value'],key)

    def test_count_duplicates_combined_name_and_manual_refresh(self):
        r=extract_result(fixture(),'MEDICAL_REPORT')
        before=copy.deepcopy(r)
        changed=apply_review(r,{'fields':{'exam_projects':'身高体重血压；血常规；血常规；DR胸部正位检查'}},actor='test')
        self.assertEqual(fields(changed)['exam_project_count']['value'],3)
        self.assertEqual(r,before)
        with self.assertRaises(TaskError):
            apply_review(changed,{'fields':{'exam_project_count':30}},actor='test')
        with self.assertRaises(TaskError):
            apply_review(changed,{'absence_reasons':{'exam_project_count':'NOT_ON_DOCUMENT'}},actor='test')

    def test_item_correction_refreshes_auto_summary_but_preserves_manual_value(self):
        r=extract_result(fixture(),'MEDICAL_REPORT')
        item=next(i for i in r['items'] if i['name']=='潜血')
        changed=apply_review(r,{'items':[{'item_id':item['item_id'],'result_raw':'+2','diagnosis':'复核潜血阳性'}]},actor='test')
        self.assertIn('潜血 +2',fields(changed)['urine_routine']['value'])
        manual=apply_review(changed,{'fields':{'urine_routine':'人工核对的尿检结论'}},actor='test')
        again=apply_review(manual,{'items':[{'item_id':item['item_id'],'result_raw':'+1'}]},actor='test')
        self.assertEqual(fields(again)['urine_routine']['value'],'人工核对的尿检结论')
        absent=apply_review(changed,{'absence_reasons':{'urine_routine':'NOT_ON_DOCUMENT'}},actor='test')
        self.assertIsNone(fields(public_result(absent))['urine_routine']['value'])

    def test_history_projection_preserves_sources_and_does_not_reintroduce_retired_fields(self):
        r=extract_result(fixture(),'MEDICAL_REPORT')
        r['fields']=[f for f in r['fields'] if f['key'] in {'name','gender','age','exam_date','exam_projects'}]
        r['fields'].append(dict(key='phone',value='123456'))
        r['basic_info']={f['key']:f for f in r['fields']}
        before=copy.deepcopy(r)
        out=public_result(r)
        self.assertEqual(len(out['fields']),17)
        self.assertEqual(fields(out)['exam_project_count']['value'],3)
        self.assertEqual(fields(out)['height']['value'],'174 cm')
        self.assertNotIn('phone',out['basic_info'])
        self.assertEqual(r,before)

    def test_workbook_exact_template_saved_values_and_one_report_per_row(self):
        r=extract_result(fixture(),'MEDICAL_REPORT')
        r=apply_review(r,{'fields':{'institution':'修正机构','ecg':'人工心电图结论'}},actor='test')
        before=copy.deepcopy(r)
        book=load_workbook(io.BytesIO(workbook_bytes([r,r])))
        self.assertEqual(book.sheetnames,['审核记录台账'])
        sheet=book.active
        self.assertEqual(list(next(sheet.values)),['序号','姓名','性别','年龄','体检日期','体检机构','报告日期','身高','体重','血压',
            '血常规','尿常规','心电图','DR影像','腹部彩超','肝功五项（生化检测）','体检项目数量','异常指标'])
        self.assertEqual(sheet.max_row,3)
        self.assertEqual((sheet['A2'].value,sheet['A3'].value),(1,2))
        self.assertEqual(sheet['F2'].value,'修正机构')
        self.assertEqual(sheet['M2'].value,'人工心电图结论')
        self.assertEqual(sheet['Q2'].value,3)
        self.assertIn('潜血 +3',sheet['R2'].value)
        self.assertIn('尿蛋白质阳性',sheet['R2'].value)
        self.assertNotIn('身高',sheet['R2'].value)
        self.assertIsNone(sheet['K2'].value)
        self.assertFalse(any(c.data_type=='f' for row in sheet for c in row))
        self.assertEqual(r,before)

    def test_missing_catalogue_and_abnormality_are_not_claimed_zero_or_none(self):
        r=extract_result([],'MEDICAL_REPORT')
        book=load_workbook(io.BytesIO(workbook_bytes([r])))
        self.assertIsNone(book.active['Q2'].value)
        self.assertIsNone(book.active['R2'].value)

    def test_conflicting_normal_conclusion_cannot_hide_abnormal_measurement(self):
        evidence=[block('尿常规检验报告单'),block(table([['项目名称','检测结果','风险等级'],
            ['潜血','+3','E3'],['小结','未见异常','A']]),y=.2)]
        r=extract_result(evidence,'MEDICAL_REPORT')
        f=fields(r)['urine_routine']
        self.assertIsNone(f['value'])
        self.assertTrue(any(i['code']=='CONFLICT' for i in f['review_issues']))
        self.assertIn('潜血 +3',load_workbook(io.BytesIO(workbook_bytes([r]))).active['R2'].value)
        resolved=apply_review(r,{'fields':{'urine_routine':'原件矛盾，人工核对潜血+3'}},actor='test')
        self.assertEqual(fields(public_result(resolved))['urine_routine']['value'],'原件矛盾，人工核对潜血+3')

    def test_separate_pressure_sources_and_unknown_heading_are_not_guessed(self):
        r=extract_result([block('尿常规检验报告单'),block('其他检查',y=.2),
            block(table([['项目名称','检测结果','风险等级'],['潜血','+3','E3']]),y=.3),
            block(table([['项目名称','检测结果','单位'],['收缩压','120','mmHg']]),page=8),
            block(table([['项目名称','检测结果','单位'],['舒张压','80','mmHg']]),page=9)],'MEDICAL_REPORT')
        self.assertIsNone(fields(r)['urine_routine']['value'])
        self.assertIsNone(fields(r)['blood_pressure']['value'])
        self.assertEqual(sum(confirmed_abnormal(i) for i in r['items']),1)


if __name__=='__main__': unittest.main()
