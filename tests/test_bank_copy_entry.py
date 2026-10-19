"""Separate business requirements while reusing source-bound bank recognition."""
import copy
import io
import json
import unittest
import time
from unittest.mock import Mock, patch
from uuid import uuid4

from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import text

import test_business_api as base
import test_required_field_settings as settings
from test_bank_photocopy import block, fields
from ocr_app.doc_profiles import catalogue, profile
from ocr_app.document_classification import suggestion
from ocr_app.extraction import extract_result
from ocr_app.exporting import required_export_blockers, workbook_bytes
from ocr_app.person_packages import name_suggestions
from ocr_app.review import apply_review
from ocr_app.task_store import public_result


class BankCopyEntryTests(unittest.TestCase):
    def test_copy_pipeline_shares_upright_regions_and_glm_lane(self):
        from ocr_app.pipeline import process_document
        image = Image.new('RGB', (100, 150), 'white')
        image.putpixel((40, 40), (0, 0, 0))
        stream = io.BytesIO(); image.save(stream, format='PNG')
        plans = [dict(box=(0, 0, 100, 70), task='ocr', bank_role='front'),
                 dict(box=(0, 70, 100, 140), task='ocr', bank_role='supplement')]
        recognize = Mock(side_effect=[dict(text='合成银行\n4111111111111111', termination_reason='eos'),
                                     dict(text='李小明 13800000000\n东城支行', termination_reason='eos')])
        with patch('ocr_app.bank_photocopy.plan_bank_regions', return_value=(plans, [])) as planner, \
             patch('ocr_app.image_preprocessing.classify_bank_orientation', return_value=(0, .99, .8)) as orientation, \
             patch('ocr_app.image_preprocessing.estimate_skew', return_value=(0., 'already_level')):
            result = process_document(stream.getvalue(), 'BANK_CARD_COPY', recognize=recognize,
                progress=lambda *a: None, save_page=lambda *a: 'page', record_attempt=lambda *a: None,
                cancelled=lambda: False, deadline=time.monotonic() + 10)
        orientation.assert_called_once(); planner.assert_called_once()
        self.assertEqual(recognize.call_count, 2)
        self.assertTrue(all(call.args[0] == 'glm_ocr' for call in recognize.call_args_list))
        self.assertEqual(result['doc_type'], 'BANK_CARD_COPY')
        self.assertEqual(fields(result)['holder_name']['value'], '李小明')
        self.assertEqual(fields(result)['branch_name']['source']['bbox'], [0, .466667, 1, .933333])
        self.assertEqual(result['pages'][0]['image_preprocessing']['version'], 'bank-card-upright-v1')

    def test_profiles_have_independent_required_defaults_and_overrides(self):
        required = lambda kind, policy=None: {f['key'] for f in profile(kind, required_overrides=policy) if f['required']}
        self.assertEqual(required('BANK_CARD'), {'card_number'})
        self.assertEqual(required('BANK_CARD_COPY'), {'holder_name', 'card_number', 'bank_name', 'branch_name'})
        policy = {'BANK_CARD': ['card_number', 'bank_name'], 'BANK_CARD_COPY': ['card_number', 'holder_name']}
        self.assertEqual(required('BANK_CARD', policy), {'card_number', 'bank_name'})
        self.assertEqual(required('BANK_CARD_COPY', policy), {'card_number', 'holder_name'})
        self.assertEqual([f['key'] for f in profile('BANK_CARD')], [f['key'] for f in profile('BANK_CARD_COPY')])
        labels = {p['doc_type']: p['label'] for p in catalogue()}
        self.assertEqual(labels['BANK_CARD_COPY'], '银行卡复印件')
        self.assertEqual(labels['BUSINESS_LICENSE'], '营业执照')

    def test_only_pan_is_sufficient_for_ordinary_card_but_not_copy(self):
        evidence = [block('4111111111111111', 'front')]
        ordinary = extract_result(evidence, 'BANK_CARD')
        photocopy = extract_result(evidence, 'BANK_CARD_COPY')
        self.assertFalse(required_export_blockers(ordinary))
        self.assertEqual({b['field'] for b in required_export_blockers(photocopy)}, {'holder_name', 'bank_name', 'branch_name'})
        self.assertEqual(photocopy['doc_type'], 'BANK_CARD_COPY')
        self.assertTrue(all(fields(photocopy)[k]['required'] for k in ('holder_name', 'bank_name', 'branch_name')))

    def test_copy_retains_regions_candidates_and_save_resolves_only_selected_number(self):
        result = extract_result([block('合成银行\n4111111111111111', 'front'),
            block('4111111111111112\n李小明 13800000000\n东城支行', 'supplement')], 'BANK_CARD_COPY')
        pan = fields(result)['card_number']
        self.assertIsNone(pan['value'])
        self.assertEqual([c['reason'] for c in pan['candidates']], ['卡面识别', '补充文字识别'])
        fixed = apply_review(result, {'fields': {'card_number': '4111111111111112'}, 'confirm_export_disputes': True}, actor='operator')
        self.assertEqual(fixed['doc_type'], 'BANK_CARD_COPY')
        self.assertFalse(required_export_blockers(fixed))
        self.assertEqual(len(fields(fixed)['card_number']['candidates']), 2)
        self.assertEqual(fields(fixed)['branch_name']['source']['raw_text'], '东城支行')

    def test_history_projection_is_read_only_and_preserves_manual_confirmations(self):
        original = extract_result([block('4111111111111111')], 'BANK_CARD')
        original.update(run_id=str(uuid4()), result_revision=3)
        original = apply_review(original, {'fields': {'holder_name': '历史姓名', 'branch_name': '历史支行',
                                                     'card_number': '41111111111111111111'},
                                          'confirm_export_disputes': True}, actor='old-operator')
        self.assertTrue(original['_export_confirmations'])
        before = copy.deepcopy(original)
        projected = public_result(original, {'BANK_CARD_COPY': ['holder_name', 'branch_name']})
        self.assertEqual(original, before)
        self.assertEqual(projected['doc_type'], 'BANK_CARD')
        self.assertEqual(fields(projected)['holder_name']['value_origin'], 'manual')
        self.assertEqual(fields(projected)['branch_name']['value'], '历史支行')
        self.assertEqual(projected['result_revision'], 3)
        self.assertEqual(original.get('_export_confirmations'), before.get('_export_confirmations'))
        self.assertFalse(required_export_blockers(original))

    def test_copy_exports_all_bank_fields_in_own_sheet_and_person_summary(self):
        result = apply_review(extract_result([], 'BANK_CARD_COPY'), {'fields': dict(
            holder_name='李小明', card_number='4111111111111111', bank_name='合成银行', branch_name='东城支行')}, actor='operator')
        result.update(task_id=str(uuid4()), run_id=str(uuid4()), result_revision=1)
        book = load_workbook(io.BytesIO(workbook_bytes([result], [dict(person_name='李小明', results=[result])])))
        self.assertIn('银行卡复印件', book.sheetnames)
        rows = list(book['人员信息总表'].values)
        row = dict(zip(rows[0], rows[1]))
        self.assertEqual(row['银行卡复印件·姓名'], '李小明')
        self.assertEqual(row['银行卡复印件·卡号'], '4111111111111111')
        self.assertEqual(row['银行卡复印件·开户支行'], '东城支行')
        self.assertEqual(name_suggestions([result])[0]['name'], '李小明')

    def test_classification_does_not_guess_ordinary_card_from_absent_handwriting(self):
        ambiguous = suggestion('中国银行\nUNIONPAY\n4111111111111111', '1.jpg')
        self.assertIsNone(ambiguous['doc_type'])
        self.assertEqual(set(ambiguous['candidates']), {'BANK_CARD', 'BANK_CARD_COPY'})
        self.assertEqual(suggestion('中国银行\n4111111111111111\n开户支行：东城支行', '1.jpg')['doc_type'], 'BANK_CARD_COPY')
        self.assertEqual(suggestion('UNIONPAY\n4111111111111111', '银行卡复印件.jpg')['doc_type'], 'BANK_CARD_COPY')
        self.assertIsNone(suggestion('身份证正面公民身份号码 姓名 性别 民族\n银行卡复印件', '1.jpg')['doc_type'])


class BankCopyAPITests(unittest.TestCase):
    setUp = base.BusinessAPITests.setUp
    upload = base.BusinessAPITests.upload
    admin = settings.RequiredFieldSettingsTests.admin
    save = settings.RequiredFieldSettingsTests.save
    origin = settings.RequiredFieldSettingsTests.origin

    def test_upload_catalogue_and_policy_do_not_rewrite_old_records(self):
        self.admin()
        old = self.upload('BANK_CARD').json()
        task_before = self.business.store.task(old['task_ids'][0])
        with self.business.engine.connect() as connection:
            stored_before = dict(connection.execute(text('SELECT * FROM tasks WHERE id=:id'), {'id': task_before['task_id']}).mappings().one())
        self.assertEqual(self.save('BANK_CARD', ['card_number', 'bank_name']).status_code, 200)
        self.assertEqual(self.save('BANK_CARD_COPY', ['holder_name', 'card_number', 'bank_name', 'branch_name']).status_code, 200)
        policy = self.business.store.required_policy()['overrides']
        self.assertEqual(policy['BANK_CARD'], ['card_number', 'bank_name'])
        created = self.upload('BANK_CARD_COPY')
        self.assertEqual(created.status_code, 202, created.text)
        self.assertEqual(self.business.store.task(created.json()['task_ids'][0])['doc_type'], 'BANK_CARD_COPY')
        self.assertEqual(self.business.store.task(task_before['task_id'])['run_id'], task_before['run_id'])
        with self.business.engine.connect() as connection:
            stored_after = dict(connection.execute(text('SELECT * FROM tasks WHERE id=:id'), {'id': task_before['task_id']}).mappings().one())
        self.assertEqual(stored_after, stored_before)
        catalogue_response = self.client.get('/api/v1/doc-types', headers=self.headers)
        self.assertEqual(catalogue_response.status_code, 200)
        self.assertEqual(len(catalogue_response.json()), 14)

    def test_person_package_blocks_copy_missing_fields_then_exports_saved_four_fields(self):
        other = io.BytesIO(); Image.new('RGB', (20, 20), 'blue').save(other, format='PNG')
        created = self.client.post('/api/v1/batches', headers=self.headers,
            data={'file_types': json.dumps(['BANK_CARD', 'BANK_CARD_COPY'])},
            files=[('files', ('card.png', self.image, 'image/png')),
                   ('files', ('copy.png', other.getvalue(), 'image/png'))])
        self.assertEqual(created.status_code, 202, created.text)
        tasks = {}
        for _ in range(2):
            claimed = self.business.store.claim_next()
            result = extract_result([block('4111111111111111')], claimed['doc_type'])
            result.update(status='NEED_REVIEW', pages=[])
            self.assertTrue(self.business.store.publish(claimed, result))
            tasks[claimed['doc_type']] = self.business.store.task(claimed['id'])
        export_body = {'batch_id': created.json()['batch_id'], 'format': 'xlsx', 'masked': False}
        self.assertEqual(self.client.post('/api/v1/export', headers=self.headers, json=export_body).status_code, 409)
        ordinary_before = self.business.store.task(tasks['BANK_CARD']['task_id'])
        task = tasks['BANK_CARD_COPY']
        reviewed = self.client.post('/api/v1/tasks/' + task['task_id'] + '/review', headers=self.headers,
            json=dict(run_id=task['run_id'], expected_revision=task['result_revision'],
                      fields=dict(holder_name='李小明', bank_name='合成银行', branch_name='东城支行'), confirm_export_disputes=True))
        self.assertEqual(reviewed.status_code, 200, reviewed.text)
        self.assertEqual(self.business.store.task(ordinary_before['task_id']), ordinary_before)
        response = self.client.post('/api/v1/export', headers=self.headers, json=export_body)
        self.assertEqual(response.status_code, 200, response.text if response.status_code != 200 else '')
        book = load_workbook(io.BytesIO(response.content))
        self.assertTrue({'银行卡', '银行卡复印件', '人员信息总表'} <= set(book.sheetnames))
        public = self.client.get('/api/v1/tasks/' + task['task_id'] + '/result', headers=self.headers).json()
        self.assertEqual(public['doc_type'], 'BANK_CARD_COPY')
        self.assertEqual(fields(public)['holder_name']['value'], '李小明')

    def test_copy_pdf_is_rejected_before_storage(self):
        import fitz
        with fitz.open() as document:
            document.new_page()
            pdf = document.tobytes()
        response = self.client.post('/api/v1/batches', headers=self.headers,
            data={'doc_type_hint': 'BANK_CARD_COPY'}, files={'files': ('copy.pdf', pdf, 'application/pdf')})
        self.assertEqual(response.status_code, 415, response.text)
        self.assertEqual(self.business.store.batches()['total'], 0)
