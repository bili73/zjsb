import copy
from datetime import datetime, timedelta, timezone
import io
import sys
import time
import unittest
from pathlib import Path
from unittest.mock import Mock, patch
from uuid import uuid4
from zipfile import ZipFile

from openpyxl import load_workbook
from PIL import Image, ImageDraw

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'backend'))
from ocr_app.pipeline import process_document
from ocr_app.review import apply_review
from ocr_app.table_recognition import decode_table_output, effective_tables, infer_header_rows, validate_table_grid, table_blockers
from ocr_app.table_export import table_export_bundle, table_workbook_bytes
from ocr_app.task_store import TaskError
import test_business_api as api_fixtures


def table():
    return dict(table_id=str(uuid4()),title='合成表格',header_rows=1,rows=[['编号','备注'],['001','=1+1'],['','第二行\n末行']],
                merges=[],partial_acknowledged=False,incomplete=False,confidence=None,title_candidates=[],header_suggestion=1,header_suggestion_source='structure',
                source=dict(page=1,bbox=[0,0,1,1],extraction_method='ocr'),candidates=[],outside_text=[],quality_warnings=[])


def draft(t):
    return {k:copy.deepcopy(t[k]) for k in ('table_id','title','header_rows','rows','merges','partial_acknowledged')}


class PhotoTableTests(unittest.TestCase):
    def test_existing_workbooks_use_twelve_point_headers_and_data(self):
        from ocr_app.exporting import workbook_bytes
        from ocr_app.extraction import extract_result
        from ocr_app.doc_profiles import SUPPORTED
        results=[dict(extract_result([],kind),task_id=str(uuid4()),run_id=str(uuid4()),result_revision=0)
                 for kind in SUPPORTED if kind!='TABLE']
        book=load_workbook(io.BytesIO(workbook_bytes(results)))
        for sheet in book:
            for row in sheet:
                for cell in row:
                    if cell.row > 1:
                        self.assertEqual(cell.fill.fgColor.rgb[-6:],'FFFFFF',(sheet.title,cell.coordinate))
                    if cell.value is not None:
                        self.assertEqual((cell.font.name,cell.font.sz),('Microsoft YaHei',12),(sheet.title,cell.coordinate))

    def test_review_keeps_initial_values_and_rejects_forged_source(self):
        t=table();original=dict(doc_type='TABLE',tables=[t],review_issues=[])
        d=draft(t);d['rows'][1][0]='0009'
        result=apply_review(original,{'tables':[d]},actor='operator')
        self.assertEqual(original['tables'][0]['rows'][1][0],'001')
        self.assertEqual(result['tables'],original['tables'])
        self.assertEqual(result['reviewed_tables'][0]['rows'][1][0],'0009')
        self.assertEqual(result['status'],'COMPLETED')
        for extra in [{'source':{}},{'table_id':str(uuid4())},{'table_id':[]},{'incomplete':False}]:
            with self.subTest(extra=extra),self.assertRaises(TaskError):
                apply_review(original,{'tables':[dict(d,**extra)]},actor='operator')

    def test_incomplete_needs_explicit_ack_but_can_save_draft(self):
        t=table();t['incomplete']=True
        original=dict(doc_type='TABLE',tables=[t],review_issues=[])
        result=apply_review(original,{'tables':[draft(t)]},actor='operator')
        self.assertEqual(table_blockers(result)[0]['code'],'INCOMPLETE_OUTPUT')
        self.assertEqual(result['status'],'NEED_REVIEW')
        d=draft(t);d['partial_acknowledged']=True;d['title']='';d['header_rows']=None
        result=apply_review(original,{'tables':[d]},actor='operator')
        self.assertEqual([p['field'] for p in table_blockers(result)],['tables[0].header_rows'])
        d['header_rows']=0
        self.assertFalse(table_blockers(apply_review(original,{'tables':[d]},actor='operator')))

    def test_merge_cannot_destroy_values_and_grid_bounds_are_enforced(self):
        t=draft(table());t['merges']=[dict(row=0,column=0,rowspan=1,colspan=2)]
        with self.assertRaises(ValueError):validate_table_grid(t)
        t['rows'][0][1]='';validate_table_grid(t)
        t['merges']*=2
        with self.assertRaises(ValueError):validate_table_grid(t)
        for rows in [[],[['a'],['b','c']],[['x']*201],[['x']*20]*501,[['\x00']],[['x'*32768]]]:
            t=draft(table());t['rows']=rows
            with self.subTest(size=len(rows)),self.assertRaises(ValueError):validate_table_grid(t)

    def test_excel_matches_every_cell_and_never_coerces_text(self):
        t=draft(table());book=load_workbook(io.BytesIO(table_workbook_bytes(t)))
        sheet=book.active
        self.assertEqual([[cell.value or '' for cell in row] for row in list(sheet)[1:]],t['rows'])
        self.assertEqual(sheet['A3'].value,'001');self.assertEqual(sheet['B3'].data_type,'s')
        self.assertEqual(sheet['A1'].font.sz,21);self.assertEqual(sheet['A2'].font.sz,12)
        self.assertEqual(sheet['A2'].fill.fgColor.rgb[-6:],'EF760E')
        for row in sheet.iter_rows(min_row=3):
            for cell in row:self.assertEqual(cell.fill.fgColor.rgb[-6:],'FFFFFF')
        self.assertEqual(sheet['A1'].alignment.vertical,'center')
        t['rows'][0][1]='';t['merges']=[dict(row=0,column=0,rowspan=1,colspan=2)]
        self.assertIn('A2:B2',str(load_workbook(io.BytesIO(table_workbook_bytes(t))).active.merged_cells))

    def test_zip_duplicate_titles_and_path_like_titles(self):
        t=table();t['title']='../同名表';items=[dict(tables=[t],reviewed_tables=[draft(t)])]*2
        data,media,name=table_export_bundle(items)
        self.assertEqual(media,'application/zip')
        with ZipFile(io.BytesIO(data)) as archive:
            self.assertEqual(len(set(archive.namelist())),2)
            self.assertTrue(all('/' not in n and '\\' not in n for n in archive.namelist()))
            for path in archive.namelist():self.assertEqual(load_workbook(io.BytesIO(archive.read(path))).active['A3'].value,'001')

    def test_auto_header_count_can_export_and_saved_override_controls_excel(self):
        parsed=decode_table_output('<table><tr><th rowspan="2">Name</th><th colspan="2">Checks</th></tr><tr><th>A</th><th>B</th></tr><tr><td>X</td><td>1</td><td>2</td></tr></table>')[0]
        count,source=infer_header_rows(parsed)
        t=dict(table(),rows=parsed['rows'],merges=parsed['merges'],header_rows=None,header_suggestion=count,header_suggestion_source=source)
        original=dict(doc_type='TABLE',tables=[t],review_issues=[])
        self.assertFalse(table_blockers(original))
        sheet=load_workbook(io.BytesIO(table_export_bundle([original])[0])).active
        self.assertEqual(sheet.freeze_panes,f'A{count+2}')
        self.assertIsNone(original['tables'][0]['header_rows'])
        for saved_count in (count,0):
            d=dict(draft(t),header_rows=saved_count)
            reviewed=apply_review(original,{'tables':[d]},actor='operator')
            self.assertFalse(table_blockers(reviewed))
            self.assertIsNone(reviewed['tables'][0]['header_rows'])
            sheet=load_workbook(io.BytesIO(table_workbook_bytes(reviewed['reviewed_tables'][0]))).active
            self.assertEqual(sheet.freeze_panes,f'A{saved_count+2}')
            self.assertEqual(sheet['B3'].fill.fgColor.rgb[-6:],'EF760E' if saved_count else 'FFFFFF')
            self.assertEqual(sheet['B4'].fill.fgColor.rgb[-6:],'FFFFFF')
            self.assertEqual((sheet['A1'].font.sz,sheet['B3'].font.sz,sheet['B4'].font.sz),(21,12,12))
            self.assertEqual(sheet['A4'].value,'X')

    def test_export_projection_preserves_originals_and_prefers_saved_blanks(self):
        t=dict(table(),title='',header_rows=None,title_candidates=[dict(text='唯一标题',box=[0,0,10,10])])
        original=dict(doc_type='TABLE',tables=[t],review_issues=[])
        before=copy.deepcopy(original)
        self.assertFalse(table_blockers(original))
        data,_,name=table_export_bundle([original])
        self.assertEqual(name,'唯一标题.xlsx')
        sheet=load_workbook(io.BytesIO(data)).active
        self.assertEqual(sheet.freeze_panes,'A3')
        self.assertEqual(sheet['A3'].value,'001')
        self.assertEqual(sheet['B3'].data_type,'s')
        self.assertEqual(original,before)
        saved=dict(draft(t),title='',header_rows=None)
        result=dict(original,reviewed_tables=[saved])
        self.assertEqual(table_blockers(result),[dict(field='tables[0].header_rows',code='MISSING')])
        saved.update(title='人工标题',header_rows=0);saved['rows'][1][0]='人工编号'
        self.assertFalse(table_blockers(result))
        sheet=load_workbook(io.BytesIO(table_export_bundle([result])[0])).active
        self.assertEqual(sheet['A1'].value,'人工标题');self.assertEqual(sheet['A3'].value,'人工编号')
        self.assertEqual(sheet.freeze_panes,'A2')
        self.assertEqual(original,before)

    def test_automatic_export_does_not_guess_headers_or_accept_invalid_output(self):
        base=dict(table(),title='',header_rows=None,title_candidates=[dict(text='候选',box=[0,0,10,10])])
        cases=[({'header_suggestion_source':None},'tables[0].header_rows','MISSING'),
               ({'header_suggestion':99},'tables[0].header_rows','MISSING'),
               ({'incomplete':True},'tables[0]','INCOMPLETE_OUTPUT'),
               ({'rows':[['x'],['x','y']]},'tables[0]','INVALID_TABLE')]
        for change,field,code in cases:
            result=dict(tables=[dict(base,**change)])
            with self.subTest(change=change):
                self.assertIn(dict(field=field,code=code),table_blockers(result))
                with self.assertRaises(ValueError):table_export_bundle([result])
        for saved in ([],None):
            self.assertTrue(table_blockers(dict(tables=[base],reviewed_tables=saved)))

    def test_ranked_valid_titles_are_suggestions_and_never_replace_saved_values(self):
        candidates=[dict(text=text,box=[0,0,10,10]) for text in ['', 'x'*201, 'bad\x00title', '  优先标题  ', '另一标题']]
        original=dict(tables=[dict(table(),title='',title_candidates=candidates)],_source_filename='原始图片.jpg')
        before=copy.deepcopy(original)
        self.assertEqual(effective_tables(original)[0]['title'],'优先标题')
        self.assertEqual(table_export_bundle([original])[2],'优先标题.xlsx')
        for title in ['', '   ', '人工填写']:
            with self.subTest(title=title):
                result=dict(original,reviewed_tables=[dict(draft(original['tables'][0]),title=title)])
                self.assertEqual(effective_tables(result)[0]['title'],title)
                self.assertFalse(table_blockers(result))
                self.assertEqual(table_export_bundle([result])[2],('人工填写' if title.strip() else '原始图片')+'.xlsx')
        self.assertEqual(original,before)

    def test_untitled_workbooks_start_at_matrix_and_shift_all_row_metadata(self):
        for count in (0,1,2):
            for merged in (False,True):
                with self.subTest(headers=count,merged=merged):
                    t=dict(table(),title='   ',header_rows=count)
                    if merged:
                        t['rows'][0][1]='';t['merges']=[dict(row=0,column=0,rowspan=1,colspan=2)]
                    original=dict(tables=[t],_source_filename='签到表照片.jpg')
                    data,_,name=table_export_bundle([original])
                    self.assertEqual(name,'签到表照片.xlsx')
                    sheet=load_workbook(io.BytesIO(data)).active
                    self.assertEqual([[c.value or '' for c in row] for row in sheet],t['rows'])
                    self.assertEqual(sheet.freeze_panes,f'A{count+1}' if count else None)
                    self.assertEqual(sheet.print_title_rows,f'$1:${count}' if count else None)
                    self.assertEqual(sheet.auto_filter.ref,'A1:B3' if count==1 and not merged else None)
                    self.assertEqual(str(sheet.merged_cells),'A1:B1' if merged else '')
                    self.assertEqual(sheet['A1'].font.sz,12)
                    self.assertEqual(sheet['A1'].fill.fgColor.rgb[-6:],'EF760E' if count else 'FFFFFF')
                    self.assertEqual(sheet['A2'].value,'001');self.assertEqual(sheet['B2'].data_type,'s')

    def test_untitled_names_use_safe_original_basename_or_timestamp_and_zip_dedup(self):
        t=dict(table(),title='')
        for original,expected in [('目录/签到.jpg','签到.xlsx'),('C:\\资料\\签到.PNG','签到.xlsx'),('a:b?.jpg','a-b-.xlsx')]:
            with self.subTest(original=original):
                self.assertEqual(table_export_bundle([dict(tables=[t],_source_filename=original)])[2],expected)
        for original in (None,'','.jpg','NUL.png'):
            with self.subTest(original=original):
                name=table_export_bundle([dict(tables=[t],_source_filename=original)])[2]
                self.assertRegex(name,r'^表格_\d{8}_\d{6}\.xlsx$')
        inputs=[dict(tables=[t],_source_filename='Same.jpg'),dict(tables=[t],_source_filename='same.png'),dict(tables=[dict(t,title='Same')])]
        data,_,_=table_export_bundle(inputs)
        with ZipFile(io.BytesIO(data)) as archive:
            self.assertEqual(archive.namelist(),['Same.xlsx','same (2).xlsx','Same (3).xlsx'])
        with patch('ocr_app.table_export.datetime') as clock:
            clock.now.return_value=datetime(2026,10,15,16,30,1,tzinfo=timezone.utc).astimezone(timezone(timedelta(hours=8)))
            self.assertEqual(table_export_bundle([dict(tables=[t])])[2],'表格_20261016_003001.xlsx')
            clock.now.assert_called_once_with(timezone(timedelta(hours=8)))

    def test_pipeline_adds_header_evidence_without_extra_model_calls(self):
        stream=io.BytesIO();Image.new('RGB',(300,200),'white').save(stream,format='PNG')
        raw='<fcel>序号<fcel>名称<nl><fcel>1<fcel>甲<nl><fcel>2<fcel>乙<nl>'
        for reason,hint,expected in [('eos',[0,0],(1,'structure')),('eos',[9,9],(None,None)),('token_limit',[0,0],(None,None))]:
            with self.subTest(reason=reason,hint=hint):
                recognize=Mock(return_value=dict(text=raw,termination_reason=reason))
                with patch('ocr_app.photo_table_pipeline.nearby_text',return_value=[]), patch('ocr_app.photo_table_pipeline.grid_dimensions_hint',return_value=hint):
                    result=process_document(stream.getvalue(),'TABLE',recognize=recognize,progress=Mock(),save_page=lambda *a:str(uuid4()),
                        record_attempt=Mock(),cancelled=lambda:False,deadline=time.monotonic()+180)
                t=result['tables'][0]
                self.assertEqual((t['header_suggestion'],t['header_suggestion_source']),expected)
                self.assertIsNone(t['header_rows'])
                self.assertEqual(recognize.call_count,1 if reason=='eos' else 2)

    def test_pipeline_reuses_paddle_then_bounded_glm_and_retains_attempts(self):
        image=Image.new('RGB',(500,300),'white');draw=ImageDraw.Draw(image)
        for x in range(30,481,225):draw.line((x,100,x,260),fill='black',width=2)
        for y in range(100,261,80):draw.line((30,y,480,y),fill='black',width=2)
        stream=io.BytesIO();image.save(stream,format='PNG')
        attempts=[];recognize=Mock(side_effect=[dict(text='<table><tr>',termination_reason='token_limit'),dict(text='<fcel>A<fcel>B<nl><fcel>001<ecel><nl>',termination_reason='eos')])
        with patch('ocr_app.photo_table_pipeline.nearby_text',return_value=[]):
            result=process_document(stream.getvalue(),'TABLE',recognize=recognize,progress=Mock(),save_page=lambda *a:str(uuid4()),
                record_attempt=attempts.append,cancelled=lambda:False,deadline=time.monotonic()+180)
        self.assertEqual([c.args[0] for c in recognize.call_args_list],['paddleocr_vl','glm_ocr'])
        self.assertTrue(all(c.kwargs['generation_timeout_seconds']==40 for c in recognize.call_args_list))
        self.assertEqual(len(attempts),2);self.assertEqual(attempts[0]['prediction']['termination_reason'],'token_limit')
        self.assertFalse(result['tables'][0]['incomplete']);self.assertEqual(result['tables'][0]['rows'][1],['001',''])
        self.assertIsNone(result['tables'][0]['header_suggestion'])
        self.assertIsNone(result['tables'][0]['header_suggestion_source'])
        self.assertEqual(result['pages'][0]['regions'][0]['status'],'recognized')

    def test_timeout_retains_source_and_does_not_start_another_model(self):
        from ocr_app.documents import DocumentError
        stream=io.BytesIO();Image.new('RGB',(200,100),'white').save(stream,format='PNG')
        attempts=[];recognize=Mock(side_effect=DocumentError('MODEL_REQUEST_TIMEOUT'))
        with patch('ocr_app.photo_table_pipeline.nearby_text',return_value=[]):
            result=process_document(stream.getvalue(),'TABLE',recognize=recognize,progress=Mock(),save_page=lambda *a:str(uuid4()),
                record_attempt=attempts.append,cancelled=lambda:False,deadline=time.monotonic()+180)
        self.assertEqual(result['error']['code'],'EXTRACT_TIMEOUT');self.assertEqual(recognize.call_count,1)
        self.assertEqual(len(result['pages']),1);self.assertTrue(result['tables'][0]['incomplete'])
        self.assertEqual(len(attempts),1)


class PhotoTableAPITests(unittest.TestCase):
    setUp=api_fixtures.BusinessAPITests.setUp
    upload=api_fixtures.BusinessAPITests.upload
    task=api_fixtures.BusinessAPITests.task

    def test_cross_account_and_readonly_review_cannot_change_tables(self):
        import test_batch_ownership as ownership
        self.password=ownership.BatchOwnershipTests.password
        task,t=self.prepared();url=f'/api/v1/tasks/{task["task_id"]}/review'
        body=dict(run_id=task['run_id'],expected_revision=task['result_revision'],tables=[draft(t)])
        _,other,_=ownership.BatchOwnershipTests.account(self,'other')
        _,reader,_=ownership.BatchOwnershipTests.account(self,'reader','viewer')
        for headers,status in [(other,404),(reader,403)]:
            self.assertEqual(self.client.post(url,json=body,headers=headers).status_code,status)
            self.assertEqual(self.client.post('/api/v1/export',json={'task_ids':[task['task_id']]},headers=headers).status_code,status)
            self.assertEqual(self.business.store.task(task['task_id'])['result_revision'],task['result_revision'])

    def test_revoked_session_cannot_receive_a_serialized_table(self):
        import test_batch_ownership as ownership
        self.password=ownership.BatchOwnershipTests.password
        _,self.headers,token=ownership.BatchOwnershipTests.account(self,'tableowner','admin')
        task,t=self.prepared()
        self.client.post(f'/api/v1/tasks/{task["task_id"]}/review',headers=self.headers,
            json=dict(run_id=task['run_id'],expected_revision=task['result_revision'],tables=[draft(t)])).raise_for_status()
        def revoke(results):
            value=table_export_bundle(results);self.business.accounts.logout(token);return value
        with patch('ocr_app.table_export.table_export_bundle',side_effect=revoke):
            response=self.client.post('/api/v1/export',headers=self.headers,json=dict(task_ids=[task['task_id']]))
        self.assertEqual(response.status_code,401)
        self.assertNotIn('spreadsheet',response.headers['content-type'])

    def test_live_table_shapes_match_public_openapi(self):
        import yaml,jsonschema
        from test_pipeline_schema import json_schema
        document=json_schema(yaml.safe_load((Path(__file__).resolve().parents[1]/'openapi.yaml').read_text(encoding='utf-8')))
        task,t=self.prepared();url=f'/api/v1/tasks/{task["task_id"]}'
        body=dict(run_id=task['run_id'],expected_revision=task['result_revision'],tables=[draft(t)])
        for name,payload in [('ReviewRequest',body),('TaskResult',self.client.get(url+'/result',headers=self.headers).json())]:
            jsonschema.Draft7Validator(dict(document,**{'$ref':'#/components/schemas/'+name})).validate(payload)
        saved=self.client.post(url+'/review',json=body,headers=self.headers)
        self.assertEqual(saved.status_code,200,saved.text)
        jsonschema.Draft7Validator(dict(document,**{'$ref':'#/components/schemas/TaskResult'})).validate(self.client.get(url+'/result',headers=self.headers).json())

    def prepared(self,override=None):
        uploaded=self.upload('TABLE');self.assertEqual(uploaded.status_code,202,uploaded.text)
        task=self.business.store.claim_next();t=table()
        t.update(override or {})
        page_id=self.business.store.save_page(task,1,self.image,16,12)
        self.business.store.publish(task,dict(doc_type='TABLE',status='NEED_REVIEW',fields=[],tables=[t],review_issues=[],
            pages=[dict(page_index=1,page_id=page_id,width=16,height=12,regions=[],extraction_method='ocr')]))
        return self.task(task['id']),t

    def test_historical_ranked_and_untitled_results_export_without_any_review_write(self):
        from urllib.parse import unquote
        for candidates,expected in [([dict(text='最先候选',box=[0,0,1,1]),dict(text='次选候选',box=[0,0,1,1])],'最先候选'),([],None)]:
            with self.subTest(candidates=candidates):
                task,_=self.prepared(dict(title='',title_candidates=candidates))
                url=f'/api/v1/tasks/{task["task_id"]}/result'
                before=self.client.get(url,headers=self.headers).json()
                self.assertTrue(before['export_ready']);self.assertEqual(before['export_blockers'],[])
                response=self.client.post('/api/v1/export',headers=self.headers,json=dict(task_ids=[task['task_id']]))
                self.assertEqual(response.status_code,200)
                self.assertIn((expected or 'synthetic-0')+'.xlsx',unquote(response.headers['content-disposition']))
                sheet=load_workbook(io.BytesIO(response.content)).active
                self.assertEqual(sheet['A1'].value,expected or '编号')
                self.assertEqual(self.client.get(url,headers=self.headers).json(),before)
                self.assertNotIn('reviewed_tables',before)

    def test_review_export_revision_and_anonymous_boundaries(self):
        task,t=self.prepared();d=draft(t)
        url=f'/api/v1/tasks/{task["task_id"]}/review'
        body=dict(run_id=task['run_id'],expected_revision=task['result_revision'],tables=[d])
        self.assertEqual(self.client.post(url,json=body).status_code,401)
        saved=self.client.post(url,json=body,headers=self.headers)
        self.assertEqual(saved.status_code,200,saved.text)
        self.assertTrue(saved.json()['export_ready'])
        self.assertEqual(self.client.post(url,json=body,headers=self.headers).status_code,409)
        response=self.client.post('/api/v1/export',json={'task_ids':[task['task_id']],'masked':False},headers=self.headers)
        self.assertEqual(response.status_code,200,response.text[:100] if response.status_code!=200 else '')
        self.assertIn("filename*=UTF-8''",response.headers['content-disposition'])
        self.assertEqual(load_workbook(io.BytesIO(response.content)).active['A3'].value,'001')

    def test_complete_unreviewed_table_exports_without_review_mutation(self):
        task,_=self.prepared()
        url=f'/api/v1/tasks/{task["task_id"]}/result'
        before=self.client.get(url,headers=self.headers).json()
        self.assertTrue(before['export_ready'])
        body=dict(task_ids=[task['task_id']],versions=[dict(task_id=task['task_id'],run_id=task['run_id'],result_revision=task['result_revision'])])
        response=self.client.post('/api/v1/export',json=body,headers=self.headers)
        self.assertEqual(response.status_code,200)
        self.assertEqual(load_workbook(io.BytesIO(response.content)).active['A3'].value,'001')
        self.assertEqual(self.client.get(url,headers=self.headers).json(),before)
        self.assertNotIn('reviewed_tables',before);self.assertEqual(before['status'],'NEED_REVIEW')
        body['versions'][0]['result_revision']+=1
        self.assertEqual(self.client.post('/api/v1/export',json=body,headers=self.headers).status_code,409)

    def test_batch_export_blocks_only_actual_incomplete_metadata(self):
        first,_=self.prepared();second,t=self.prepared()
        ids=[first['task_id'],second['task_id']]
        body=dict(task_ids=ids)
        response=self.client.post('/api/v1/export',json=body,headers=self.headers)
        self.assertEqual(response.status_code,200)
        with ZipFile(io.BytesIO(response.content)) as archive:self.assertEqual(len(archive.namelist()),2)
        d=dict(draft(t),title='')
        self.client.post(f'/api/v1/tasks/{second["task_id"]}/review',headers=self.headers,
            json=dict(run_id=second['run_id'],expected_revision=second['result_revision'],tables=[d])).raise_for_status()
        response=self.client.post('/api/v1/export',json=body,headers=self.headers)
        self.assertEqual(response.status_code,200)
        with ZipFile(io.BytesIO(response.content)) as archive:
            self.assertIn('synthetic-0.xlsx',archive.namelist())
            self.assertEqual(load_workbook(io.BytesIO(archive.read('synthetic-0.xlsx'))).active['A1'].value,'编号')
        second=self.task(second['task_id']);d['header_rows']=None
        self.client.post(f'/api/v1/tasks/{second["task_id"]}/review',headers=self.headers,
            json=dict(run_id=second['run_id'],expected_revision=second['result_revision'],tables=[d])).raise_for_status()
        response=self.client.post('/api/v1/export',json=body,headers=self.headers)
        self.assertEqual(response.status_code,409)
        self.assertIn('tables[0].header_rows',response.text)

    def test_table_excluded_from_packages_pdf_and_required_configuration(self):
        for data in [dict(file_types='["TABLE"]'),dict(doc_type_hint='TABLE',person_name='合成人员')]:
            response=self.client.post('/api/v1/batches',headers=self.headers,data=data,files=[('files',('synthetic.png',self.image,'image/png'))])
            self.assertEqual(response.status_code,422,response.text)
        import fitz
        pdf=fitz.open();pdf.new_page()
        self.assertEqual(self.upload('TABLE',data=pdf.tobytes()).status_code,415)
        from ocr_app.required_fields import configuration,save_configuration
        self.assertNotIn('TABLE',[item['doc_type'] for item in configuration(dict(revision=0,overrides={}))['items']])
        with self.assertRaises(TaskError):
            save_configuration(self.business.store,dict(doc_type='TABLE',required_fields=[],expected_revision=0),'owner',authorize=lambda c:None)


if __name__=='__main__':unittest.main()

