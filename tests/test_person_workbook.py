"""Complete person rows preserve selected document and nested record boundaries."""
import copy
import io
from pathlib import Path
import sys
import unittest
from uuid import uuid4

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
from ocr_app.doc_profiles import LABELS, MEMBER_PROFILE, profile
from ocr_app.extraction import extract_result
from ocr_app.exporting import workbook_bytes
from ocr_app.task_store import TaskError


def document(kind, name='合成姓名', filename=None, **values):
    result = extract_result([], kind)
    for field in result['fields']:
        field['value'] = values.get(field['key'], name if field['key'] == 'name' else None)
    result['_source_filename'] = filename or LABELS[kind] + '.png'
    result['task_id'] = str(uuid4())
    return result


def exported(*packages):
    results = [result for package in packages for result in package['results']]
    return load_workbook(io.BytesIO(workbook_bytes(results, packages)))


def rows(sheet):
    values = list(sheet.values)
    return [dict(zip(values[0], row)) for row in values[1:]]


class PersonWorkbookTests(unittest.TestCase):
    def test_all_profile_fields_including_blank_optional_and_medical_catalogue(self):
        selected = [document(kind, gender='F', school_name='修正学校',
                    certificate_no='00123456789012345678', exam_projects='内科；眼科',
                    age=0, major='=1+2') for kind in LABELS if kind != 'GENERIC']
        package = dict(person_name='人工姓名', results=selected)
        before = copy.deepcopy(package)
        book = exported(package)
        self.assertEqual(book.sheetnames[0], '人员信息总表')
        sheet = book.worksheets[0]
        row = rows(sheet)[0]
        self.assertEqual(row['人员姓名'], '人工姓名')
        self.assertEqual(row['材料份数'], len(selected))
        for result in selected:
            kind = result['doc_type']
            for field in profile(kind):
                self.assertIn(LABELS[kind] + '·' + field['label'], row)
            self.assertEqual(row[LABELS[kind] + '·来源文件'], result['_source_filename'])
        self.assertIsNone(row['毕业证·发证日期'])
        self.assertEqual(row['毕业证·学校名称'], '修正学校')
        self.assertEqual(row['毕业证·证书编号'], '00123456789012345678')
        self.assertEqual(row['毕业证·专业'], '=1+2')
        self.assertEqual(row['体检报告·体检项目'], '内科；眼科')
        self.assertEqual(row['体检报告·体检项目数量'], 2)
        self.assertEqual(row['体检报告·年龄'], 0)
        self.assertEqual(row['体检报告·性别'], '女')
        self.assertIn('体检报告·异常指标', row)
        self.assertEqual(book['审核记录台账'].max_column, 18)
        self.assertEqual(sheet.freeze_panes, 'C2')
        self.assertEqual(sheet['A1'].fill.fgColor.rgb[-6:], 'EF760E')
        for label in ['毕业证·证书编号', '毕业证·专业']:
            cell = sheet.cell(2, list(row).index(label) + 1)
            self.assertEqual(cell.data_type, 's')
            self.assertEqual(cell.number_format, '@')
        self.assertEqual(package, before)

    def test_repeated_documents_and_cross_package_union_do_not_merge_values(self):
        first = dict(person_name='甲', results=[
            document('DIPLOMA', '甲', '本科.png', school_name='甲校', major='专业一'),
            document('DIPLOMA', '甲', '硕士.png', school_name='乙校', major=None)])
        second = dict(person_name='乙', results=[
            document('DIPLOMA', '乙', '专科.png', school_name='丙校'),
            document('BANK_CARD', filename='银行.png', card_number='0001234567890123456')])
        book = exported(first, second)
        a, b = rows(book.worksheets[0])
        self.assertEqual(a['毕业证1·学校名称'], '甲校')
        self.assertEqual(a['毕业证1·专业'], '专业一')
        self.assertEqual(a['毕业证2·学校名称'], '乙校')
        self.assertIsNone(a['毕业证2·专业'])
        self.assertEqual(a['毕业证2·来源文件'], '硕士.png')
        self.assertEqual(b['毕业证1·学校名称'], '丙校')
        self.assertIsNone(b['毕业证2·来源文件'])
        self.assertIsNone(a['银行卡·卡号'])
        self.assertEqual(b['银行卡·卡号'], '0001234567890123456')
        self.assertEqual(book['毕业证'].max_row, 4)
        self.assertEqual(len(a), len(b))

    def test_member_fields_generic_long_text_and_duplicate_keys_are_complete(self):
        household = document('HOUSEHOLD_REGISTER')
        household['members'] = [dict(fields=[dict(key='name', value='成员一'),
            dict(key='gender', value='M')]), dict(name=dict(key='name', value='成员二'))]
        generic = document('GENERIC')
        generic['generic'] = dict(full_text='旧正文', key_values=[])
        body = '新正文\n' * 18000
        generic['reviewed_generic'] = dict(full_text=body, key_values=[
            dict(key='同名键', value='=1+2'), dict(key='同名键', value='000123')])
        row = rows(exported(dict(person_name=None, results=[household, generic])).worksheets[0])[0]
        for member in (1, 2):
            for field in MEMBER_PROFILE:
                self.assertIn(f'户口本·成员{member}·{field["label"]}', row)
        self.assertEqual(row['户口本·成员1·姓名'], '成员一')
        self.assertEqual(row['户口本·成员1·性别'], '男')
        self.assertEqual(row['户口本·成员2·姓名'], '成员二')
        self.assertEqual(''.join(value for key, value in row.items()
                         if key.startswith('通用识别·正文')), body)
        self.assertEqual(row['通用识别·键值1·字段'], '同名键')
        self.assertEqual(row['通用识别·键值1·内容'], '=1+2')
        self.assertEqual(row['通用识别·键值2·内容'], '000123')

    def test_absence_reasons_render_and_empty_review_does_not_restore_old_generic(self):
        diploma = document('DIPLOMA')
        for field in diploma['fields']:
            if field['key'] == 'issue_date': field['absence_reason'] = 'NOT_ON_DOCUMENT'
            if field['key'] == 'duration': field['absence_reason'] = 'NOT_APPLICABLE'
        generic = document('GENERIC')
        generic['generic'] = dict(full_text='不得恢复', key_values=[dict(key='旧', value='旧')])
        generic['reviewed_generic'] = dict(full_text='', key_values=[])
        row = rows(exported(dict(person_name=None, results=[diploma, generic])).worksheets[0])[0]
        self.assertEqual(row['毕业证·发证日期'], '未载明')
        self.assertEqual(row['毕业证·学制'], '不适用')
        self.assertIsNone(row['通用识别·正文1'])
        self.assertNotIn('不得恢复', str(row))
        self.assertNotIn('通用识别·键值1·字段', row)

    def test_long_structured_value_and_excessive_column_count_fail_loudly(self):
        with self.assertRaises(TaskError) as caught:
            exported(dict(person_name=None, results=[document('DIPLOMA', major='x' * 32768)]))
        self.assertEqual(caught.exception.code, 'EXPORT_VALUE_TOO_LONG')
        documents = [document('BUSINESS_LICENSE', filename=str(i)) for i in range(1100)]
        with self.assertRaises(TaskError) as caught:
            exported(dict(person_name=None, results=documents))
        self.assertEqual(caught.exception.code, 'EXPORT_VALUE_TOO_LONG')


if __name__ == '__main__': unittest.main()
