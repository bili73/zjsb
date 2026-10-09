"""Business content, saved corrections, association and spreadsheet safety."""
import copy
import io
from pathlib import Path
import sys
import unittest
from uuid import uuid4
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
from ocr_app.extraction import extract_result, _medical_item
from ocr_app.exporting import workbook_bytes
from ocr_app.task_store import TaskError


def records(sheet):
    rows=list(sheet.values)
    return [dict(zip(rows[0], row)) for row in rows[1:]]


class BusinessWorkbookTests(unittest.TestCase):
    def book(self, *results):
        return load_workbook(io.BytesIO(workbook_bytes(results)))

    def test_saved_certificate_fields_and_text_identifiers_have_business_labels(self):
        r=extract_result([], 'DIPLOMA')
        values=dict(name='修正姓名', school_name='修正大学', major='=SUM(A1:A2)',
                    enroll_date='2020-09', graduate_date='2024-06', issue_date='2024-07-01',
                    certificate_no='00123456789012345678')
        for f in r['fields']: f['value']=values.get(f['key'])
        r['_source_filename']='毕业证.jpg'
        r['fields'].extend([dict(key='principal',value='removed'),dict(key='study_mode',value='removed')])
        r.update(task_id=str(uuid4()),run_id=str(uuid4()),result_revision=5,status='NEED_REVIEW')
        before=copy.deepcopy(r)
        book=self.book(r)
        self.assertEqual(book.sheetnames,['毕业证'])
        row=records(book['毕业证'])[0]
        self.assertEqual(row['来源文件'],'毕业证.jpg')
        self.assertEqual(row['姓名'],'修正姓名')
        self.assertEqual(row['学校名称'],'修正大学')
        self.assertEqual(row['毕业日期'],'2024-06')
        self.assertEqual(row['证书编号'],values['certificate_no'])
        self.assertFalse({'task_id','status','校长','学习形式'} & row.keys())
        for key in ['专业','证书编号']:
            col=list(row).index(key)+1
            self.assertEqual(book['毕业证'].cell(2,col).data_type,'s')
        self.assertEqual(r,before)

    def test_medical_only_abnormal_rows_with_linked_diagnosis_and_unlinked_evidence(self):
        r=extract_result([], 'MEDICAL_REPORT')
        for f in r['fields']: f['value']={'name':'合成人员','gender':'M','exam_date':'2026-10-01','exam_projects':'内科；眼科'}.get(f['key'])
        items=[_medical_item(name,value,'u',ref,None,None,dict(page=2,bbox=[.1,y,.8,y+.04]),[name,value],risk_grade=grade)
               for name,value,ref,grade,y in [('身高','180','150-175','C',.1),('体重','80','40-70','C',.2),
                                            ('小结','异常','/','C',.3),('体重指数','26','18-24','B',.4),
                                            ('潜血','+3','/','E3',.5),('正常项目','4','0-9','A',.6)]]
        r['items']=items
        r['report_summary']['findings']=[dict(finding_id=str(uuid4()),text='潜血：异常',project_name='潜血',
            result_item_id=items[4]['item_id'],result_sources=[items[4]['source']],result_raw='+3',risk_grade='E3',diagnosis='潜血异常'),
            dict(finding_id=str(uuid4()),text='蛋白质：异常',project_name='蛋白质',result_sources=[],result_raw=None,risk_grade='E3',diagnosis='蛋白质异常')]
        book=self.book(r)
        self.assertEqual(book.sheetnames,['审核记录台账'])
        row=records(book.active)[0]
        self.assertEqual(row['体检项目数量'],2)
        self.assertNotIn('人工录取意见',row)
        self.assertIn('体重指数 26',row['异常指标'])
        self.assertIn('潜血 +3',row['异常指标'])
        self.assertIn('潜血异常',row['异常指标'])
        self.assertIn('蛋白质 未关联检测结果',row['异常指标'])
        self.assertNotIn('身高',row['异常指标'])
        self.assertEqual(len(row['异常指标'].splitlines()),3)
        self.assertNotIn('、',row['异常指标'])
        self.assertTrue(book.active['R2'].alignment.wrap_text)
        self.assertGreaterEqual(book.active.row_dimensions[2].height,59)
        self.assertEqual(book.active['A1'].fill.fgColor.rgb[-6:],'4472C4')
        self.assertEqual(book.active['K1'].fill.fgColor.rgb[-6:],'FF6C0D')
        # Legacy source-backed links must not produce a second copy.
        r['report_summary']['findings'][0].pop('result_item_id')
        self.assertEqual(records(self.book(r).active)[0]['异常指标'].count('潜血 +3'),1)

    def test_household_members_generic_chunks_and_layout(self):
        household=extract_result([], 'HOUSEHOLD_REGISTER')
        household['members']=[dict(name=dict(key='name',value='成员'),id_number=dict(key='id_number',value='001234567890123456'))]
        generic=dict(doc_type='GENERIC',generic=dict(full_text='旧文本',key_values=[]),
                     reviewed_generic=dict(full_text='文本'*18000,key_values=[dict(key='键',value='=1+2')]))
        book=self.book(household,generic)
        self.assertEqual(''.join(row[-1] for row in list(book['通用识别文本'].values)[1:]),generic['reviewed_generic']['full_text'])
        self.assertEqual(records(book['户口成员'])[0]['公民身份号码'],'001234567890123456')
        for sheet in book:
            self.assertEqual(sheet.freeze_panes,'A2')
            self.assertTrue(sheet.auto_filter.ref)
            self.assertEqual(sheet['A1'].font.bold,True)
            self.assertTrue(all(c.fill.fgColor.rgb[-6:]=='EF760E' for c in sheet[1]))
            self.assertTrue(sheet.column_dimensions['A'].width>=8)
            self.assertEqual(sheet.print_options.horizontalCentered,True)
        with self.assertRaises(TaskError):
            self.book(dict(doc_type='DIPLOMA',fields=[dict(key='major',value='a'*32768)]))

if __name__=='__main__': unittest.main()
