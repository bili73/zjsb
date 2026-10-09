"""Synthetic review/export safety checks; no inference or production data."""
import copy
import io
from pathlib import Path
import sys
import unittest
from uuid import uuid4

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
from ocr_app.extraction import extract_result
from ocr_app.review import apply_review, blockers
from ocr_app.task_store import TaskError
from ocr_app.exporting import workbook_bytes


class ReviewTests(unittest.TestCase):
    def result(self):
        return extract_result([], 'DIPLOMA')

    def test_empty_review_does_not_approve_missing(self):
        result = apply_review(self.result(), {}, actor='synthetic')
        self.assertTrue(blockers(result))
        self.assertEqual(result['status'], 'NEED_REVIEW')

    def test_unknown_and_masked_values_rejected(self):
        for patch in ({'unknown': 'x'}, {'certificate_no': '****1234'}):
            with self.assertRaises(TaskError):
                apply_review(self.result(), {'fields': patch}, actor='synthetic')

    def test_required_absence_cannot_approve(self):
        with self.assertRaises(TaskError):
            apply_review(self.result(), {'absence_reasons': {'name': 'NOT_ON_DOCUMENT'}}, actor='synthetic')

    def test_unacceptable_issue_rejected(self):
        result = self.result()
        identifier = result['fields'][0]['review_issues'][0]['issue_id']
        with self.assertRaises(TaskError):
            apply_review(result, {'accepted_issues': [{'issue_id': identifier, 'reason': 'checked'}]}, actor='synthetic')

    def test_export_strings_are_not_formulas_or_numbers(self):
        from openpyxl import load_workbook
        result = dict(task_id=str(uuid4()), run_id=str(uuid4()), result_revision=1, status='COMPLETED',
                      doc_type='DIPLOMA', fields=[dict(key='certificate_no', value='0012345678901234567'),
                                               dict(key='name', value='=1+1')])
        book = load_workbook(io.BytesIO(workbook_bytes([result])))
        cells = [cell for row in book['毕业证'] for cell in row]
        self.assertTrue(any(cell.value == '0012345678901234567' and cell.data_type == 's' for cell in cells))
        self.assertFalse(any(cell.data_type == 'f' for cell in cells))

    def test_format_exception_stays_accepted_but_cross_field_conflict_cannot(self):
        result = apply_review(self.result(), {'fields': {'graduate_date': '2020-99'}}, actor='human')
        field = next(f for f in result['fields'] if f['key'] == 'graduate_date')
        problem = next(p for p in field['review_issues'] if p['code'] == 'VALIDATION_FAILED')
        result = apply_review(result, {'accepted_issues': [{'issue_id': problem['issue_id'], 'reason': '原件确实如此，保留例外'}]}, actor='human')
        again = apply_review(result, {}, actor='human')
        self.assertFalse(next(f for f in again['fields'] if f['key'] == 'graduate_date')['review_issues'])
        identity = extract_result([], 'ID_CARD_FRONT')
        identity = apply_review(identity, {'fields': {'id_number': '11010519491231002X', 'birth_date': '2000-01-01', 'gender': 'F'}}, actor='human')
        target = next(f for f in identity['fields'] if f['key'] == 'id_number')
        conflict = next(p for p in target['review_issues'] if p['code'] == 'CONFLICT')
        self.assertFalse(conflict['can_accept'])
        with self.assertRaises(TaskError):
            apply_review(identity, {'accepted_issues': [{'issue_id': conflict['issue_id'], 'reason': 'ignore'}]}, actor='human')
        corrected = apply_review(identity, {'fields': {'birth_date': '1949-12-31'}}, actor='human')
        self.assertFalse(any(p['code'] == 'CONFLICT' for p in next(f for f in corrected['fields'] if f['key'] == 'id_number')['review_issues']))

    def test_medical_edit_recalculates_and_preserves_original_source(self):
        from ocr_app.extraction import _medical_item
        original = extract_result([], 'MEDICAL_REPORT')
        item = _medical_item('测试指标', '3', 'u', '1-5', None, None, {'page': 1}, ['original row'])
        original['items'] = [item]
        updated = apply_review(original, {'items': [{'item_id': item['item_id'], 'result_raw': '9↑',
                                                      'risk_grade': None, 'diagnosis': None,
                                                      'absence_reasons': {'risk_grade': 'NOT_ON_DOCUMENT', 'diagnosis': 'NOT_ON_DOCUMENT'}}]}, actor='human')
        actual = updated['items'][0]
        self.assertEqual(actual['value_num'], 9)
        self.assertEqual(actual['abnormal_flag'], 'HIGH')
        self.assertFalse(actual['review_issues'])
        self.assertEqual(actual['source'], item['source'])
        self.assertEqual(original['items'][0]['result_raw'], '3')
        self.assertIsNone(updated['stats']['total'])

    def test_medical_cannot_edit_derived_confidence_or_duplicate_ids(self):
        from ocr_app.extraction import _medical_item
        result = extract_result([], 'MEDICAL_REPORT')
        result['items'] = [_medical_item('指标', '3', None, None, None, None, {}, [])]
        identifier = result['items'][0]['item_id']
        for body in ({'items': [{'item_id': identifier, 'confidence': 1}]},
                     {'items': [{'item_id': identifier}], 'removed_item_ids': [identifier]},
                     {'removed_item_ids': [str(uuid4())]}):
            with self.assertRaises(TaskError):
                apply_review(result, body, actor='human')

    def test_generic_long_text_export_keeps_all_characters(self):
        from openpyxl import load_workbook
        original = extract_result([], 'GENERIC')
        text = '测试文本' * 20000
        result = apply_review(original, {'generic': {'full_text': text, 'key_values': []}}, actor='human')
        self.assertEqual(result['generic']['full_text'], original['generic']['full_text'])
        result.update(task_id=str(uuid4()), run_id=str(uuid4()), result_revision=1)
        book = load_workbook(io.BytesIO(workbook_bytes([result])))
        exported = ''.join(row[-1].value for row in list(book['通用识别文本'])[1:])
        self.assertEqual(exported, text)

    def test_optional_absence_and_member_required_rules(self):
        result = apply_review(self.result(), {'absence_reasons': {'gender': 'NOT_ON_DOCUMENT'}}, actor='human')
        gender = next(f for f in result['fields'] if f['key'] == 'gender')
        self.assertFalse(gender['need_review'])
        self.assertEqual(gender['absence_reason'], 'NOT_ON_DOCUMENT')
        household = extract_result([], 'HOUSEHOLD_REGISTER')
        household['members'] = [dict(name=dict(key='name', label='姓名', value=None, review_issues=[], confidence=None, confidence_method='unavailable'))]
        household = apply_review(household, {}, actor='human')
        self.assertTrue(household['members'][0]['name']['need_review'])
        with self.assertRaises(TaskError):
            apply_review(household, {'members': [{'index': 0, 'absence_reasons': {'name': 'NOT_APPLICABLE'}}]}, actor='human')

    def test_xuexin_wrong_kind_cannot_be_confirmed(self):
        result = extract_result([], 'XUEXIN_DIPLOMA_TABLE')
        result['report_kind'] = 'OTHER_REPORT'
        result = apply_review(result, {}, actor='human')
        self.assertTrue(any(p['field'] == 'task' and p['code'] == 'UNKNOWN' for p in blockers(result)))

    def test_accept_and_patch_same_original_issue_is_valid(self):
        original = extract_result([dict(text='姓名：测试姓名', page=1)], 'DIPLOMA')
        name = next(f for f in original['fields'] if f['key'] == 'name')
        problem = name['review_issues'][0]
        updated = apply_review(original, {'fields': {'name': '测试新姓名'}, 'accepted_issues': [dict(issue_id=problem['issue_id'], reason='核对并更正')]}, actor='human')
        self.assertFalse(next(f for f in updated['fields'] if f['key'] == 'name')['review_issues'])

    def test_finding_repair_recomputes_count_and_links(self):
        from ocr_app.extraction import _medical_item, issue
        result = extract_result([], 'MEDICAL_REPORT')
        result['pages'] = [{'page_index': 1}]
        item = _medical_item('测试指标', '7↑', 'u', '1-5', None, None, {'page': 1}, ['original'])
        result['items'] = [item]
        result['report_summary'] = dict(scope='follow_up_required', declared_count=1, count_check='mismatch', findings=[], sources=[{'page': 1, 'raw_text': '原标题'}])
        result['review_issues'].append(issue('CONFLICT', '报告声明数量与该标题下清单不一致，请核对原页。'))
        fixed = apply_review(result, {'report_summary': {'source_page': 1, 'reason': '补录原页遗漏条目',
                                'added_findings': [{'text': '测试指标异常', 'risk_grade': 'B', 'diagnosis': '指标偏高', 'result_item_id': item['item_id']}]}}, actor='human')
        self.assertEqual(fixed['report_summary']['count_check'], 'match')
        finding = fixed['report_summary']['findings'][0]
        self.assertEqual(finding['result_raw'], '7↑')
        self.assertEqual(finding['result_sources'], [item['source']])
        self.assertEqual(fixed['report_summary']['sources'], result['report_summary']['sources'])
        self.assertFalse(any(p['code'] == 'CONFLICT' for p in fixed['review_issues']))

    def test_summary_count_requires_explicit_followup_title(self):
        original = extract_result([], 'MEDICAL_REPORT')
        original['pages'] = [{'page_index': 1}]
        for summary in ({'source_page': 1, 'reason': '核对', 'declared_count': 0},
                        {'source_page': 1, 'reason': '核对', 'declared_count': 0, 'scope': 'follow_up_required'}):
            with self.assertRaises(TaskError):
                apply_review(original, {'report_summary': summary}, actor='human')
        updated = apply_review(original, {'report_summary': {'source_page': 1, 'reason': '核对该标题',
                               'declared_count': 0, 'scope': 'follow_up_required', 'title': '异常复检项目共零项'}}, actor='human')
        self.assertEqual(updated['report_summary']['count_check'], 'match')
        self.assertEqual(updated['report_summary']['recognition_scope'], 'unidentified')
        updated = apply_review(updated, {'report_summary': {'source_page': 1, 'reason': '更正为总体结论',
                               'scope': 'overall_conclusion'}}, actor='human')
        self.assertIsNone(updated['report_summary']['declared_count'])
        self.assertEqual(updated['report_summary']['count_check'], 'unverified')

    def test_household_cover_alone_never_completes(self):
        original = extract_result([], 'HOUSEHOLD_REGISTER')
        original['pages'] = [{'page_index': 1}]
        top = dict(household_type='家庭户', householder='测试人员', household_no='00001', address='测试地址')
        result = apply_review(original, {'fields': top}, actor='human')
        self.assertTrue(any(b['code'] == 'MISSING' and b['field'] == 'task' for b in blockers(result)))
        self.assertTrue(any(b['code'] == 'INCOMPLETE_OUTPUT' for b in blockers(result)))
        result = apply_review(result, {'added_members': [{'fields': {'name': '测试成员'}, 'source_page': 1}],
                              'coverage': {'pages': [1], 'reason': '已核对扉页和成员'}}, actor='human', coverage_verified=True)
        self.assertEqual(result['status'], 'COMPLETED')
        self.assertEqual(len(result['members']), 1)
        self.assertTrue(result['members'][0]['member_id'])
        removed = apply_review(result, {'removed_member_indexes': [0]}, actor='human')
        self.assertEqual(removed['status'], 'NEED_REVIEW')

    def test_chinese_gender_is_normalized_before_identity_cross_checks(self):
        result = extract_result([], 'ID_CARD_FRONT')
        result = apply_review(result, {'fields': {'id_number': '11010519491231002X', 'birth_date': '1949-12-31', 'gender': '女'}}, actor='human')
        gender = next(f for f in result['fields'] if f['key'] == 'gender')
        identity = next(f for f in result['fields'] if f['key'] == 'id_number')
        self.assertEqual(gender['value'], 'F')
        self.assertFalse(any(p['code'] == 'CONFLICT' for p in identity['review_issues']))
        result = apply_review(result, {'fields': {'gender': '男'}}, actor='human')
        self.assertTrue(any(p['code'] == 'CONFLICT' for p in next(f for f in result['fields'] if f['key'] == 'id_number')['review_issues']))

    def test_accept_without_reason_never_accepts_conflict_or_incomplete(self):
        from ocr_app.extraction import issue
        for code in ('CONFLICT', 'INCOMPLETE_OUTPUT'):
            result = self.result()
            problem = issue(code, 'synthetic issue')
            result['review_issues'].append(problem)
            with self.assertRaises(TaskError):
                apply_review(result, {'accepted_issues': [{'issue_id': problem['issue_id']}]}, actor='human')


if __name__ == '__main__':
    unittest.main()
