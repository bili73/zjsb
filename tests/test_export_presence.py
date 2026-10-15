"""Presence checks with valid synthetic formats, independent of review status."""
import copy
import io
import unittest
from unittest.mock import patch
import test_business_api as business_tests
from sqlalchemy import text
from ocr_app.doc_profiles import profile
from ocr_app.extraction import extract_result
from ocr_app.exporting import required_export_blockers, workbook_bytes


def complete(kind):
    return dict(doc_type=kind, status='NEED_REVIEW', fields=[dict(key=f['key'], value='2020-06-01' if f['type']=='date' else 'synthetic')
                for f in profile(kind) if f['required']], review_issues=[dict(code='INCOMPLETE_OUTPUT')])


class PresenceTests(unittest.TestCase):
    def test_only_missing_values_block_and_input_is_not_mutated(self):
        result = complete('DIPLOMA')
        before = copy.deepcopy(result)
        self.assertEqual(required_export_blockers(result), [])
        self.assertEqual(result, before)
        for value in [None, '', ' \t\n']:
            result['fields'][0]['value'] = value
            self.assertEqual(required_export_blockers(result), [{'field': 'name', 'code': 'MISSING'}])
        for value in [0, False, 'invalid-format-but-present']:
            result['fields'][0]['value'] = value
            self.assertEqual(required_export_blockers(result), [])
        self.assertEqual(required_export_blockers(None), [{'field': 'task', 'code': 'MISSING'}])

    def test_conditional_profile_and_household_members(self):
        result = complete('BUSINESS_LICENSE')
        values = {f['key']: f for f in result['fields']}
        values['company_type']['value'] = '个体工商户'
        values['registered_capital']['value'] = None
        values['unified_social_credit_code']['value'] = None
        result['fields'].append(dict(key='registration_no', value='legacy'))
        self.assertEqual(required_export_blockers(result), [])
        result = complete('HOUSEHOLD_REGISTER')
        for f in result['fields']:
            if f['key'] == 'household_type': f['value'] = '集体户口'
            if f['key'] == 'householder': f['value'] = None
        self.assertEqual(required_export_blockers(result), [{'field': 'members[0].name', 'code': 'MISSING'}])
        result['members'] = [dict(name=dict(key='name', value='synthetic')), dict(fields=[dict(key='name', value=' ')])]
        self.assertEqual(required_export_blockers(result), [{'field': 'members[1].name', 'code': 'MISSING'}])

    def test_basic_info_fallback_is_exported_but_explicit_null_field_wins(self):
        result = complete('MEDICAL_REPORT')
        result['basic_info'] = {f['key']: f for f in result.pop('fields')}
        self.assertEqual(required_export_blockers(result), [])
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(workbook_bytes([result])))
        self.assertIn('synthetic', [c.value for row in book['审核记录台账'] for c in row])
        result['fields'] = [dict(key='name', value=None)]
        self.assertEqual(required_export_blockers(result), [{'field': 'name', 'code': 'MISSING'}])

    def test_recognized_generic_exports_all_text_but_empty_review_does_not_fall_back(self):
        from openpyxl import load_workbook
        body = 'synthetic=' * 9000
        result = dict(doc_type='GENERIC', generic=dict(full_text=body, key_values=[]))
        self.assertEqual(required_export_blockers(result), [])
        book = load_workbook(io.BytesIO(workbook_bytes([result])))
        self.assertEqual(''.join(row[-1].value for row in list(book['通用识别文本'])[1:]), body)
        result['reviewed_generic'] = dict(full_text=' ', key_values=[])
        self.assertEqual(required_export_blockers(result), [{'field': 'generic.full_text', 'code': 'MISSING'}])


class PresenceHTTPTests(unittest.TestCase):
    setUp = business_tests.BusinessAPITests.setUp
    upload = business_tests.BusinessAPITests.upload

    def saved(self, kind, status='NEED_REVIEW'):
        created = self.upload(kind).json()
        worker = self.business.store.claim_next()
        result = extract_result([], kind)
        result.update(complete(kind), status=status)
        self.business.store.publish(worker, result)
        return created['task_ids'][0]

    def test_failed_and_unreviewed_export_without_review_revalidation_or_state_changes(self):
        ids = [self.saved('DIPLOMA', status) for status in ['FAILED', 'NEED_REVIEW']]
        with self.business.engine.connect() as c:
            original = c.execute(text('SELECT id,status,result_enc,revision FROM tasks ORDER BY id')).all()
        for identifier in ids:
            task = self.client.get('/api/v1/tasks/' + identifier, headers=self.headers).json()
            self.assertTrue(task['export_ready'])
            result = self.client.get('/api/v1/tasks/' + identifier + '/result', headers=self.headers).json()
            self.assertTrue(result['export_ready'])
        with patch('ocr_app.review.apply_review', side_effect=AssertionError('export must not revalidate review')):
            response = self.client.post('/api/v1/export', headers=self.headers, json={'task_ids': ids})
        self.assertEqual(response.status_code, 200)
        with self.business.engine.connect() as c:
            self.assertEqual(c.execute(text('SELECT id,status,result_enc,revision FROM tasks ORDER BY id')).all(), original)
        missing = self.upload().json()['task_ids'][0]
        task = self.client.get('/api/v1/tasks/' + missing, headers=self.headers).json()
        self.assertFalse(task['export_ready'])
        response = self.client.post('/api/v1/export', headers=self.headers, json={'task_ids': ids + [missing]})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()['details']['tasks'][0]['blockers'], [{'field': 'task', 'code': 'MISSING'}])
        self.assertNotIn('synthetic', response.text)

    def test_medical_filter_does_not_change_eligibility(self):
        identifier = self.saved('MEDICAL_REPORT')
        for selected in ['all', 'abnormal', 'unknown']:
            response = self.client.get('/api/v1/tasks/' + identifier + '/result', headers=self.headers,
                                       params={'items_filter': selected})
            self.assertEqual(response.status_code, 200)
            self.assertTrue(response.json()['export_ready'])
            self.assertEqual(response.json()['export_blockers'], [])

    def test_completed_missing_value_blocks_and_recognized_generic_can_export(self):
        identifier = self.saved('DIPLOMA', 'COMPLETED')
        store = self.business.store
        with store.transaction() as c:
            row = store._row(c, identifier)
            result = self.business.storage.unpack(row['result_enc'])
            result['fields'][0]['value'] = '  '
            c.execute(text('UPDATE tasks SET result_enc=:result WHERE id=:id'),
                      dict(id=identifier, result=self.business.storage.pack(result)))
        response = self.client.post('/api/v1/export', headers=self.headers, json={'task_ids': [identifier]})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()['details']['tasks'][0]['blockers'], [{'field': 'name', 'code': 'MISSING'}])
        created = self.upload('GENERIC').json()
        worker = store.claim_next()
        raw = extract_result([], 'GENERIC')
        raw['generic']['full_text'] = '=synthetic recognized text'
        raw['status'] = 'NEED_REVIEW'
        store.publish(worker, raw)
        response = self.client.post('/api/v1/export', headers=self.headers, json={'task_ids': created['task_ids']})
        self.assertEqual(response.status_code, 200)
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(response.content))
        cell = list(book['通用识别文本'])[1][-1]
        self.assertEqual(cell.value, raw['generic']['full_text'])
        self.assertEqual(cell.data_type, 's')
