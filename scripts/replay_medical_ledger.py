"""Serial native-PDF ledger replay; original samples and production results stay intact."""
import argparse
import contextlib
import hashlib
import io
import json
from pathlib import Path
import sys
import time

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'backend'))
from ocr_app.documents import decode_pages
from ocr_app.extraction import extract_result, confirmed_abnormal
from ocr_app.exporting import workbook_bytes
from ocr_app.medical_ledger import HEADERS, GROUPS


def replay(directory, out):
    out.mkdir(parents=True,exist_ok=False)
    results=[]
    summary=dict(mode='serial_native_pdf_mapping',model_calls=0,production_writes=0,reports=[])
    for path in sorted(directory.glob('*.pdf')):
        data=path.read_bytes()
        checksum=hashlib.sha256(data).hexdigest()
        evidence=[]; pending=0; pages=0
        begin=time.monotonic()
        with contextlib.redirect_stdout(io.StringIO()):
            for page in decode_pages(data):
                evidence.extend(page.evidence)
                pending+=len(page.ocr_boxes)
                pages+=1
                page.image.close()
        native_seconds=time.monotonic()-begin
        begin=time.monotonic()
        r=extract_result(evidence,'MEDICAL_REPORT',complete=False)
        mapping_seconds=time.monotonic()-begin
        results.append(r)
        fields={f['key']:f for f in r['fields']}
        assert len(fields)==17 and r['stats']['total'] is None
        assert not any(i['name']=='小结' for i in r['items'])
        assert hashlib.sha256(path.read_bytes()).hexdigest()==checksum
        (out/(path.stem+'.json')).write_text(json.dumps(r,ensure_ascii=False,indent=2),encoding='utf-8')
        summary['reports'].append(dict(file_name=path.name,sha256=checksum,pages=pages,pending_ocr_boxes=pending,
            native_seconds=native_seconds,mapping_seconds=mapping_seconds,
            fields={key:dict(value=f['value'],source_page=(f.get('source') or {}).get('page'),
                            issues=[i['code'] for i in f['review_issues']]) for key,f in fields.items()},
            measured_items=len(r['items']),confirmed_abnormal=sum(confirmed_abnormal(i) for i in r['items']),
            category_counts={label:sum(i.get('category')==label for i in r['items']) for label,_ in GROUPS.values()},
            declared_count=r['report_summary']['declared_count'],findings=len(r['report_summary']['findings']),
            linked_findings=sum(bool(f.get('result_item_id')) for f in r['report_summary']['findings'])))
    target=out/'入职体检审核记录台账.xlsx'
    target.write_bytes(workbook_bytes(results))
    from openpyxl import load_workbook
    book=load_workbook(target)
    assert book.sheetnames==['审核记录台账']
    assert list(next(book.active.values))==HEADERS
    assert book.active.max_row==len(results)+1
    assert not any(c.data_type=='f' for row in book.active for c in row)
    summary['workbook']=target.name
    (out/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({**summary,'reports':[dict(file_name=r['file_name'],project_count=r['fields']['exam_project_count']['value'],
        abnormal_count=r['confirmed_abnormal'],groups=r['category_counts'],pending_ocr_boxes=r['pending_ocr_boxes']) for r in summary['reports']]},ensure_ascii=False))


if __name__=='__main__':
    p=argparse.ArgumentParser()
    p.add_argument('--input',type=Path,required=True)
    p.add_argument('--out',type=Path,required=True)
    a=p.parse_args()
    replay(a.input,a.out)
