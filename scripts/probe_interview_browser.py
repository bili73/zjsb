"""Synthetic real-HTTP/Edge check of the interview entry; no recognition model."""
import argparse
import io
import json
from pathlib import Path
import secrets
import socket
import sys
import threading
import time
from uuid import uuid4

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT/'backend'), str(ROOT/'tests')]

BROWSER = r'''
const fs=require('node:fs'),assert=require('node:assert/strict');
const {chromium}=require(process.env.PLAYWRIGHT_MODULE_PATH || 'C:/Users/JBS/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/node_modules/playwright');
(async()=>{
 const a=JSON.parse(fs.readFileSync(0,'utf8')), report={status:'failed',checks:[]};
 let browser,page,stage='launch';
 try {
  browser=await chromium.launch({channel:'msedge',headless:true});
  const context=await browser.newContext({viewport:{width:1500,height:1100},acceptDownloads:true});
  page=await context.newPage();page.setDefaultTimeout(20000);const errors=[];
  page.on('pageerror',e=>errors.push(e.message));
  const login=async(p,name)=>{await p.goto(a.url);await p.getByPlaceholder('账号',{exact:true}).fill(name);await p.getByPlaceholder('密码',{exact:true}).fill(a.password);await p.getByRole('button',{name:'进入工作台',exact:true}).click();await p.getByRole('button',{name:'自定义入口',exact:true}).waitFor();};
  assert.equal((await context.request.get(a.url+'/api/v1/doc-types')).status(),401);
  await login(page,'interview_qa');stage='upload';
  await page.locator('.type-more summary').click();
  await page.locator('.type-more .type-card').filter({hasText:'面试登记表'}).click();
  await page.locator('input[type=file]').setInputFiles(a.input);
  const upload=page.waitForResponse(r=>r.url().endsWith('/api/v1/batches')&&r.request().method()==='POST');
  await page.getByRole('button',{name:/开始识别/}).click();
  const response=await upload;assert.equal(response.status(),202);const batch=await response.json();
  const path='/api/v1/batches/'+batch.batch_id+'/tasks';let task;
  for(let i=0;i<80;i++){const r=await context.request.get(a.url+path);const rows=await r.json();task=rows[0];if(task?.status==='NEED_REVIEW')break;await page.waitForTimeout(200);}
  assert.equal(task.doc_type,'INTERVIEW_FORM');assert.equal(task.export_ready,false);
  report.checks.push('upload_recognition_and_missing_required_gate');
  stage='open';await page.getByRole('tab',{name:'任务批次'}).click();
  const batchPanel=page.locator('.batch-collapse .el-collapse-item').filter({hasText:'面试登记表'}).first();
  const header=batchPanel.locator('.el-collapse-item__header');
  if(await header.getAttribute('aria-expanded')!=='true')await header.click();
  await batchPanel.getByRole('button',{name:/^(查看|补填|核对)$/}).first().click();
  const editor=page.locator('.review-editor');await editor.waitFor();
  const field=label=>editor.locator('.el-form-item').filter({has:page.locator('.el-form-item__label').filter({hasText:new RegExp('^'+label+'（必填）$')})});
  assert.equal(await editor.locator('.field-grid > .el-form-item').count(),5);
  stage='candidates';
  await field('姓名').getByRole('button',{name:/采用/}).click();
  await field('身份证号').getByRole('button',{name:/采用/}).click();
  await field('手机号').locator('input').fill('13800138000');
  await field('身份证号').getByRole('button',{name:'查看来源',exact:true}).click();
  await page.locator('.source-preview figure img').waitFor();
  const overlay=page.locator('.source-preview .bbox-overlay');await overlay.waitFor();
  const box=await overlay.boundingBox(),picture=await page.locator('.source-preview figure img').first().boundingBox();
  assert.ok(Math.abs(box.width/picture.width-.9)<.02 && Math.abs(box.height/picture.height-.65)<.02);
  await page.screenshot({path:a.out+'/source-proof.png'});
  report.checks.push('five_required_fields_candidate_adoption_and_source_preview');
  stage='id_length';
  const idInput=field('身份证号').locator('input'),lengthHint=field('身份证号').locator('.number-length-hint');
  const originalResult=await (await context.request.get(a.url+'/api/v1/tasks/'+task.task_id+'/result')).json();
  const originalEvidence=originalResult.generic.full_text;
  await idInput.fill('11010519491231002');await lengthHint.waitFor();assert.match(await lengthHint.innerText(),/17位，比18位少1位/);
  await idInput.fill('110105194912310002X');assert.match(await lengthHint.innerText(),/19位，比18位多1位/);
  assert.equal(await idInput.inputValue(),'110105194912310002X');
  const beforeSave=await (await context.request.get(a.url+'/api/v1/tasks/'+task.task_id+'/result')).json();
  assert.equal(beforeSave.fields.find(f=>f.key==='id_number').value,null);
  const kept=page.waitForResponse(r=>r.url().endsWith('/review')&&r.request().method()==='POST');
  await editor.getByRole('button',{name:'保存修正并重新校验',exact:true}).click();assert.equal((await kept).status(),200);
  const resultHint=page.locator('.result-view .number-length-hint');await resultHint.waitFor();
  assert.match(await resultHint.innerText(),/19位，比18位多1位/);
  const preserved=await (await context.request.get(a.url+'/api/v1/tasks/'+task.task_id+'/result')).json();
  assert.equal(preserved.fields.find(f=>f.key==='id_number').value,'110105194912310002X');
  assert.equal(preserved.generic.full_text,originalEvidence);
  assert.equal(preserved.export_ready,true);assert.deepEqual(preserved.export_disputes,[]);
  assert.equal(await page.getByRole('button',{name:'确认与原件一致',exact:true}).count(),0);
  // A later patch without the save acknowledgment must invalidate the old one.
  const headers={'X-CSRF-Token':(await (await context.request.get(a.url+'/api/v1/auth/me')).json()).csrf_token,'Origin':a.url};
  const unreviewed=await context.request.post(a.url+'/api/v1/tasks/'+task.task_id+'/review',{headers,data:{run_id:preserved.run_id,expected_revision:preserved.result_revision,fields:{id_number:'110105194912310003X'}}});
  assert.equal(unreviewed.status(),200);assert.equal((await unreviewed.json()).export_ready,false);
  await page.reload();await page.getByRole('tab',{name:'任务批次'}).click();
  if(await header.getAttribute('aria-expanded')!=='true')await header.click();
  await batchPanel.getByText('1 个任务有异议',{exact:true}).waitFor();
  await batchPanel.locator('.el-checkbox').filter({hasText:'只看有异议'}).click();
  assert.equal(await batchPanel.locator('.el-table__body-wrapper tbody tr').count(),1);
  await page.screenshot({path:a.out+'/batch-dispute-filter.png'});
  await batchPanel.getByRole('button',{name:'核对',exact:true}).click();await editor.waitFor();
  assert.equal(await page.locator('.drawer-head').getByRole('button',{name:'导出 Excel',exact:true}).isDisabled(),true);
  assert.equal((await context.request.post(a.url+'/api/v1/export',{headers,data:{task_ids:[task.task_id]}})).status(),409);
  await page.screenshot({path:a.out+'/dispute-blocked.png'});
  const ordinary=page.waitForResponse(r=>r.url().endsWith('/review')&&r.request().method()==='POST');
  await editor.getByRole('button',{name:'保存修正并重新校验',exact:true}).click();assert.equal((await ordinary).status(),200);
  const confirmed=await (await context.request.get(a.url+'/api/v1/tasks/'+task.task_id+'/result')).json();
  assert.equal(confirmed.export_ready,true);assert.equal(confirmed.fields.find(f=>f.key==='id_number').value,'110105194912310003X');
  assert.equal(confirmed.generic.full_text,originalEvidence);
  await batchPanel.getByText('1 个任务有异议',{exact:true}).waitFor({state:'hidden'});
  await page.screenshot({path:a.out+'/dispute-confirmed.png'});
  report.checks.push('one_save_confirms_edited_and_unchanged_values_updates_batch_and_keeps_source');
  await idInput.fill('11010519491231002X');await lengthHint.waitFor({state:'hidden'});
  // The saved warning remains until the corrected draft is explicitly saved.
  assert.equal(await resultHint.count(),1);
  report.checks.push('id_length_live_and_saved_hints_preserve_values_and_original_evidence');
  stage='save';const saved=page.waitForResponse(r=>r.url().endsWith('/review')&&r.request().method()==='POST');
  await editor.getByRole('button',{name:'保存修正并重新校验',exact:true}).click();assert.equal((await saved).status(),200);
  const r=await context.request.get(a.url+'/api/v1/tasks/'+task.task_id+'/result');const result=await r.json();
  assert.equal(result.fields.find(f=>f.key==='name').value,'测试甲');assert.equal(result.fields.find(f=>f.key==='mobile_number').value,'13800138000');
  assert.equal(result.fields.find(f=>f.key==='id_number').value,'11010519491231002X');assert.equal(result.generic.full_text,originalEvidence);
  await resultHint.waitFor({state:'hidden'});
  const latest=await (await context.request.get(a.url+'/api/v1/tasks/'+task.task_id)).json();assert.equal(latest.export_ready,true);
  stage='export';const download=page.waitForEvent('download');
  await page.locator('.drawer-actions').getByRole('button',{name:'导出 Excel',exact:true}).click();
  await (await download).saveAs(a.out+'/interview-reviewed.xlsx');
  await page.screenshot({path:a.out+'/reviewed.png'});report.checks.push('saved_values_unlock_real_xlsx_download');
  for(const sample of a.number_cases){
   stage='number_hint_'+sample.kind;
   await page.locator('.el-drawer:visible .el-drawer__close-btn').click();await editor.waitFor({state:'hidden'});
   const panel=page.locator('.batch-collapse .el-collapse-item').filter({hasText:sample.batch_label}).first();
   const heading=panel.locator('.el-collapse-item__header');if(await heading.getAttribute('aria-expanded')!=='true')await heading.click();
   await panel.getByRole('button',{name:/^(查看|补填|核对)$/}).first().click();await editor.waitFor();
   const target=field(sample.label),hint=target.locator('.number-length-hint');await target.scrollIntoViewIfNeeded();await hint.waitFor();
   assert.match(await hint.innerText(),new RegExp(sample.expect));
   assert.equal(await target.locator('input').inputValue(),sample.original);
   await page.locator('.result-view .number-length-hint.warning').waitFor();
   await page.screenshot({path:a.out+'/'+sample.kind+'-length-hint.png'});
   await target.locator('input').fill(sample.corrected);await hint.waitFor({state:'hidden'});
   const corrected=page.waitForResponse(r=>r.url().endsWith('/review')&&r.request().method()==='POST');
   await editor.getByRole('button',{name:'保存修正并重新校验',exact:true}).click();assert.equal((await corrected).status(),200);
   const updated=await (await context.request.get(a.url+'/api/v1/tasks/'+sample.task_id+'/result')).json();
   const number=updated.fields.find(f=>f.key===sample.key);assert.equal(number.value,sample.corrected);assert.equal(number.source.raw_text,sample.original);
   await page.locator('.result-view .number-length-hint.warning').waitFor({state:'hidden'});
   report.checks.push(sample.kind+'_saved_and_live_number_hints_preserve_source');
  }
  stage='ownership';const other=await browser.newContext();const p=await other.newPage();await login(p,'interview_other');
  assert.equal((await other.request.get(a.url+'/api/v1/tasks/'+task.task_id+'/result')).status(),404);
  report.checks.push('anonymous_and_cross_account_access_denied');assert.equal(errors.length,0);report.status='passed';
 }catch(e){report.stage=stage;report.error=String(e.message).replaceAll(a.password,'[secret]').slice(0,260);if(page)await page.screenshot({path:a.out+'/failure.png'}).catch(()=>{});}
 finally {if(browser)await browser.close();fs.writeFileSync(a.out+'/browser.json',JSON.stringify(report,null,2));process.exitCode=report.status==='passed'?0:1;}
})();
'''


def main():
    import uvicorn
    from PIL import Image, ImageDraw, ImageFont
    from openpyxl import load_workbook
    from ocr_app.api import create_app
    from ocr_app.business import Business
    from ocr_app.extraction import extract_result
    from ocr_app.winjob import JobProcess
    from test_business_api import FakeRuntime
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--base',type=Path,required=True)
    parser.add_argument('--static',type=Path)
    args=parser.parse_args();base=args.base.resolve();out=base/('interview-browser-'+time.strftime('%H%M%S'));out.mkdir()
    business=Business(out/'private/data',key=secrets.token_bytes(32));password=secrets.token_urlsafe(24)
    for name in ('interview_qa','interview_other'):business.accounts.set_user(name,password,name,'operator')
    image=Image.new('RGB',(1200,800),'white');draw=ImageDraw.Draw(image)
    text='面试登记表（合成测试）\n基本信息 测试甲\n身份证姓名 11010519491231002X 性别 女 年龄 43\n身份证号码\n手机号码\n婚姻状况 未婚'
    draw.multiline_text((80,100),text,font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',28),fill='black',spacing=12)
    image.save(out/'synthetic.png');payload=(out/'synthetic.png').read_bytes()
    from ocr_app.doc_profiles import LABELS
    owner=next(user['user_id'] for user in business.accounts.users() if user['username']=='interview_qa')
    number_cases=[]
    for kind,key,original,corrected,expect in [
        ('BANK_CARD','card_number','621785'+'0'*12,'621785'+'0'*13,'少1位'),
        ('DIPLOMA','certificate_no','0'*19,'0'*18,'常见17或18位'),
        ('DEGREE_CERT','certificate_no','0'*17,'T'+'0'*16,'历史T/Z'),
        ('XUEXIN_DIPLOMA_TABLE','certificate_no','0'*19,'0'*17,'常见17或18位'),
        ('BUSINESS_LICENSE','unified_social_credit_code','9'*17,'9'*18,'少1位'),
    ]:
        created=business.store.create_batch([dict(data=payload,name='hint-'+kind+'.png',mime='image/png',actual_source_page_count=1)],kind,owner)
        task=business.store.claim_next();assert task['id']==created['task_ids'][0]
        page_id=business.store.save_page(task,1,payload,1200,800)
        result=extract_result([],kind);field=next(f for f in result['fields'] if f['key']==key)
        field.update(value=original,source=dict(page=1,raw_text=original),review_issues=[],need_review=False)
        result.update(status='NEED_REVIEW',pages=[dict(page_id=page_id,page_index=1,width=1200,height=800,extraction_method='ocr',regions=[])],warnings=[])
        assert business.store.publish(task,result)
        number_cases.append(dict(kind=kind,key=key,original=original,corrected=corrected,expect=expect,task_id=task['id'],batch_label=LABELS[kind],label=field['label']))
    class Runtime(FakeRuntime):
        def start(self):
            super().start()
            def loop():
                while self.alive:
                    task=business.store.claim_next()
                    if task:
                        page_id=business.store.save_page(task,1,payload,1200,800)
                        region=dict(region_id=str(uuid4()),bbox=[.05,.1,.95,.75],extraction_method='ocr',status='recognized')
                        result=extract_result([dict(text=text,page=1,**{k:v for k,v in region.items() if k!='status'})],'INTERVIEW_FORM')
                        result.update(status='NEED_REVIEW',pages=[dict(page_id=page_id,page_index=1,width=1200,height=800,extraction_method='ocr',regions=[region])],warnings=[])
                        business.store.publish(task,result)
                    time.sleep(.1)
            self.thread=threading.Thread(target=loop,daemon=True);self.thread.start()
        def close(self,*a,**kw):
            super().close();self.thread.join(timeout=5)
    runtime=Runtime();app=create_app(runtime,args.static.resolve() if args.static else base/'static',out/'private/session',business=business)
    sock=socket.socket();sock.bind(('127.0.0.1',0));url='http://127.0.0.1:'+str(sock.getsockname()[1])
    server=uvicorn.Server(uvicorn.Config(app,access_log=False,log_config=None,log_level='critical'))
    thread=threading.Thread(target=lambda:server.run(sockets=[sock]),daemon=True);thread.start();process=None
    try:
        for _ in range(100):
            if server.started:break
            time.sleep(.1)
        assert server.started
        process=JobProcess(['node','-e',BROWSER],cwd=ROOT)
        for stream in (process.stdout,process.stderr):threading.Thread(target=stream.read,daemon=True).start()
        process.stdin.write(json.dumps(dict(url=url,password=password,input=str(out/'synthetic.png'),out=str(out),number_cases=number_cases)));process.stdin.close()
        code=process.wait(timeout=150)
    finally:
        if process:process.close()
        server.should_exit=True;thread.join(timeout=10)
    report=json.loads((out/'browser.json').read_text());print(json.dumps(report,ensure_ascii=False));assert not thread.is_alive()
    if code:raise SystemExit(code)
    sheet=load_workbook(out/'interview-reviewed.xlsx')['面试登记表'];headers=[c.value for c in sheet[1]]
    for label,value in [('姓名','测试甲'),('身份证号','11010519491231002X'),('手机号','13800138000')]:
        cell=sheet.cell(2,headers.index(label)+1);assert cell.value==value and cell.data_type=='s'
    assert sheet.cell(1,headers.index('姓名')+1).fill.fgColor.rgb[-6:].lower()=='ef760e'
    print('XLSX saved values, text numbers and header style verified')


if __name__=='__main__':main()
