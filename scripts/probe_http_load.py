"""Isolated real-HTTP capacity probe; runtime is stubbed, never a GPU benchmark.

Run with .venvs/app/Scripts/python.exe scripts/probe_http_load.py
Evidence is retained under output/http-load-<timestamp>. No production data is read.
"""
from __future__ import annotations

import argparse
import asyncio
from collections import Counter, defaultdict
from datetime import datetime, timezone
import hashlib
import io
import json
import os
from pathlib import Path
import platform
import random
import shutil
import socket
import subprocess
import sys
import threading
import time
import zipfile

import httpx
import psutil
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]


def save(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def picture(width=32, height=32):
    stream = io.BytesIO()
    pixels = random.Random(42).randbytes(width * height * 3)
    Image.frombytes('RGB', (width, height), pixels).save(stream, 'JPEG', quality=85)
    return stream.getvalue()


class IdleRuntime:
    """Only replace inference; deliberately leave uploaded tasks QUEUED."""
    def start(self):
        self.alive = True

    def close(self, graceful=True):
        self.alive = False

    def is_alive(self):
        return self.alive

    def snapshot(self):
        return dict(configured_mode='dual_resident', effective_mode='dual_resident',
                    state='ready', models=dict(glm_ocr='ready', paddleocr_vl='ready'), reason_code=None)

    def request(self, *args, **kwargs):
        raise AssertionError('Inference is outside this HTTP probe')


def serve(out):
    trace_connections = json.loads((out / 'manifest.json').read_text('utf-8')).get('trace_connections', False)
    sys.path.insert(0, str(out / 'source' / 'backend'))
    import uvicorn
    from sqlalchemy import text
    from ocr_app.api import create_app
    from ocr_app.business import Business
    from ocr_app.diagnostics import init_diagnostics
    from ocr_app.extraction import extract_result
    from ocr_app.exporting import required_export_blockers

    key = os.urandom(32)
    (out / 'private' / 'test.key').write_bytes(key)
    business = Business(out / 'private' / 'data',
                        logger=init_diagnostics(out / 'diagnostics', 'api'), key=key)
    users = []
    password = 'synthetic-' + os.urandom(16).hex()
    images = [picture(), picture(33, 32)]  # Upload admission deduplicates identical bytes within a batch.
    for index in range(51):
        username = f'load{index:03d}'
        account = business.accounts.set_user(username, password, username,
                                             'admin' if index == 50 else 'operator')
        user = dict(account, password=password, batches=[])
        for batch_index in range(10 if index < 50 else 0):
            created = business.store.create_batch(
                [dict(data=image, name=f'synthetic-{n}.jpg', mime='image/jpeg',
                      actual_source_page_count=1) for n, image in enumerate(images)], 'DIPLOMA', account['user_id'])
            assert len(created['task_ids']) == 2, 'Fixture must contain two distinct documents'
            user['batches'].append(created['batch_id'])
            for task_id in created['task_ids']:
                result = extract_result([], 'DIPLOMA')
                values = dict(name=f'LoadUser{index:03d}', gender='男', birth_date='1990-01-01',
                              school_name='Synthetic University', major='计算机科学与技术',
                              education_level='本科', duration='4', enroll_date='2010-09-01',
                              graduate_date='2014-06-01', issue_date='2014-06-01',
                              certificate_no='100012014000000001')
                for field in result['fields']:
                    field.update(value=values[field['key']], need_review=False, review_issues=[])
                result.update(status='COMPLETED', review_issues=[])
                assert not required_export_blockers(result)
                with business.store.transaction() as c:
                    c.execute(text("UPDATE tasks SET status='COMPLETED',result_enc=:result WHERE id=:id"),
                              dict(result=business.storage.pack(result), id=task_id))
                    c.execute(text("UPDATE runs SET status='COMPLETED' WHERE task_id=:id"), dict(id=task_id))
            if batch_index == 0:
                user['task'] = business.store.task(created['task_ids'][0])
            if batch_index == 1:
                user['export_batch'] = created['batch_id']
        users.append(user)
        if index % 10 == 0:
            print(f'Seeded {index + 1}/51 accounts', flush=True)
    save(out / 'private' / 'fixtures.json', users)
    with business.engine.connect() as c:
        settings = {key: c.exec_driver_sql('PRAGMA ' + key).scalar()
                    for key in ['journal_mode', 'synchronous', 'busy_timeout']}
        settings['migration'] = c.exec_driver_sql('SELECT version_num FROM alembic_version').scalar()
    save(out / 'database-settings.json', dict(settings, pool=business.engine.pool.status()))
    sock = socket.socket()
    sock.bind(('127.0.0.1', 0))
    port = sock.getsockname()[1]
    app = create_app(IdleRuntime(), out / 'no-frontend', business.root, business, business.logger)
    @app.middleware('http')
    async def record_probe_arrival(request, call_next):
        identifier = request.headers.get('X-Load-Sequence', '')
        if trace_connections and identifier.isdigit() and len(identifier) < 12:
            with (out / 'arrivals.jsonl').open('a', encoding='utf-8') as arrival:
                arrival.write(json.dumps({'sequence': int(identifier), 'at': time.time(), 'port': request.client.port}) + '\n')
        return await call_next(request)
    from uvicorn.protocols.http.h11_impl import H11Protocol
    class TracedProtocol(H11Protocol):
        def timeout_keep_alive_handler(self):
            with (out / 'connection-closes.jsonl').open('a', encoding='utf-8') as stream:
                stream.write(json.dumps({'at': time.time(), 'port': self.client[1], 'reason': 'keepalive'}) + '\n')
            super().timeout_keep_alive_handler()
    server = uvicorn.Server(uvicorn.Config(app, log_level='error', access_log=False,
                                         http=TracedProtocol if trace_connections else 'auto'))

    def stop_when_requested():
        while not (out / 'STOP').exists():
            time.sleep(.2)
        server.should_exit = True

    threading.Thread(target=stop_when_requested, daemon=True).start()
    save(out / 'ready.json', dict(port=port, pid=os.getpid(), runtime='STUB_NO_INFERENCE'))
    try:
        server.run(sockets=[sock])
    finally:
        sock.close()
    # Reopen only this probe's isolated database after app shutdown.
    business = Business(out / 'private' / 'data', key=key)
    failures = []
    expected_path = out / 'expected.json'
    expected = json.loads(expected_path.read_text('utf-8')) if expected_path.exists() else {}
    for accepted in expected.get('uploads', []):
        for identifier in accepted['task_ids']:
            task = business.store.task(identifier)
            if task['status'] != 'QUEUED' or task['batch_id'] != accepted['batch_id']:
                failures.append('upload_identity_or_state')
        with business.engine.connect() as c:
            owner = c.execute(text('SELECT owner_id FROM batches WHERE id=:id'),
                              dict(id=accepted['batch_id'])).scalar_one()
            if owner != accepted['owner_id']:
                failures.append('upload_owner')
    for user in expected.get('reviews', []):
        result = business.store.result(user['task_id'])
        if result['result_revision'] != user['revision']:
            failures.append('review_revision')
        if user.get('name') and next(f['value'] for f in result['fields'] if f['key'] == 'name') != user['name']:
            failures.append('review_value')
    with business.engine.connect() as c:
        checks = dict(sqlite=c.exec_driver_sql('PRAGMA integrity_check').scalars().all(),
                      foreign_keys=[list(r) for r in c.exec_driver_sql('PRAGMA foreign_key_check')],
                      counts={table: c.exec_driver_sql('SELECT count(*) FROM ' + table).scalar()
                              for table in ['web_users', 'batches', 'tasks', 'audit']},
                      states=dict(c.exec_driver_sql('SELECT status,count(*) FROM tasks GROUP BY status').all()),
                      audit_events=dict(c.exec_driver_sql('SELECT event,count(*) FROM audit GROUP BY event').all()))
    uploaded = sum(len(row['task_ids']) for row in expected.get('uploads', []))
    if checks['counts']['tasks'] != 1000 + uploaded:
        failures.append('unacknowledged_or_missing_task')
    checks.update(failures=failures, accepted_uploads=uploaded,
                  reviews_checked=len(expected.get('reviews', [])))
    with business.engine.connect() as c:
        references = set(c.exec_driver_sql('SELECT reference FROM artifacts WHERE reference IS NOT NULL').scalars())
    blobs = {p.stem for p in business.storage.blobs.glob('*.bin')}
    corrupt = 0
    for reference in references & blobs:
        try:
            business.storage.read(reference)
        except Exception:
            corrupt += 1
    checks['artifacts'] = dict(references=len(references), files=len(blobs),
                              missing=len(references - blobs), orphaned=len(blobs - references), corrupt=corrupt)
    if references != blobs or corrupt:
        failures.append('artifact_integrity')
    save(out / 'integrity.json', checks)
    business.close()


def percentile(values, fraction):
    values = sorted(values)
    if not values:
        return None
    position = (len(values) - 1) * fraction
    low = int(position)
    return round(values[low] + (values[min(low + 1, len(values) - 1)] - values[low]) * (position - low), 2)


class Load:
    def __init__(self, out, base, client):
        self.out, self.base, self.client = out, base, client
        self.log = (out / 'requests.jsonl').open('w', encoding='utf-8', buffering=1)
        self.rows, self.phases = [], []
        self.expected = dict(uploads=[], reviews=[])
        self.epoch = time.perf_counter()
        self.epoch_wall = time.time()
        self.sequence = 0
        self.trace_connections = json.loads((out / 'manifest.json').read_text('utf-8')).get('trace_connections', False)

    async def request(self, phase, op, user, method, path, expected=(200,), **kwargs):
        started = time.perf_counter()
        response, error, status, transport_codes = None, None, 0, []
        self.sequence += 1
        sequence = self.sequence
        trace_events = []
        client_port = None
        async def trace(event, info):
            nonlocal client_port
            # Trace event names only: never headers, bodies, exception messages or URLs.
            if len(trace_events) < 40:
                trace_events.append(event)
            request = info.get('request')
            if request is not None:
                client_port = request.extensions.get('load_client_port', client_port)
        try:
            headers = dict(user['headers'])
            if self.trace_connections:
                headers['X-Load-Sequence'] = str(sequence)
                kwargs['extensions'] = {'trace': trace}
            response = await self.client.request(method, self.base + path, headers=headers, **kwargs)
            status = response.status_code
            if status not in expected:
                try:
                    error = response.json().get('code', 'HTTP_ERROR')
                except (ValueError, AttributeError):
                    error = 'HTTP_ERROR'
        except httpx.HTTPError as exception:
            error = type(exception).__name__
            current = exception
            seen = set()
            for _ in range(10):
                if current is None:
                    break
                if id(current) in seen:
                    break
                seen.add(id(current))
                transport_codes.append({'type': type(current).__name__})
                for key in ('errno', 'winerror'):
                    value = getattr(current, key, None)
                    if isinstance(value, int):
                        transport_codes.append({key: value})
                current = current.__cause__ or current.__context__
        row = dict(phase=phase, op=op, user=user['username'], start_s=round(started - self.epoch, 6),
                   ms=round((time.perf_counter() - started) * 1000, 3), status=status,
                   bytes=len(response.content) if response is not None else 0,
                   error=error, expected=status in expected)
        row['transport_codes'] = transport_codes
        row['sequence'] = sequence
        row['client_port'] = client_port
        row['start_epoch'] = self.epoch_wall + started - self.epoch
        if error is not None:
            row['transport_trace'] = trace_events
        row['controlled_rejection'] = status in {429, 503} and error in {'SERVICE_BUSY', 'OPERATION_THROTTLED'}
        row['timings_ms'] = {name: float(response.headers[header]) for name, header in [
            ('queue_wait', 'X-Queue-Wait-Ms'), ('db_wait', 'X-DB-Wait-Ms'),
            ('db_hold', 'X-DB-Hold-Ms'), ('file_write', 'X-File-Write-Ms')]
            if response is not None and header in response.headers}
        self.rows.append(row)
        self.log.write(json.dumps(row) + '\n')
        return response

    async def phase(self, name, users, work):
        print(f'PHASE {name} users={len(users)}', flush=True)
        start = time.perf_counter()
        start_epoch = time.time()
        await work(name, users)
        elapsed = time.perf_counter() - start
        rows = [r for r in self.rows if r['phase'] == name]
        record = dict(name=name, users=len(users), start_epoch=start_epoch, seconds=round(elapsed, 3), requests=len(rows),
                      rps=round(len(rows) / elapsed, 2), errors=sum(not r['expected'] for r in rows),
                      statuses=dict(Counter(str(r['status']) for r in rows)))
        record.update(successful=sum(200 <= r['status'] < 300 for r in rows),
                      expected_conflicts=sum(r['status'] == 409 and r['expected'] for r in rows),
                      controlled_rejections=sum(r['controlled_rejection'] for r in rows),
                      unexpected_failures=sum(not r['expected'] and not r['controlled_rejection'] for r in rows))
        record['timings_p95_ms'] = {key: percentile([r['timings_ms'][key] for r in rows if key in r['timings_ms']], .95)
                                   for key in ('queue_wait', 'db_wait', 'db_hold', 'file_write')}
        for label, fraction in [('p50', .5), ('p95', .95), ('p99', .99), ('max', 1)]:
            record[label + '_ms'] = percentile([r['ms'] for r in rows], fraction) if rows else None
        self.phases.append(record)
        save(self.out / 'phases.json', self.phases)
        print(json.dumps(record), flush=True)
        if psutil.virtual_memory().available < 2 * 1024**3:
            raise RuntimeError('Stopped: host available memory below 2 GiB')

    async def reads(self, phase, users, duration, pace):
        deadline = time.perf_counter() + duration

        async def reader(index, user):
            await asyncio.sleep(index * pace / len(users))
            routes = [('batches', '/api/v1/batches?page_size=50'),
                      ('batch_tasks', '/api/v1/batches/' + user['batches'][0] + '/tasks'),
                      ('result', '/api/v1/tasks/' + user['task']['task_id'] + '/result'),
                      ('health', '/api/v1/health')]
            while time.perf_counter() < deadline:
                start = time.perf_counter()
                for op, path in routes:
                    await self.request(phase, op, user, 'GET', path)
                await asyncio.sleep(max(0, pace - (time.perf_counter() - start)))
        await asyncio.gather(*(reader(i, u) for i, u in enumerate(users)))

    async def burst(self, phase, users, op, rounds):
        payload = picture(640, 480) if op == 'upload' else None

        async def action(user, round_index):
            if op == 'upload':
                r = await self.request(phase, op, user, 'POST', '/api/v1/batches', expected=(202,),
                                       data={'doc_type_hint': 'DIPLOMA'},
                                       files={'files': ('synthetic.jpg', payload, 'image/jpeg')})
                if r is not None and r.status_code == 202:
                    self.expected['uploads'].append(dict(r.json(), owner_id=user['user_id']))
            elif op == 'review':
                task = user['task']
                name = f"LoadUser{user['username']}-{phase}-{round_index}"
                r = await self.request(phase, op, user, 'POST', '/api/v1/tasks/' + task['task_id'] + '/review',
                                       json=dict(run_id=task['run_id'], expected_revision=task['result_revision'],
                                                 fields={'name': name}, note='Synthetic capacity probe'))
                if r is not None and r.status_code == 200:
                    task['result_revision'] += 1
                    user['last_name'] = name
            elif op == 'export':
                r = await self.request(phase, op, user, 'POST', '/api/v1/export',
                                       json=dict(batch_id=user['export_batch'], masked=True))
                if r is not None and r.status_code == 200:
                    from openpyxl import load_workbook
                    def validate():
                        with zipfile.ZipFile(io.BytesIO(r.content)) as book:
                            assert book.testzip() is None
                        book = load_workbook(io.BytesIO(r.content), read_only=True)
                        assert any(sheet.max_row >= 3 for sheet in book), 'Expected two material rows plus header'
                        book.close()
                    await asyncio.to_thread(validate)
            else:
                await self.request(phase, 'audit', user, 'GET', '/api/v1/admin/audit/events?page_size=50')
        for round_index in range(rounds):
            await asyncio.gather(*(action(user, round_index) for user in users))

    async def collision(self, phase, users):
        user = users[0]
        task = user['task']
        body = dict(run_id=task['run_id'], expected_revision=task['result_revision'], note='Concurrent revision probe')
        replies = await asyncio.gather(*(self.request(phase, 'revision_collision', user, 'POST',
            '/api/v1/tasks/' + task['task_id'] + '/review', expected=(200, 409), json=body) for _ in range(50)))
        statuses = Counter(r.status_code if r is not None else 0 for r in replies)
        save(self.out / 'collision.json', dict(statuses=dict(statuses), passed=statuses == {200: 1, 409: 49}))
        task['result_revision'] += statuses[200]
        assert statuses == {200: 1, 409: 49}, statuses

    def finish(self, users):
        self.expected['reviews'] = [dict(task_id=u['task']['task_id'], revision=u['task']['result_revision'],
                                         name=u.get('last_name')) for u in users[:50]]
        save(self.out / 'expected.json', self.expected)
        groups = defaultdict(list)
        for row in self.rows:
            groups[(row['phase'], row['op'])].append(row)
        save(self.out / 'operations.json', [dict(phase=p, operation=op, requests=len(rows),
            errors=sum(not r['expected'] for r in rows), p95_ms=percentile([r['ms'] for r in rows], .95),
            error_codes=dict(Counter(r['error'] for r in rows if r['error']))) for (p, op), rows in groups.items()])
        self.log.close()


async def exercise(out, port, smoke, writes_only=False):
    # Diagnostic-only hook in this probe process, never a library/file modification.
    from httpcore._async.http11 import AsyncHTTP11Connection
    original_headers = AsyncHTTP11Connection._send_request_headers
    async def traced_headers(connection, request):
        address = connection._network_stream.get_extra_info('client_addr')
        request.extensions['load_client_port'] = address[1] if address else None
        return await original_headers(connection, request)
    manifest = json.loads((out / 'manifest.json').read_text('utf-8'))
    if manifest.get('trace_connections'):
        AsyncHTTP11Connection._send_request_headers = traced_headers
    users = json.loads((out / 'private' / 'fixtures.json').read_text('utf-8'))
    base = f'http://127.0.0.1:{port}'
    async with httpx.AsyncClient(trust_env=False, timeout=30,
                                 limits=httpx.Limits(max_connections=100, max_keepalive_connections=100,
                                     keepalive_expiry=manifest['client_keepalive_s'])) as client:
        for user in users:
            client.cookies.clear()  # Never log out the previous user's session during setup.
            r = await client.post(base + '/api/v1/auth/login', headers={'Origin': base},
                                  json=dict(username=user['username'], password=user.pop('password')))
            assert r.status_code == 200, ('Login preflight status', r.status_code)
            user['headers'] = {'Origin': base, 'Cookie': 'ocr_session=' + r.cookies['ocr_session'],
                               'X-CSRF-Token': r.json()['csrf_token']}
        client.cookies.clear()
        load = Load(out, base, client)
        async def heartbeat():
            while True:
                await load.request('background', 'health', users[0], 'GET', '/api/v1/health')
                await asyncio.sleep(1)
        monitor = asyncio.create_task(heartbeat())
        try:
            # Untimed functional checks distinguish fixture/harness failures from load failures.
            r = await load.request('preflight', 'isolation', users[0], 'GET',
                                   '/api/v1/tasks/' + users[1]['task']['task_id'], expected=(404,))
            assert r.status_code == 404
            await load.phase('smoke-read', users[:1], lambda p, u: load.reads(p, u, 1, 1))
            for op in ['upload', 'review', 'export']:
                await load.phase('smoke-' + op, users[:1], lambda p, u, op=op: load.burst(p, u, op, 1))
            await load.phase('smoke-audit', users[50:], lambda p, u: load.burst(p, u, 'audit', 1))
            assert not any(r['error'] for r in load.rows), 'Functional preflight failed'
            if not smoke:
                if writes_only:
                    await load.phase('connection-warmup-50', users[:50], lambda p, u: load.reads(p, u, 1, 0))
                else:
                    for n in [20, 50]:
                        await load.phase(f'paced-read-{n}', users[:n], lambda p, u: load.reads(p, u, 60, 2))
                    await load.phase('saturated-read-50', users[:50], lambda p, u: load.reads(p, u, 30, 0))
                for op in ['upload', 'review', 'export']:
                    for n in [20, 50]:
                        await load.phase(f'{op}-{n}', users[:n],
                                         lambda p, u, op=op: load.burst(p, u, op, 2 if op == 'export' else 3))
                await load.phase('audit-50', [users[50]] * 50, lambda p, u: load.burst(p, u, 'audit', 2))
                await load.phase('revision-collision-50', [users[0]] * 50, load.collision)
        finally:
            monitor.cancel()
            try:
                await monitor
            except asyncio.CancelledError:
                pass
            load.finish(users)
            if manifest.get('trace_connections'):
                AsyncHTTP11Connection._send_request_headers = original_headers


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--serve', type=Path, help=argparse.SUPPRESS)
    parser.add_argument('--smoke', action='store_true', help='Run functional smoke only, with an independent fixture')
    parser.add_argument('--writes-only', action='store_true', help='Focused write/connection regression; omit full sustained reading phases')
    parser.add_argument('--client-keepalive', type=float, default=4, help='Expire idle client connections before the server default of 5s; use 5 to reproduce the original race')
    parser.add_argument('--trace-connections', action='store_true', help='Diagnostic-only numeric connection/arrival instrumentation')
    args = parser.parse_args()
    if args.serve:
        serve(args.serve.resolve())
        return
    stamp = datetime.now().strftime('%Y%m%d-%H%M%S-%f')
    out = ROOT / 'output' / ('http-load-' + stamp)
    (out / 'private').mkdir(parents=True)
    print('OUTPUT ' + str(out), flush=True)
    shutil.copytree(ROOT / 'backend', out / 'source' / 'backend', ignore=shutil.ignore_patterns('__pycache__', '*.pyc'))
    shutil.copy2(__file__, out / 'probe_http_load.py')
    sources = {str(p.relative_to(out / 'source')).replace('\\', '/'): hashlib.sha256(p.read_bytes()).hexdigest()
               for p in (out / 'source').rglob('*') if p.is_file()}
    manifest = dict(created_at=datetime.now(timezone.utc).isoformat(), platform=platform.platform(),
                    python=platform.python_version(), cpu_logical=psutil.cpu_count(),
                    cpu_physical=psutil.cpu_count(logical=False), memory=psutil.virtual_memory()._asdict(),
                    commit=subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip(),
                    working_tree=subprocess.check_output(['git', 'status', '--porcelain'], cwd=ROOT, text=True),
                    source_sha256=sources, runtime='STUB_NO_INFERENCE', seed_users=51,
                    seed_batches=500, seed_tasks=1000, upload_bytes=len(picture(640, 480)),
                    paced_cycle='4 sequential reads per user every 2s; closed loop; staggered start',
                    timeout_s=30, http_connections=100, client_keepalive_s=args.client_keepalive,
                    server_keepalive_s=5, trace_connections=args.trace_connections, smoke=args.smoke, writes_only=args.writes_only)
    save(out / 'manifest.json', manifest)
    stop_sampling = threading.Event()
    with (out / 'server.log').open('w', encoding='utf-8') as log:
        server = subprocess.Popen([sys.executable, str(out / 'probe_http_load.py'), '--serve', str(out)],
                                   stdout=log, stderr=subprocess.STDOUT, cwd=out,
                                   creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
        def sample():
            process, generator = psutil.Process(server.pid), psutil.Process()
            observed = {process.pid: process}
            with (out / 'resources.jsonl').open('w', encoding='utf-8', buffering=1) as stream:
                while not stop_sampling.wait(1):
                    try:
                        descendants = [process, *process.children(recursive=True)]
                        active = [observed.setdefault(p.pid, p) for p in descendants]
                        stream.write(json.dumps(dict(at=time.time(), server_cpu=sum(p.cpu_percent() for p in active),
                            server_pids=[p.pid for p in active], server_rss=sum(p.memory_info().rss for p in active), generator_cpu=generator.cpu_percent(),
                            generator_rss=generator.memory_info().rss, host_cpu=psutil.cpu_percent(),
                            available=psutil.virtual_memory().available)) + '\n')
                    except psutil.NoSuchProcess:
                        return
        sampler = threading.Thread(target=sample, daemon=True)
        sampler.start()
        outcome = dict(completed=False)
        try:
            deadline = time.monotonic() + 240
            while not (out / 'ready.json').exists():
                if server.poll() is not None:
                    raise RuntimeError('Test server exited; inspect server.log')
                if time.monotonic() > deadline:
                    raise TimeoutError('Fixture startup exceeded 240 seconds')
                time.sleep(.5)
            ready = json.loads((out / 'ready.json').read_text('utf-8'))
            for _ in range(50):
                try:
                    r = httpx.get(f"http://127.0.0.1:{ready['port']}/health", trust_env=False, timeout=2)
                    if r.status_code == 200:
                        break
                except httpx.HTTPError:
                    pass
                time.sleep(.2)
            asyncio.run(exercise(out, ready['port'], args.smoke, args.writes_only))
            outcome['completed'] = True
        except Exception as error:
            outcome['failure'] = type(error).__name__ + ': ' + str(error)[:300]
            print(outcome['failure'], flush=True)
        finally:
            (out / 'STOP').touch()
            try:
                server.wait(timeout=40)
            except subprocess.TimeoutExpired:
                server.terminate()  # Only the process created above; never a process-name kill.
                server.wait(timeout=10)
                outcome['forced_stop'] = True
            stop_sampling.set()
            sampler.join(timeout=3)
            outcome['server_exit_code'] = server.returncode
            if (out / 'integrity.json').exists():
                integrity = json.loads((out / 'integrity.json').read_text('utf-8'))
                phases = json.loads((out / 'phases.json').read_text('utf-8')) if (out / 'phases.json').exists() else []
                requests = [json.loads(line) for line in (out / 'requests.jsonl').read_text('utf-8').splitlines()]
                outcome['unexpected_requests'] = sum(not row['expected'] for row in requests)
                outcome['passed'] = (outcome['completed'] and server.returncode == 0
                    and integrity['sqlite'] == ['ok'] and not integrity['foreign_keys']
                    and not integrity['failures'] and not outcome['unexpected_requests']
                    and all(p['errors'] == 0 for p in phases))
            save(out / 'outcome.json', outcome)
    print('DONE ' + json.dumps(outcome), flush=True)
    if not outcome['completed'] or server.returncode:
        sys.exit(1)
    if not outcome.get('passed'):
        sys.exit(2)  # A completed measurement can still fail capacity/integrity acceptance.


if __name__ == '__main__':
    main()
